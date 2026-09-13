"""Provider-neutral 10x5 XiaoAn model/judge matrix runner.

The module keeps provider transport separate from evaluation aggregation.  API
credentials are read from environment variables and are never written to the
public workbook.  Provider failures are retained as UNAVAILABLE observations.
"""

from __future__ import annotations
from .oracle_judge import contract as oracle_contract, validate as validate_oracle, summarize as summarize_oracles, response_schema as oracle_response_schema

from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, is_dataclass
import fcntl
import hashlib
import json
import os
from pathlib import Path
import re
from statistics import median
import threading
import time
from typing import Any, Callable, Mapping, Sequence

from .coverage import case_coverage, coverage_rows
from .rules import RatingRule
from .scoring import SCORING_CONTRACT_VERSION, TurnQuality, score_case
from .methodology_metrics import robust_dimension_summary, self_judging
from .methodology_runtime import build_matrix_summaries, run_attribution_pass

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None


PROVIDERS = ("claude", "gpt", "gemini", "qwen", "kimi")
DEFAULT_MODEL_PAIRS = {
    "claude": ("claude-opus-5", "claude-sonnet-5"),
    "gpt": ("gpt-5.6-sol", "o4-mini-2025-04-16"),
    "gemini": ("gemini-3-pro-preview-thinking", "gemini-3.8-flash"),
    "qwen": ("qwen3.8-max", "qwen3.7-max"),
    "kimi": ("kimi-k3", "kimi-k2.6"),
}
JUDGE_EVIDENCE_SCHEMA_VERSION = "judge-evidence/v1"
MATRIX_JUDGE_REQUEST_SCHEMA_VERSION = "matrix-judge-request/v3"
MATRIX_JUDGE_PROMPT_SCHEMA_VERSION = "matrix-judge-prompt/v2"
JUDGE_CONTEXT_LAYERS = frozenset({"CAPSULE", "WIKI", "SOURCE"})


class JudgeEvidenceError(ValueError):
    """The compact projection cannot substantiate trace claims safely."""


@dataclass(frozen=True)
class ModelSpec:
    provider: str
    model: str
    tier: str
    reasoning_effort: str | None = None

    @property
    def id(self) -> str:
        return f"{self.provider}:{self.model}:{self.tier}:{self.reasoning_effort or 'na'}"


@dataclass(frozen=True)
class ProviderResponse:
    provider: str
    model: str
    status: str
    text: str = ""
    first_char: str = ""
    first_character_ms: float | None = None
    elapsed_ms: float | None = None
    input_tokens: int | None = None
    output_tokens: int | None = None
    total_tokens: int | None = None
    request_id: str | None = None
    error: str | None = None
    trace: Mapping[str, Any] | None = None
    cached_input_tokens: int | None = None
    cache_write_tokens: int | None = None
    cache_hit_ratio: float | None = None
    queue_ms: float = 0.0
    error_type: str | None = None
    attempt_count: int = 1
    retry_errors: tuple[str, ...] = ()


def _provider_response_from_mapping(value: Mapping[str, Any]) -> ProviderResponse:
    """Restore JSON checkpoint data to the canonical response value types."""
    data = dict(value)
    retry_errors = data.get("retry_errors", ())
    if not isinstance(retry_errors, Sequence) or isinstance(retry_errors, (str, bytes)):
        raise ValueError("retry_errors must be an array")
    data["retry_errors"] = tuple(str(error) for error in retry_errors)
    return ProviderResponse(**data)


def default_subject_specs() -> tuple[ModelSpec, ...]:
    """Return the requested 10 subject rows; override model names via env vars."""
    specs: list[ModelSpec] = []
    for provider in PROVIDERS:
        latest, second = DEFAULT_MODEL_PAIRS[provider]
        latest = os.getenv(f"XIAOAN_{provider.upper()}_LATEST_MODEL", latest)
        second = os.getenv(f"XIAOAN_{provider.upper()}_SECOND_MODEL", second)
        specs.extend((ModelSpec(provider, latest, "latest", "medium"), ModelSpec(provider, second, "second", "high")))
    return tuple(specs)


def default_judge_specs() -> tuple[ModelSpec, ...]:
    """Return one current judge per provider, aligned with the latest subject tier."""
    return tuple(
        ModelSpec(
            provider,
            os.getenv(
                f"XIAOAN_{provider.upper()}_JUDGE_MODEL",
                os.getenv(
                    f"XIAOAN_{provider.upper()}_LATEST_MODEL",
                    DEFAULT_MODEL_PAIRS[provider][0],
                ),
            ),
            "judge",
            "medium",
        )
        for provider in PROVIDERS
    )


def _usage(raw: Mapping[str, Any]) -> tuple[int | None, int | None, int | None]:
    usage = raw.get("usage", raw.get("usageMetadata", {}))
    if not isinstance(usage, Mapping):
        return None, None, None
    def number(*names: str) -> int | None:
        for name in names:
            value = usage.get(name)
            if isinstance(value, int):
                return value
        return None
    prompt = number("prompt_tokens", "input_tokens", "promptTokenCount")
    completion = number("completion_tokens", "output_tokens", "candidatesTokenCount")
    total = number("total_tokens", "totalTokenCount")
    if total is None and prompt is not None and completion is not None:
        total = prompt + completion
    return prompt, completion, total


def _cache_usage(raw: Mapping[str, Any], prompt_tokens: int | None) -> tuple[int | None, int | None, float | None]:
    usage = raw.get("usage", raw.get("usageMetadata", {}))
    if not isinstance(usage, Mapping):
        return None, None, None
    details = usage.get("input_tokens_details", usage.get("prompt_tokens_details", {}))
    details = details if isinstance(details, Mapping) else {}
    cached = details.get(
        "cached_tokens",
        usage.get("cache_read_input_tokens", usage.get("cachedContentTokenCount")),
    )
    cache_write = details.get(
        "cache_write_tokens",
        usage.get("cache_creation_input_tokens"),
    )
    cached = cached if isinstance(cached, int) and not isinstance(cached, bool) else None
    cache_write = cache_write if isinstance(cache_write, int) and not isinstance(cache_write, bool) else None
    ratio = cached / prompt_tokens if cached is not None and prompt_tokens else None
    return cached, cache_write, ratio


def normalize_response(provider: str, model: str, raw: Mapping[str, Any], elapsed_ms: float) -> ProviderResponse:
    text = raw.get("text")
    if not isinstance(text, str):
        content = raw.get("content")
        if isinstance(content, Sequence):
            text = "".join(str(item.get("text", "")) for item in content if isinstance(item, Mapping))
        choices = raw.get("choices")
        if not isinstance(text, str) and isinstance(choices, Sequence) and choices and isinstance(choices[0], Mapping):
            message = choices[0].get("message", choices[0])
            text = message.get("content", "") if isinstance(message, Mapping) else ""
        elif not isinstance(text, str):
            candidates = raw.get("candidates")
            text = (((candidates[0] or {}).get("content") or {}).get("parts") or [{}])[0].get("text", "") if candidates else ""
    text = text if isinstance(text, str) else str(text or "")
    prompt, completion, total = _usage(raw)
    cached, cache_write, cache_ratio = _cache_usage(raw, prompt)
    trace = raw.get("chatflow_debug")
    first_character = raw.get("_xiaoan_first_character_ms", raw.get("first_character_ms"))
    if isinstance(first_character, bool) or not isinstance(first_character, (int, float)) or first_character < 0:
        first_character = None
    return ProviderResponse(
        provider=provider,
        model=model,
        status="PASS",
        text=text,
        first_char=text[:1],
        first_character_ms=float(first_character) if first_character is not None else None,
        elapsed_ms=elapsed_ms,
        input_tokens=prompt,
        output_tokens=completion,
        total_tokens=total,
        request_id=str(raw.get("id") or raw.get("responseId") or "") or None,
        trace=dict(trace) if isinstance(trace, Mapping) else None,
        cached_input_tokens=cached,
        cache_write_tokens=cache_write,
        cache_hit_ratio=cache_ratio,
        attempt_count=int(raw.get("_xiaoan_attempt_count", 1) or 1),
        retry_errors=tuple(str(item) for item in raw.get("_xiaoan_retry_errors", ()) or ()),
    )


def invoke(spec: ModelSpec, prompt: str, *, transport: Callable[[ModelSpec, str], Mapping[str, Any]] | None = None, queue_ms: float = 0.0) -> ProviderResponse:
    started = time.perf_counter()
    try:
        if transport is None:
            raise RuntimeError("no provider transport configured; pass a transport or use the provider plugin")
        raw = transport(spec, prompt)
        if not isinstance(raw, Mapping):
            raise RuntimeError("provider response must be a JSON object")
        response = normalize_response(spec.provider, spec.model, raw, (time.perf_counter() - started) * 1000)
        response = ProviderResponse(**{**asdict(response), "queue_ms": queue_ms})
        if not response.text.strip():
            return ProviderResponse(**{**asdict(response), "status": "UNAVAILABLE", "error": "provider returned empty text", "error_type": "EMPTY_RESPONSE"})
        return response
    except Exception as exc:  # noqa: BLE001 - preserve operational state in matrix
        text = f"{type(exc).__name__}: {exc}"
        lowered = text.lower()
        retry_errors = tuple(getattr(exc, "retry_errors", ()))
        error_type = (
            "RATE_LIMIT" if "429" in lowered
            else "PROVIDER_5XX" if any(item in lowered for item in ("500", "502", "503", "504"))
            else "TIMEOUT" if "timeout" in lowered or "timed out" in lowered
            else "CONNECTION" if retry_errors or any(item in lowered for item in ("disconnect", "connection", "network", "urlerror"))
            else "PROVIDER_ERROR"
        )
        return ProviderResponse(
            spec.provider, spec.model, "UNAVAILABLE",
            elapsed_ms=(time.perf_counter() - started) * 1000,
            error=text, queue_ms=queue_ms, error_type=error_type,
            attempt_count=int(getattr(exc, "attempt_count", 1)),
            retry_errors=retry_errors,
        )


def weighted_score(scores: Mapping[str, Any], rule: RatingRule, quality_focus: Sequence[str] = ()) -> float | None:
    if not scores:
        return None
    return score_case(rule, [TurnQuality(scores)], quality_focus).weighted_total


def matrix_pair_summary(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Case-macro weighted quality from complete cases; shared by both renderers."""
    cases: dict[str, list[Mapping[str, Any]]] = {}
    for row in rows:
        cases.setdefault(str(row.get("case_id")), []).append(row)
    scores = []
    for case_rows in cases.values():
        expected = set().union(*(set(row.get("expected_turns", ())) for row in case_rows))
        observed = [row.get("turn") for row in case_rows]
        if len(observed) != len(set(observed)) or (expected and set(observed) != expected):
            continue
        if not all(row.get("status") == "PASS" and row.get("primary_eligible", True)
                   and isinstance(row.get("weighted_score"), (int, float)) for row in case_rows):
            continue
        score = 0.0 if any(row.get("triggered_red_lines") for row in case_rows) else sum(row["weighted_score"] for row in case_rows) / len(case_rows)
        scores.append(score)
    return {"value": sum(scores) / len(scores) if scores else None,
            "eligible_cases": len(scores), "attempted_cases": len(cases),
            "status": "AVAILABLE" if scores else "UNAVAILABLE",
            "aggregation": "case_macro_dynamic_weighted_mean"}


def _field(value: Any, name: str, default: Any = None) -> Any:
    """Read a field from legacy mappings and the typed case dataclasses."""
    if isinstance(value, Mapping):
        return value.get(name, default)
    return getattr(value, name, default)


def _json_value(value: Any) -> Any:
    if is_dataclass(value) and not isinstance(value, type):
        return asdict(value)
    if isinstance(value, Mapping):
        return dict(value)
    return value


def _canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )


def _selected_fields(value: Any, names: Sequence[str]) -> dict[str, Any]:
    if not isinstance(value, Mapping):
        return {}
    return {name: value[name] for name in names if name in value}


def _composer_invocation(trace: Mapping[str, Any]) -> Mapping[str, Any]:
    snapshot = trace.get("effective_context_snapshot")
    if not isinstance(snapshot, Mapping):
        return {}
    invocations = snapshot.get("invocations")
    if isinstance(invocations, Mapping) and isinstance(invocations.get("composer"), Mapping):
        return invocations["composer"]
    composer = snapshot.get("composer")
    return composer if isinstance(composer, Mapping) else {}


def _compact_context_units(trace: Mapping[str, Any]) -> list[dict[str, Any]]:
    invocation = _composer_invocation(trace)
    raw_units = invocation.get("context_units", invocation.get("units", ()))
    if not isinstance(raw_units, Sequence) or isinstance(raw_units, (str, bytes)):
        return []
    units: list[dict[str, Any]] = []
    for raw in raw_units:
        if not isinstance(raw, Mapping):
            continue
        layer = str(raw.get("layer", ""))
        field_path = str(raw.get("field_path", raw.get("unit_id", "")))
        keep_safety_prompt = layer == "PROMPT" and field_path == "safety_message"
        if layer not in JUDGE_CONTEXT_LAYERS and not keep_safety_prompt:
            continue
        if raw.get("inclusion_state", "EXPOSED") != "EXPOSED":
            continue
        unit = _selected_fields(
            raw,
            (
                "ref", "layer", "entity_id", "field_path", "item_id", "content",
                "content_hash", "content_sha256", "parent_ref", "policy_ids",
            ),
        )
        units.append(unit)
    return units


def _normalized_ground_ref(value: Any) -> str:
    ref = str(value or "")
    if ref.startswith("composer:"):
        ref = ref[len("composer:"):]
    if ref.startswith("knowledge/wiki/nodes/"):
        ref = ref[len("knowledge/wiki/nodes/"):]
    return ref


def _ground_item_refs(item: Mapping[str, Any]) -> set[str]:
    content = item.get("content")
    semantic_text = item.get("text")
    if isinstance(content, Mapping):
        semantic_text = content.get("text", content.get("content"))
    elif isinstance(content, str):
        semantic_text = content
    if not isinstance(semantic_text, str) or not semantic_text.strip():
        return set()
    refs = {
        _normalized_ground_ref(item.get("ref")),
        _normalized_ground_ref(item.get("entity_id")),
    }
    if isinstance(content, Mapping):
        refs.add(_normalized_ground_ref(content.get("ref")))
    return {ref for ref in refs if ref}


def build_judge_evidence(trace: Mapping[str, Any] | None) -> dict[str, Any]:
    """Project a full Chatflow trace into small, semantic Judge evidence.

    This is deliberately an allowlist. Provider requests, router candidate
    catalogs, prompts, retries, timings, token counts, and response IDs remain
    in the private checkpoint but never reach a Judge.
    """
    if trace is None:
        return {
            "schema_version": JUDGE_EVIDENCE_SCHEMA_VERSION,
            "trace_sha256": None,
            "trace_status": "ABSENT",
            "route": {},
            "safety": {},
            "capsule": {},
            "ground": {},
            "state": {},
            "output_guard": {},
            "context": {"snapshot_id": None, "context_kind": None, "status": "ABSENT", "units": []},
            "evidence_integrity": {"status": "NOT_APPLICABLE", "unresolved_refs": []},
        }
    if not isinstance(trace, Mapping):
        raise JudgeEvidenceError("chatflow trace must be a mapping")

    snapshot = trace.get("effective_context_snapshot")
    snapshot = snapshot if isinstance(snapshot, Mapping) else {}
    invocation = _composer_invocation(trace)
    units = _compact_context_units(trace)
    ground_units = [unit for unit in units if str(unit.get("layer")) in {"WIKI", "SOURCE"}]
    ground = trace.get("ground")
    ground = ground if isinstance(ground, Mapping) else {}
    raw_refs = ground.get("resolved_ground", ground.get("resolved_refs", ()))
    if isinstance(raw_refs, (str, bytes)):
        if raw_refs:
            raise JudgeEvidenceError("resolved ground refs must be an array")
        resolved_refs = []
    elif isinstance(raw_refs, Sequence):
        resolved_refs = [str(item) for item in raw_refs]
    else:
        raise JudgeEvidenceError("resolved ground refs must be an array")
    raw_items = ground.get("resolved_items", ())
    resolved_items = (
        [dict(item) for item in raw_items if isinstance(item, Mapping)]
        if isinstance(raw_items, Sequence) and not isinstance(raw_items, (str, bytes))
        else []
    )
    evidence_items = ground_units or resolved_items
    available_refs = set().union(
        *(_ground_item_refs(item) for item in evidence_items),
        set(),
    )
    unresolved = [ref for ref in resolved_refs if _normalized_ground_ref(ref) not in available_refs]
    if unresolved:
        raise JudgeEvidenceError("unresolved ground refs: " + ", ".join(unresolved))

    return {
        "schema_version": JUDGE_EVIDENCE_SCHEMA_VERSION,
        "trace_sha256": "sha256:" + hashlib.sha256(_canonical_json(trace).encode("utf-8")).hexdigest(),
        "trace_status": "AVAILABLE",
        "route": _selected_fields(
            trace.get("route"),
            ("capsule_id", "capsule_title", "confidence", "method", "reason", "fallback_reason", "should_continue_active_capsule"),
        ),
        "safety": _selected_fields(trace.get("safety"), ("level", "reason", "triggered_rules", "response_key")),
        "capsule": _selected_fields(trace.get("capsule"), ("id", "version")),
        "ground": {
            **_selected_fields(ground, ("loaded", "policy_reason", "warnings")),
            "resolved_refs": resolved_refs,
            "evidence_items": [] if ground_units else resolved_items,
        },
        "state": _selected_fields(trace.get("state"), ("active_capsule_id", "ttl_turns")),
        "output_guard": _selected_fields(trace.get("output_guard"), ("passed", "warnings")),
        "context": {
            "snapshot_id": snapshot.get("snapshot_id"),
            "context_kind": snapshot.get("context_kind"),
            "status": invocation.get("status", "ABSENT"),
            "units": units,
        },
        "evidence_integrity": {"status": "COMPLETE", "unresolved_refs": []},
    }


def _answer_id(subject: ModelSpec, case_id: str, turn: int) -> str:
    identity = _canonical_json({"subject_id": subject.id, "case_id": case_id, "turn": turn})
    return "answer:" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24]


def _judge_prompt(stable: Mapping[str, Any], dynamic: Mapping[str, Any]) -> str:
    """Encode a provider-neutral stable-prefix/dynamic-suffix contract."""
    digest = hashlib.sha256(_canonical_json(stable).encode("utf-8")).hexdigest()
    return json.dumps({
        "xiaoan_prompt_contract": MATRIX_JUDGE_PROMPT_SCHEMA_VERSION,
        "prompt_cache_key": f"xiaoan-matrix-judge-v1:{digest[:16]}",
        "stable_prefix": stable,
        "dynamic_input": dynamic,
    }, ensure_ascii=False, default=str)


def _spec_id(value: Mapping[str, Any]) -> str:
    if value.get("id"):
        return str(value["id"])
    return ModelSpec(
        str(value.get("provider", "")),
        str(value.get("model", "")),
        str(value.get("tier", "")),
        str(value["reasoning_effort"]) if value.get("reasoning_effort") else None,
    ).id


def _checkpoint_state(path: Path, contract_sha256: str, *, allow_legacy: bool) -> tuple[dict[str, dict[str, Any]], dict[tuple[str, str], dict[str, Any]], dict[tuple[str, str], dict[str, Any]], dict[str, dict[str, Any]]]:
    answers: dict[str, dict[str, Any]] = {}
    judgements: dict[tuple[str, str], dict[str, Any]] = {}
    cells: dict[tuple[str, str], dict[str, Any]] = {}
    attributions: dict[str, dict[str, Any]] = {}
    truncate_at: int | None = None
    with path.open(encoding="utf-8") as source:
        line_number = 0
        while True:
            offset = source.tell()
            line = source.readline()
            if not line:
                break
            line_number += 1
            if not line.strip():
                continue
            try:
                event = json.loads(line)
            except json.JSONDecodeError as exc:
                if not line.endswith("\n"):
                    truncate_at = offset
                    break
                raise ValueError(f"invalid matrix checkpoint JSON at line {line_number}") from exc
            if not isinstance(event, Mapping):
                raise ValueError(f"invalid matrix checkpoint event at line {line_number}")
            stored_contract = event.get("contract_sha256")
            if stored_contract is None and not allow_legacy:
                raise ValueError("legacy matrix checkpoint has no contract hash; migrate it or opt in explicitly")
            if stored_contract is not None and stored_contract != contract_sha256:
                raise ValueError("matrix checkpoint contract does not match cases, models, or rating rule")
            event_type = event.get("event")
            if event_type == "answer" and isinstance(event.get("answer"), Mapping):
                subject = event.get("subject", {})
                if not isinstance(subject, Mapping):
                    continue
                answer_id = str(event.get("answer_id") or _answer_id(
                    ModelSpec(str(subject.get("provider", "")), str(subject.get("model", "")), str(subject.get("tier", "")), subject.get("reasoning_effort")),
                    str(event.get("case_id", "")),
                    int(event.get("turn", 0)),
                ))
                for key in tuple(judgements):
                    if key[0] == answer_id:
                        del judgements[key]
                for key in tuple(cells):
                    if key[0] == answer_id:
                        del cells[key]
                attributions.pop(answer_id, None)
                answers[answer_id] = {**dict(event), "answer_id": answer_id}
            elif event_type == "judgement" and isinstance(event.get("judgement"), Mapping):
                judge = event.get("judge", {})
                subject = event.get("subject", {})
                if not isinstance(judge, Mapping) or not isinstance(subject, Mapping):
                    continue
                answer_id = str(event.get("answer_id") or _answer_id(
                    ModelSpec(str(subject.get("provider", "")), str(subject.get("model", "")), str(subject.get("tier", "")), subject.get("reasoning_effort")),
                    str(event.get("case_id", "")),
                    int(event.get("turn", 0)),
                ))
                judgements[(answer_id, _spec_id(judge))] = {**dict(event), "answer_id": answer_id}
            elif event_type == "cell" and isinstance(event.get("row"), Mapping):
                row = dict(event["row"])
                subject = row.get("subject", {})
                judge = row.get("judge", {})
                if not isinstance(subject, Mapping) or not isinstance(judge, Mapping):
                    continue
                answer_id = str(row.get("answer_id") or _answer_id(
                    ModelSpec(str(subject.get("provider", "")), str(subject.get("model", "")), str(subject.get("tier", "")), subject.get("reasoning_effort")),
                    str(row.get("case_id", "")),
                    int(row.get("turn", 0)),
                ))
                cells[(answer_id, _spec_id(judge))] = {**row, "answer_id": answer_id}
            elif event_type == "attribution" and event.get("answer_id"):
                answer_id = str(event["answer_id"])
                attributions[answer_id] = {
                    key: value for key, value in event.items()
                    if key not in {"event", "task_id", "completed_at", "contract_sha256"}
                }
    if truncate_at is not None:
        with path.open("r+b") as target:
            target.truncate(truncate_at)
    return answers, judgements, cells, attributions


def _resume_row(answer_event: Mapping[str, Any], judgement_event: Mapping[str, Any] | None, cell: Mapping[str, Any]) -> dict[str, Any]:
    if isinstance(cell.get("answer"), Mapping) and isinstance(cell.get("judgement"), Mapping):
        row = dict(cell)
        row["answer"] = asdict(_provider_response_from_mapping(cell["answer"]))
        row["judgement"] = asdict(_provider_response_from_mapping(cell["judgement"]))
        return row
    answer = asdict(_provider_response_from_mapping(answer_event["answer"]))
    judged_data = dict(judgement_event["judgement"]) if judgement_event else {
        "provider": cell.get("judge", {}).get("provider", ""),
        "model": cell.get("judge", {}).get("model", ""),
        "status": cell.get("judgement_status", "UNAVAILABLE"),
        "error": cell.get("judgement_error"),
        "error_type": cell.get("judgement_error_type"),
    }
    judged_data["status"] = cell.get("judgement_status", judged_data.get("status", "UNAVAILABLE"))
    if cell.get("judgement_error") is not None:
        judged_data["error"] = cell.get("judgement_error")
    if cell.get("judgement_error_type") is not None:
        judged_data["error_type"] = cell.get("judgement_error_type")
    judged = asdict(_provider_response_from_mapping(judged_data))
    return {
        "oracle_assessment": cell.get("oracle_assessment"),
        "answer_id": cell["answer_id"],
        "case_id": cell.get("case_id"),
        "turn": cell.get("turn"),
        "subject": dict(cell.get("subject", {})),
        "judge": dict(cell.get("judge", {})),
        "answer": answer,
        "judgement": judged,
        "scores": dict(cell.get("scores", {})),
        "expected_red_line_ids": tuple(cell.get("expected_red_line_ids", ())),
        "triggered_red_lines": tuple(cell.get("triggered_red_lines", ())),
        "red_line_evidence": {
            str(key): tuple(value) for key, value in dict(cell.get("red_line_evidence", {})).items()
        },
        "self_judging": bool(cell.get("self_judging", False)),
        "primary_eligible": bool(cell.get("primary_eligible", True)),
        "coverage": dict(cell.get("coverage", {})),
        "memory_metrics": list(cell.get("memory_metrics", ())),
        "attribution": dict(cell.get("attribution", {})),
        "weighted_score": cell.get("weighted_score"),
        "expected_turns": list(cell.get("expected_turns", ())),
        "scoring_contract_version": cell.get("scoring_contract_version"),
        "status": cell.get("status", "UNAVAILABLE"),
    }


def _validated_scores(judged: ProviderResponse, rating_rule: RatingRule) -> tuple[ProviderResponse, dict[str, Any], tuple[str, ...], dict[str, tuple[str, ...]]]:
    if judged.status != "PASS":
        return judged, {}, (), {}
    try:
        parsed = json.loads(judged.text)
        if not isinstance(parsed, Mapping):
            raise ValueError("judge response must be a JSON object")
        dimensions = parsed.get("dimensions", parsed.get("scores", {}))
        if isinstance(dimensions, Mapping):
            scores = dict(dimensions)
        elif isinstance(dimensions, Sequence) and not isinstance(dimensions, (str, bytes)):
            scores = {
                str(item.get("module")): item.get("score")
                for item in dimensions
                if isinstance(item, Mapping) and item.get("module") is not None
            }
        else:
            scores = {}
        expected_modules = {module.name for module in rating_rule.modules}
        if set(scores) != expected_modules or any(
            isinstance(value, bool) or not isinstance(value, int) or value not in {0, 1, 2, 3}
            for value in scores.values()
        ):
            raise ValueError("judge dimensions must exactly match rating rule with integer scores 0-3")
        red_lines = parsed.get("red_lines")
        if not isinstance(red_lines, Sequence) or isinstance(red_lines, (str, bytes)):
            raise ValueError("judge red_lines must be an array")
        expected_red_lines = {item.id for item in rating_rule.red_lines}
        normalized: dict[str, bool] = {}
        evidence_by_id: dict[str, tuple[str, ...]] = {}
        for item in red_lines:
            if not isinstance(item, Mapping) or not isinstance(item.get("id"), str) or not isinstance(item.get("triggered"), bool):
                raise ValueError("judge red_lines entries require id and boolean triggered")
            if item["id"] in normalized:
                raise ValueError("judge red_lines IDs must be unique")
            evidence = item.get("evidence")
            if not isinstance(evidence, Sequence) or isinstance(evidence, (str, bytes)) or any(not isinstance(value, str) for value in evidence):
                raise ValueError("judge red_lines evidence must be an array of strings")
            normalized[item["id"]] = item["triggered"]
            evidence_by_id[item["id"]] = tuple(evidence)
            if item["triggered"] and not evidence:
                raise ValueError("triggered judge red_lines require non-empty evidence")
        if set(normalized) != expected_red_lines:
            raise ValueError("judge red_lines must exactly match rating rule IDs")
        triggered = tuple(sorted(identifier for identifier, hit in normalized.items() if hit))
        if triggered:
            scores = {module.name: 0 for module in rating_rule.modules}
        return judged, scores, triggered, evidence_by_id
    except (json.JSONDecodeError, ValueError):
        return ProviderResponse(
            provider=judged.provider,
            model=judged.model,
            status="UNAVAILABLE",
            elapsed_ms=judged.elapsed_ms,
            input_tokens=judged.input_tokens,
            output_tokens=judged.output_tokens,
            total_tokens=judged.total_tokens,
            request_id=judged.request_id,
            first_character_ms=judged.first_character_ms,
            error="judge response was not valid JSON",
            cached_input_tokens=judged.cached_input_tokens,
            cache_write_tokens=judged.cache_write_tokens,
            cache_hit_ratio=judged.cache_hit_ratio,
            queue_ms=judged.queue_ms,
            error_type="INVALID_JUDGE_JSON",
            attempt_count=judged.attempt_count,
            retry_errors=judged.retry_errors,
        ), {}, (), {}


def _memory_observations(case: Any, turn: int, trace: Mapping[str, Any] | None) -> list[dict[str, Any]]:
    from .cases import MemoryCheckpoint
    from .memory_metrics import memory_observation

    output = []
    for item in tuple(_field(case, "memory_checkpoints", ()) or ()):
        if int(_field(item, "after_turn", 0)) != turn:
            continue
        checkpoint = MemoryCheckpoint(
            turn, tuple(_field(item, "facts", ()) or ()), str(_field(item, "usage", "")),
            str(_field(item, "check_type", _field(item, "type", "use"))),
        )
        observation = memory_observation(checkpoint, trace or {})
        if hasattr(case, "oracle_gate_eligible") and not case.oracle_gate_eligible:
            observation = {"status": "skip", "reason": "memory oracle is not reviewed", "check_type": checkpoint.check_type}
        output.append(observation)
    return output


def _string_sequence(value: Any) -> tuple[str, ...] | None:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        return None
    return tuple(str(item) for item in value)


def run_matrix(cases: Sequence[Any], subjects: Sequence[ModelSpec], judges: Sequence[ModelSpec], *, subject_transport: Callable[[ModelSpec, str], Mapping[str, Any]], judge_transport: Callable[[ModelSpec, str], Mapping[str, Any]], rating_rule: RatingRule, checkpoint_path: Path | None = None, resume: bool = False, retry_unavailable: bool = False, allow_legacy_checkpoint: bool = False, subject_concurrency: int = 1, judge_concurrency: int = 1, max_in_flight: int | None = None, per_provider_concurrency: int = 1, attribution_provider: Callable[[Mapping[str, Any]], str] | None = None, attribution_judge_version: str = "attribution-judge/v1", isolate_self_judging: bool = True) -> list[dict[str, Any]]:
    """Run the matrix while appending durable answer/judgement/cell checkpoints.

    Checkpoints are JSONL events, intentionally separate from the public
    workbook so an interrupted or provider-failed run can be audited or
    rebuilt without exposing credentials.
    """
    if subject_concurrency < 1 or judge_concurrency < 1:
        raise ValueError("subject_concurrency and judge_concurrency must be at least 1")
    if max_in_flight is not None and max_in_flight < 1:
        raise ValueError("max_in_flight must be at least 1")
    if per_provider_concurrency < 1:
        raise ValueError("per_provider_concurrency must be at least 1")
    if retry_unavailable and not resume:
        raise ValueError("retry_unavailable requires resume")
    case_list = tuple(cases)
    subject_list = tuple(subjects)
    judge_list = tuple(judges)
    for spec in (*subject_list, *judge_list):
        if spec.provider not in PROVIDERS:
            raise ValueError(f"provider must be one of {PROVIDERS} using canonical lowercase: {spec.provider}")
    contract_sha256 = "sha256:" + hashlib.sha256(_canonical_json({
        "judge_evidence_schema": JUDGE_EVIDENCE_SCHEMA_VERSION,
        "judge_request_schema": MATRIX_JUDGE_REQUEST_SCHEMA_VERSION,
        "judge_prompt_schema": MATRIX_JUDGE_PROMPT_SCHEMA_VERSION,
        "cases": [_json_value(case) for case in case_list],
        "subjects": [asdict(subject) for subject in subject_list],
        "judges": [asdict(judge) for judge in judge_list],
        "rating_rule": _json_value(rating_rule),
        "attribution": {"enabled": attribution_provider is not None, "judge_version": attribution_judge_version},
        "isolate_self_judging": isolate_self_judging,
        "scoring_contract_version": SCORING_CONTRACT_VERSION,
    }).encode("utf-8")).hexdigest()
    prior_answers: dict[str, dict[str, Any]] = {}
    prior_judgements: dict[tuple[str, str], dict[str, Any]] = {}
    prior_cells: dict[tuple[str, str], dict[str, Any]] = {}
    prior_attributions: dict[str, dict[str, Any]] = {}
    checkpoint = None
    if checkpoint_path is not None:
        checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            checkpoint = checkpoint_path.open("a" if resume else "x", encoding="utf-8")
        except FileExistsError as exc:
            raise ValueError(f"matrix checkpoint already exists; use a new output directory: {checkpoint_path}") from exc
    if checkpoint is not None:
        try:
            fcntl.flock(checkpoint.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            checkpoint.close()
            raise RuntimeError(f"matrix checkpoint is already locked by another process: {checkpoint_path}") from exc
        if resume and checkpoint_path is not None and checkpoint_path.stat().st_size:
            prior_answers, prior_judgements, prior_cells, prior_attributions = _checkpoint_state(
                checkpoint_path, contract_sha256, allow_legacy=allow_legacy_checkpoint,
            )
    if checkpoint_path is not None:
        os.chmod(checkpoint_path, 0o600)

    checkpoint_lock = threading.Lock()
    global_gate = threading.BoundedSemaphore(max_in_flight or max(subject_concurrency, judge_concurrency))
    provider_gates = {
        provider: threading.BoundedSemaphore(per_provider_concurrency)
        for provider in {spec.provider for spec in (*subject_list, *judge_list)}
    }

    def savepoint(event: Mapping[str, Any]) -> None:
        if checkpoint is None:
            return
        payload = {"completed_at": time.time(), "contract_sha256": contract_sha256, **dict(event)}
        with checkpoint_lock:
            checkpoint.write(json.dumps(payload, ensure_ascii=False, default=str) + "\n")
            checkpoint.flush()
            os.fsync(checkpoint.fileno())

    def limited_invoke(spec: ModelSpec, prompt: str, transport: Callable[[ModelSpec, str], Mapping[str, Any]]) -> ProviderResponse:
        queued = time.perf_counter()
        with provider_gates[spec.provider], global_gate:
            queue_ms = (time.perf_counter() - queued) * 1000
            return invoke(spec, prompt, transport=transport, queue_ms=queue_ms)

    def generate_lane(lane: tuple[int, int, ModelSpec, Any]) -> tuple[int, list[tuple[str, dict[str, Any]]]]:
        """Phase 1: finish one stateful subject/case lane without invoking judges."""
        lane_index, _case_index, subject, case = lane
        turns = tuple(_field(case, "turns", ()) or ())
        case_id = str(_field(case, "id", ""))
        lane_id = f"lane:{subject.id}:{case_id}"
        answer_ids = [_answer_id(subject, case_id, int(_field(turn, "turn", 0))) for turn in turns]
        completed = [answer_id in prior_answers for answer_id in answer_ids]
        retry_subject_lane = bool(
            retry_unavailable
            and any(completed)
            and (
                not all(completed)
                or any(
                    _provider_response_from_mapping(prior_answers[answer_id]["answer"]).status != "PASS"
                    for answer_id in answer_ids
                    if answer_id in prior_answers
                )
            )
        )
        if retry_subject_lane:
            for turn, answer_id in zip(turns, answer_ids):
                invalidated = ProviderResponse(
                    subject.provider,
                    subject.model,
                    "UNAVAILABLE",
                    error="subject lane retry has not completed",
                    error_type="RETRY_IN_PROGRESS",
                )
                savepoint({
                    "event": "answer",
                    "task_id": f"{lane_id}:turn:{int(_field(turn, 'turn', 0))}:retry-invalidation",
                    "answer_id": answer_id,
                    "case_id": case_id,
                    "turn": int(_field(turn, "turn", 0)),
                    "subject": asdict(subject),
                    "answer": asdict(invalidated),
                })
        start_case = getattr(subject_transport, "start_case", None)
        if callable(start_case) and any(completed) and not all(completed) and not retry_subject_lane:
            raise ValueError(f"cannot safely resume partially completed stateful case {case_id} for {subject.id}")
        start_error: str | None = None
        if callable(start_case) and (retry_subject_lane or not all(completed)):
            try:
                with provider_gates[subject.provider], global_gate:
                    start_case(subject, case_id)
            except Exception as exc:
                start_error = f"{type(exc).__name__}: {exc}"
        generated: list[tuple[str, dict[str, Any]]] = []
        try:
            for turn in turns:
                turn_number = int(_field(turn, "turn", 0))
                answer_id = _answer_id(subject, case_id, turn_number)
                if answer_id in prior_answers and not retry_subject_lane:
                    continue
                prompt = str(_field(turn, "user", ""))
                answer = (
                    ProviderResponse(subject.provider, subject.model, "UNAVAILABLE", error=start_error, error_type="START_CASE")
                    if start_error
                    else limited_invoke(subject, prompt, subject_transport)
                )
                event = {
                    "event": "answer", "task_id": f"{lane_id}:turn:{turn_number}",
                    "answer_id": answer_id, "case_id": case_id, "turn": turn_number,
                    "subject": asdict(subject), "answer": asdict(answer),
                }
                savepoint(event)
                generated.append((answer_id, event))
        finally:
            end_case = getattr(subject_transport, "end_case", None)
            if callable(end_case) and (retry_subject_lane or not all(completed)):
                end_case(subject, case_id)
        return lane_index, generated

    def run_lane(lane: tuple[int, int, ModelSpec, Any]) -> tuple[int, list[dict[str, Any]]]:
        lane_index, _case_index, subject, case = lane
        lane_rows: list[dict[str, Any]] = []
        turns = tuple(_field(case, "turns", ()) or ())
        case_id = str(_field(case, "id", ""))
        lane_id = f"lane:{subject.id}:{case_id}"
        start_case = getattr(subject_transport, "start_case", None)
        start_error: str | None = None
        answer_ids = [_answer_id(subject, case_id, int(_field(turn, "turn", 0))) for turn in turns]
        completed_answers = [answer_id in prior_answers for answer_id in answer_ids]
        if callable(start_case) and any(completed_answers) and not all(completed_answers):
            raise ValueError(f"cannot safely resume partially completed stateful case {case_id} for {subject.id}")
        if callable(start_case) and not all(completed_answers):
            try:
                with provider_gates[subject.provider], global_gate:
                    start_case(subject, case_id)
            except Exception as exc:  # preserve provider/runtime failures in output
                start_error = f"{type(exc).__name__}: {exc}"
        history: list[dict[str, str]] = []
        try:
            for turn in turns:
                prompt = str(_field(turn, "user", ""))
                turn_number = int(_field(turn, "turn", 0))
                answer_id = _answer_id(subject, case_id, turn_number)
                prior_answer = prior_answers.get(answer_id)
                if prior_answer is not None:
                    answer = _provider_response_from_mapping(prior_answer["answer"])
                    answer_data = asdict(answer)
                else:
                    answer = ProviderResponse(subject.provider, subject.model, "UNAVAILABLE", error=start_error, error_type="START_CASE") if start_error else limited_invoke(subject, prompt, subject_transport)
                    answer_data = asdict(answer)
                    savepoint({"event": "answer", "task_id": f"{lane_id}:turn:{turn_number}", "answer_id": answer_id, "case_id": case_id, "turn": turn_number, "subject": asdict(subject), "answer": answer_data})

                memory_metrics = _memory_observations(case, turn_number, answer.trace)
                attribution = {"answer_id": answer_id, "status": "NOT_RUN", "result": None}
                if attribution_provider is not None:
                    attribution = (
                        {} if answer_id in regenerated_answer_ids
                        else prior_attributions.get(answer_id, {})
                    )
                    if not attribution and answer_id not in regenerated_answer_ids:
                        attribution = next((
                            dict(cell.get("attribution", {}))
                            for (cell_answer_id, _), cell in prior_cells.items()
                            if cell_answer_id == answer_id and cell.get("attribution")
                        ), {})
                    if not attribution:
                        attribution = run_attribution_pass(
                            [{"answer_id": answer_id, "answer": answer_data}],
                            attribution_provider,
                            judge_version=attribution_judge_version,
                        )[0]
                        savepoint({"event": "attribution", "task_id": f"{lane_id}:turn:{turn_number}:attribution", **attribution})

                judge_prompt: str | None = None
                evidence_error: str | None = None
                if answer.status == "PASS":
                    try:
                        judge_evidence = build_judge_evidence(answer.trace)
                        stable_judge_request = {
                            "schema_version": MATRIX_JUDGE_REQUEST_SCHEMA_VERSION,
                            "instruction": "Return JSON only with red_lines and dimensions. Return every supplied red-line ID exactly once as {id, triggered, evidence}; evaluate red lines first. Return every module exactly once with an integer 0, 1, 2, or 3. If any red line is triggered, all dimensions must be 0.",
                            "score_scale": [asdict(anchor) for anchor in rating_rule.score_scale],
                            "red_lines": [asdict(item) for item in rating_rule.red_lines],
                            "modules": [asdict(module) for module in rating_rule.modules],
                        }
                        dynamic_judge_request = {
                            "case_id": _field(case, "id"), "turn": _field(turn, "turn"),
                            "user": prompt, "answer": answer.text, "history": list(history),
                            "quality_focus": list(_field(case, "quality_focus", ()) or ()),
                            "judge_evidence": judge_evidence,
                            "expected": _json_value(_field(turn, "expected")),
                        }
                        oracle_request = {"assistant_answer": answer.text, "redacted_user_input": prompt, "redacted_conversation_history": list(history), "expected": dynamic_judge_request["expected"]}
                        dynamic_judge_request["oracle_contract"] = oracle_contract(oracle_request)
                        stable_judge_request["oracle_response_schema"] = oracle_response_schema()
                        stable_judge_request["instruction"] += " Also return oracle_assessment matching oracle_response_schema and oracle_contract instructions."
                        judge_prompt = _judge_prompt(stable_judge_request, dynamic_judge_request)
                    except JudgeEvidenceError as exc:
                        evidence_error = f"JudgeEvidenceError: {exc}"
                        savepoint({"event": "judge_evidence_error", "task_id": f"{lane_id}:turn:{turn_number}:evidence", "answer_id": answer_id, "case_id": case_id, "turn": turn_number, "error": evidence_error})

                turn_rows: dict[int, dict[str, Any]] = {}

                def record_judgement(index: int, judge: ModelSpec, judged: ProviderResponse, *, provider_called: bool) -> None:
                    if provider_called:
                        savepoint({"event": "judgement", "task_id": f"{lane_id}:turn:{turn_number}:judge:{judge.id}", "answer_id": answer_id, "case_id": case_id, "turn": turn_number, "subject": asdict(subject), "judge": asdict(judge), "judgement": asdict(judged), "dag_node_id": f"judge:{answer_id}:{judge.id}", "parent_node_ids": [f"answer:{answer_id}"]})
                    judged, scores, triggered_red_lines, red_line_evidence = _validated_scores(judged, rating_rule)
                    try:
                        oracle_value = validate_oracle(json.loads(judged.text).get("oracle_assessment") if judged.status == "PASS" else None, {"assistant_answer": answer.text, "redacted_user_input": prompt, "redacted_conversation_history": list(history), "expected": _json_value(_field(turn, "expected"))})
                    except (ValueError, TypeError) as exc:
                        oracle_value = {"status": "UNAVAILABLE", "items": [], "reason": str(exc)}
                    is_self_judging = self_judging(asdict(subject), asdict(judge))
                    row = {
                        "answer_id": answer_id, "case_id": case_id, "turn": turn_number,
                        "subject": {**asdict(subject), "id": subject.id},
                        "judge": {**asdict(judge), "id": judge.id},
                        "answer": answer_data, "judgement": asdict(judged), "scores": scores,
                        "oracle_assessment": oracle_value,
                        "expected_red_line_ids": tuple(item.id for item in rating_rule.red_lines),
                        "triggered_red_lines": triggered_red_lines,
                        "red_line_evidence": red_line_evidence,
                        "self_judging": is_self_judging,
                        "primary_eligible": not (isolate_self_judging and is_self_judging),
                        "coverage": case_coverage(case) if hasattr(case, "oracle_gate_eligible") else {},
                        "memory_metrics": memory_metrics,
                        "attribution": attribution,
                        "weighted_score": weighted_score(scores, rating_rule, _field(case, "quality_focus", ()) or ()),
                        "expected_turns": [int(_field(item, "turn")) for item in _field(case, "turns", ())],
                        "scoring_contract_version": SCORING_CONTRACT_VERSION,
                        "status": "PASS" if answer.status == judged.status == "PASS" else "UNAVAILABLE",
                    }
                    turn_rows[index] = row
                    savepoint({"event": "cell", "task_id": f"{lane_id}:turn:{turn_number}:cell:{judge.id}", "dag_node_id": f"cell:{answer_id}:{judge.id}", "parent_node_ids": [f"answer:{answer_id}", f"judge:{answer_id}:{judge.id}"], "row": {"answer_id": answer_id, "case_id": case_id, "turn": turn_number, "subject": row["subject"], "judge": row["judge"], "judgement_status": judged.status, "judgement_error": judged.error, "judgement_error_type": judged.error_type, "oracle_assessment": row["oracle_assessment"], "scores": row["scores"], "expected_red_line_ids": row["expected_red_line_ids"], "triggered_red_lines": row["triggered_red_lines"], "red_line_evidence": row["red_line_evidence"], "self_judging": row["self_judging"], "primary_eligible": row["primary_eligible"], "coverage": row["coverage"], "memory_metrics": row["memory_metrics"], "attribution": row["attribution"], "weighted_score": row["weighted_score"], "expected_turns": row["expected_turns"], "scoring_contract_version": SCORING_CONTRACT_VERSION, "status": row["status"]}})

                pending: list[tuple[int, ModelSpec]] = []
                for index, judge in enumerate(judge_list):
                    prior_cell = prior_cells.get((answer_id, judge.id))
                    invalidate_prior_cell = bool(
                        answer_id in regenerated_answer_ids
                        or (
                            retry_unavailable
                            and prior_cell is not None
                            and prior_cell.get("status") != "PASS"
                            and answer.status == "PASS"
                            and evidence_error is None
                        )
                    )
                    prior_judgement = prior_judgements.get((answer_id, judge.id))
                    if prior_cell is not None and not invalidate_prior_cell:
                        turn_rows[index] = _resume_row(prior_answer or prior_answers[answer_id], prior_judgements.get((answer_id, judge.id)), prior_cell)
                        turn_rows[index]["answer"] = answer_data
                    elif prior_judgement is not None and not invalidate_prior_cell:
                        prior_judged = _provider_response_from_mapping(prior_judgement["judgement"])
                        record_judgement(index, judge, prior_judged, provider_called=False)
                    elif answer.status != "PASS":
                        record_judgement(index, judge, ProviderResponse(judge.provider, judge.model, "UNAVAILABLE", error="subject response unavailable", error_type="SUBJECT_UNAVAILABLE"), provider_called=False)
                    elif evidence_error is not None:
                        record_judgement(index, judge, ProviderResponse(judge.provider, judge.model, "UNAVAILABLE", error=evidence_error, error_type="EVIDENCE_ERROR"), provider_called=False)
                    else:
                        pending.append((index, judge))

                workers = min(judge_concurrency, len(pending))
                if workers == 1:
                    for index, judge in pending:
                        record_judgement(index, judge, limited_invoke(judge, judge_prompt or "", judge_transport), provider_called=True)
                elif workers > 1:
                    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="xiaoan-judge") as executor:
                        futures = {executor.submit(limited_invoke, judge, judge_prompt or "", judge_transport): (index, judge) for index, judge in pending}
                        for future in as_completed(futures):
                            index, judge = futures[future]
                            record_judgement(index, judge, future.result(), provider_called=True)
                lane_rows.extend(turn_rows[index] for index in range(len(judge_list)))
                if answer.status == "PASS":
                    history.append({"user": prompt, "assistant": answer.text})
        finally:
            end_case = getattr(subject_transport, "end_case", None)
            if callable(end_case) and not all(completed_answers):
                end_case(subject, case_id)
        return lane_index, lane_rows

    try:
        lanes = [
            (subject_index * len(case_list) + case_index, case_index, subject, case)
            for subject_index, subject in enumerate(subject_list)
            for case_index, case in enumerate(case_list)
        ]
        # Global phase barrier: no Judge or attribution provider is invoked
        # until every subject answer artifact is complete or unavailable.
        if subject_concurrency == 1:
            generated_lanes = [generate_lane(lane) for lane in lanes]
        else:
            with ThreadPoolExecutor(max_workers=subject_concurrency, thread_name_prefix="xiaoan-subject") as executor:
                generated_lanes = list(executor.map(generate_lane, lanes))
        generated_lanes.sort(key=lambda item: item[0])
        for _, generated in generated_lanes:
            for answer_id, event in generated:
                prior_answers[answer_id] = event
        regenerated_answer_ids = {
            answer_id for _, generated in generated_lanes for answer_id, _ in generated
        }

        if subject_concurrency == 1:
            completed = [run_lane(lane) for lane in lanes]
        else:
            with ThreadPoolExecutor(max_workers=subject_concurrency, thread_name_prefix="xiaoan-lane") as executor:
                completed = list(executor.map(run_lane, lanes))
        completed.sort(key=lambda item: item[0])
        return [row for _, lane_rows in completed for row in lane_rows]
    finally:
        if checkpoint is not None:
            checkpoint.close()


def render_matrix_report(rows: Sequence[Mapping[str, Any]]) -> str:
    subjects = sorted({str(row["subject"]["id"]) for row in rows})
    judges = sorted({str(row["judge"]["id"]) for row in rows})
    unique_rows = list({str(row.get("answer_id")): row for row in rows}.values())
    summaries = build_matrix_summaries(
        rows,
        memory_results=[item for row in unique_rows for item in row.get("memory_metrics", ())],
        attribution_results=[row.get("attribution", {}) for row in unique_rows if row.get("attribution", {}).get("status") != "NOT_RUN"],
    )
    by_pair: dict[tuple[str, str], list[float]] = {}
    isolated_pairs = {
        (str(row["subject"]["id"]), str(row["judge"]["id"]))
        for row in rows if row.get("self_judging") and not row.get("primary_eligible", True)
    }
    for row in rows:
        score = row.get("weighted_score")
        if row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(score, (int, float)):
            by_pair.setdefault((str(row["subject"]["id"]), str(row["judge"]["id"])), []).append(float(score))
    lines = ["# Multimodel XiaoAn evaluation", "", "Case-macro dynamic weighted score (0-3); incomplete or unavailable cases are excluded. Dimension tables show turn-score medians.", "", "| XiaoAn subject \\ Judge | " + " | ".join(judges) + " |", "| --- | " + " | ".join("---:" for _ in judges) + " |"]
    for subject in subjects:
        cells = []
        for judge in judges:
            values = by_pair.get((subject, judge), [])
            summary = matrix_pair_summary([row for row in rows if str(row["subject"]["id"]) == subject and str(row["judge"]["id"]) == judge])
            cells.append(f"{summary['value']:.4f}" if summary["value"] is not None else "SELF_ISOLATED" if (subject, judge) in isolated_pairs else "UNAVAILABLE")
        lines.append("| " + subject + " | " + " | ".join(cells) + " |")
    denominator = summaries["primary_denominator"]
    lines.extend([
        "", "## Measurement contract", "",
        f"Primary eligible: {denominator['eligible_n']}; self-judging isolated: {denominator['self_excluded_n']}; operationally unavailable: {denominator['operationally_unavailable_n']}.",
        "Semantic oracle: " + json.dumps(summaries["semantic_oracle"], ensure_ascii=False),
        f"Memory: {summaries['memory']['status']} (eligible={summaries['memory']['eligible_n']}, missing={summaries['memory']['missing_n']}).",
        f"Dedicated attribution: {summaries['attribution']['status']} (eligible={summaries['attribution']['eligible_n']}, missing={summaries['attribution']['missing_n']}).",
        "Agreement statistics are DESCRIPTIVE_ONLY and do not establish correctness. Missing ranks are not imputed; constant or insufficient data yield UNAVAILABLE alpha/W. Different subjects may have different judge panels after self-exclusion or provider failures, so cross-subject comparisons are exploratory.",
    ])
    coverage = coverage_rows([row.get("coverage", {}) for row in rows])
    lines.extend(["", "## Oracle coverage", "", "| Metric | Reviewed units | Status | Detail |", "| --- | ---: | --- | --- |"] + [f"| {item['metric']} | {item['value']} | {item['status']} | {item['interpretation']} |" for item in coverage])
    self_rows = [row for row in rows if row.get("self_judging")]
    if self_rows:
        lines.extend(["", "## Self-judging cells", "", "| Subject | Judge | Case | Turn | Weighted score | Status |", "| --- | --- | --- | ---: | ---: | --- |"])
        for row in self_rows:
            lines.append(f"| {row['subject']['id']} | {row['judge']['id']} | {row.get('case_id')} | {row.get('turn')} | {row.get('weighted_score')} | {row.get('status')} |")
    dimensions = sorted({str(name) for row in rows for name in row.get("scores", {})})
    if dimensions:
        lines.extend(["", "## Dimension medians by subject", "", "| Subject | " + " | ".join(dimensions) + " |", "| --- | " + " | ".join("---:" for _ in dimensions) + " |"])
        for subject in subjects:
            cells = []
            for dimension in dimensions:
                values = [row["scores"].get(dimension) for row in rows if row["subject"]["id"] == subject and row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(row.get("scores", {}).get(dimension), (int, float))]
                cells.append(f"{median(values):.4f}" if values else "UNAVAILABLE")
            lines.append("| " + subject + " | " + " | ".join(cells) + " |")
        lines.extend(["", "## Dimension medians by subject and judge", "", "| Subject | Judge | " + " | ".join(dimensions) + " |", "| --- | --- | " + " | ".join("---:" for _ in dimensions) + " |"])
        for subject in subjects:
            for judge in judges:
                cells = []
                for dimension in dimensions:
                    values = [row["scores"].get(dimension) for row in rows if row["subject"]["id"] == subject and row["judge"]["id"] == judge and row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(row.get("scores", {}).get(dimension), (int, float))]
                    cells.append(f"{median(values):.4f}" if values else "UNAVAILABLE")
                lines.append("| " + subject + " | " + judge + " | " + " | ".join(cells) + " |")
        lines.extend(["", "## Dimension distribution by subject and judge", "", "| Subject | Judge | Dimension | Median | MAD | IQR | Range | N |", "| --- | --- | --- | ---: | ---: | ---: | ---: | ---: |"])
        for subject in subjects:
            for judge in judges:
                matching = [row for row in rows if row["subject"]["id"] == subject and row["judge"]["id"] == judge]
                for dimension in dimensions:
                    summary = robust_dimension_summary(matching, dimension)
                    lines.append(f"| {subject} | {judge} | {dimension} | {summary['median'] if summary['median'] is not None else 'UNAVAILABLE'} | {summary['mad'] if summary['mad'] is not None else 'UNAVAILABLE'} | {summary['iqr'] if summary['iqr'] is not None else 'UNAVAILABLE'} | {summary['range'] if summary['range'] is not None else 'UNAVAILABLE'} | {summary['n']} |")
    lines.extend(["", "Self-judging cells are displayed separately and excluded by default. Explicit inclusion is exploratory; agreement uses non-self observations. Scores are descriptive and do not establish quality acceptance or Judge validity.", "", "## Operational telemetry", "", "| Subject | Judge | First character | First-character ms | Answer queue ms | Answer ms | Answer tokens | Answer attempts | Answer cached/write | Answer error | Judge queue ms | Judge ms | Judge tokens | Judge attempts | Judge cached/write | Cache hit ratio | Judge error | Status |", "| --- | --- | :---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- | --- |"])
    for row in rows:
        answer, judge = row["answer"], row["judgement"]
        lines.append(f"| {row['subject']['id']} | {row['judge']['id']} | {answer.get('first_char','')} | {answer.get('first_character_ms','')} | {answer.get('queue_ms','')} | {answer.get('elapsed_ms','')} | {answer.get('total_tokens','')} | {answer.get('attempt_count','')} | {answer.get('cached_input_tokens','')}/{answer.get('cache_write_tokens','')} | {answer.get('error_type','')} | {judge.get('queue_ms','')} | {judge.get('elapsed_ms','')} | {judge.get('total_tokens','')} | {judge.get('attempt_count','')} | {judge.get('cached_input_tokens','')}/{judge.get('cache_write_tokens','')} | {judge.get('cache_hit_ratio','')} | {judge.get('error_type','')} | {row.get('status')} |")
    return "\n".join(lines) + "\n"


def write_matrix_workbook(rows: Sequence[Mapping[str, Any]], path: Path) -> None:
    """Write de-duplicated answers and answer-linked judgement facts."""
    if Workbook is None:
        raise RuntimeError("openpyxl is required to write matrix workbooks")
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Matrix"
    subjects = sorted({str(row["subject"]["id"]) for row in rows})
    judges = sorted({str(row["judge"]["id"]) for row in rows})
    def append_text_safe(sheet: Any, values: Sequence[Any]) -> None:
        sheet.append(list(values))
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    append_text_safe(overview, ["XiaoAn subject \\ Judge", *judges])
    for subject in subjects:
        values = []
        for judge in judges:
            matching_pair = [row for row in rows if row["subject"]["id"] == subject and row["judge"]["id"] == judge]
            scores = [row["weighted_score"] for row in matching_pair if row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(row.get("weighted_score"), (int, float))]
            summary = matrix_pair_summary(matching_pair)
            values.append(summary["value"] if summary["value"] is not None else "SELF_ISOLATED" if any(row.get("self_judging") and not row.get("primary_eligible", True) for row in matching_pair) else "UNAVAILABLE")
        append_text_safe(overview, [subject, *values])
    coverage_sheet = workbook.create_sheet("Oracle_Coverage")
    append_text_safe(coverage_sheet, ["metric", "reviewed_units", "status", "detail"])
    for item in coverage_rows([row.get("coverage", {}) for row in rows]):
        append_text_safe(coverage_sheet, [item["metric"], item["value"], item["status"], item["interpretation"]])
    dimensions = sorted({str(name) for row in rows for name in row.get("scores", {})})
    answers = workbook.create_sheet("All_Answers")
    append_text_safe(answers, ["answer_id", "case_id", "turn", "subject_id", "status", "answer", "answer_first_char", "answer_first_character_ms", "answer_queue_ms", "answer_elapsed_ms", "answer_input_tokens", "answer_output_tokens", "answer_total_tokens", "answer_attempt_count", "answer_retry_errors", "answer_cached_input_tokens", "answer_cache_write_tokens", "answer_cache_hit_ratio", "trace_sha256", "error_type", "error"])
    unique_answers: dict[str, Mapping[str, Any]] = {}
    for row in rows:
        unique_answers.setdefault(str(row["answer_id"]), row)
    for answer_id, row in unique_answers.items():
        answer = row["answer"]
        trace = answer.get("trace")
        trace_sha256 = (
            "sha256:" + hashlib.sha256(_canonical_json(trace).encode("utf-8")).hexdigest()
            if isinstance(trace, Mapping)
            else ""
        )
        append_text_safe(answers, [answer_id, row.get("case_id"), row.get("turn"), row["subject"]["id"], answer.get("status"), answer.get("text", ""), answer.get("first_char", ""), answer.get("first_character_ms"), answer.get("queue_ms"), answer.get("elapsed_ms"), answer.get("input_tokens"), answer.get("output_tokens"), answer.get("total_tokens"), answer.get("attempt_count"), json.dumps(answer.get("retry_errors", ()), ensure_ascii=False), answer.get("cached_input_tokens"), answer.get("cache_write_tokens"), answer.get("cache_hit_ratio"), trace_sha256, answer.get("error_type"), answer.get("error")])

    judgements = workbook.create_sheet("All_Judgements")
    append_text_safe(judgements, ["answer_id", "case_id", "turn", "subject_id", "judge_id", "status", "self_judging", "primary_eligible", "triggered_red_lines", "red_line_evidence", "judge_text", "judge_queue_ms", "judge_elapsed_ms", "judge_input_tokens", "judge_output_tokens", "judge_total_tokens", "judge_attempt_count", "judge_retry_errors", "judge_cached_input_tokens", "judge_cache_write_tokens", "judge_cache_hit_ratio", *[f"dimension:{name}" for name in dimensions], "weighted_score", "error_type", "error"])
    for row in rows:
        judged = row["judgement"]
        append_text_safe(judgements, [row["answer_id"], row.get("case_id"), row.get("turn"), row["subject"]["id"], row["judge"]["id"], row.get("status"), row.get("self_judging"), row.get("primary_eligible"), ";".join(row.get("triggered_red_lines", ())), json.dumps(row.get("red_line_evidence", {}), ensure_ascii=False), judged.get("text", ""), judged.get("queue_ms"), judged.get("elapsed_ms"), judged.get("input_tokens"), judged.get("output_tokens"), judged.get("total_tokens"), judged.get("attempt_count"), json.dumps(judged.get("retry_errors", ()), ensure_ascii=False), judged.get("cached_input_tokens"), judged.get("cache_write_tokens"), judged.get("cache_hit_ratio"), *[row.get("scores", {}).get(name, "UNAVAILABLE") for name in dimensions], row.get("weighted_score"), judged.get("error_type") or row["answer"].get("error_type"), judged.get("error") or row["answer"].get("error")])
    dimension_sheet = workbook.create_sheet("Dimension_By_Judge")
    append_text_safe(dimension_sheet, ["subject_id", "judge_id", *dimensions, "weighted_score_average", "available_count", "attempted_count"])
    for subject in subjects:
        for judge in judges:
            matching = [row for row in rows if row["subject"]["id"] == subject and row["judge"]["id"] == judge]
            cells = []
            for dimension in dimensions:
                values = [row["scores"].get(dimension) for row in matching if row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(row.get("scores", {}).get(dimension), (int, float))]
                cells.append(median(values) if values else "UNAVAILABLE")
            weighted = [row["weighted_score"] for row in matching if row.get("status") == "PASS" and row.get("primary_eligible", True) and isinstance(row.get("weighted_score"), (int, float))]
            append_text_safe(dimension_sheet, [subject, judge, *cells, matrix_pair_summary(matching)["value"] if matrix_pair_summary(matching)["value"] is not None else "UNAVAILABLE", len(weighted), len(matching)])
    self_sheet = workbook.create_sheet("Self_Judging_Isolated")
    append_text_safe(self_sheet, ["answer_id", "case_id", "turn", "subject_id", "judge_id", "weighted_score", "status"])
    for row in rows:
        if row.get("self_judging"):
            append_text_safe(self_sheet, [row.get("answer_id"), row.get("case_id"), row.get("turn"), row["subject"]["id"], row["judge"]["id"], row.get("weighted_score"), row.get("status")])
    unique_rows = list({str(row.get("answer_id")): row for row in rows}.values())
    summaries = build_matrix_summaries(
        rows,
        memory_results=[item for row in unique_rows for item in row.get("memory_metrics", ())],
        attribution_results=[row.get("attribution", {}) for row in unique_rows if row.get("attribution", {}).get("status") != "NOT_RUN"],
    )
    contract_sheet = workbook.create_sheet("Measurement_Contract")
    append_text_safe(contract_sheet, ["section", "json"])
    for name in ("primary_denominator", "memory", "attribution", "semantic_oracle"):
        append_text_safe(contract_sheet, [name, json.dumps(summaries[name], ensure_ascii=False, sort_keys=True)])
    agreement_sheet = workbook.create_sheet("Judge_Agreement")
    append_text_safe(agreement_sheet, ["dimension", "interpretation", "eligible_n", "missing_n", "krippendorff_alpha_ordinal", "kendall_w", "pairwise_spearman_json", "alpha_status", "alpha_reason", "kendall_strata_json"])
    for dimension, report in summaries["agreement"].items():
        append_text_safe(agreement_sheet, [dimension, report["interpretation"], report["eligible_n"], report["missing_n"], report["krippendorff_alpha_ordinal"], report["kendall_w"], json.dumps(report["pairwise_spearman"], ensure_ascii=False), report["alpha_status"], report["alpha_reason"], json.dumps(report["kendall_strata"], ensure_ascii=False)])
    red_line_sheet = workbook.create_sheet("Red_Line_Agreement")
    append_text_safe(red_line_sheet, ["red_line_id", "interpretation", "eligible_n", "missing_n", "krippendorff_alpha_nominal", "pairwise_exact_json"])
    for identifier, report in summaries["red_line_agreement"].items():
        append_text_safe(red_line_sheet, [identifier, report["interpretation"], report["eligible_n"], report["missing_n"], report["krippendorff_alpha_nominal"], json.dumps(report["pairwise_exact_agreement"], ensure_ascii=False)])
    stats_sheet = workbook.create_sheet("Dimension_Statistics")
    append_text_safe(stats_sheet, ["subject_id", "judge_id", "dimension", "raw_scores_json", "median", "mad", "iqr", "range", "n"])
    for subject in subjects:
        for judge in judges:
            matching = [row for row in rows if row["subject"]["id"] == subject and row["judge"]["id"] == judge]
            for dimension in dimensions:
                summary = robust_dimension_summary(matching, dimension)
                append_text_safe(stats_sheet, [subject, judge, dimension, json.dumps(summary["raw_scores"], ensure_ascii=False), summary["median"], summary["mad"], summary["iqr"], summary["range"], summary["n"]])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_pair_workbooks(rows: Sequence[Mapping[str, Any]], directory: Path) -> list[Path]:
    """Write one auditable workbook for every subject/judge pair."""
    grouped: dict[tuple[str, str], list[Mapping[str, Any]]] = {}
    for row in rows:
        key = (str(row["subject"]["id"]), str(row["judge"]["id"]))
        grouped.setdefault(key, []).append(row)
    directory.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for (subject, judge), pair_rows in sorted(grouped.items()):
        safe_subject = re.sub(r"[^A-Za-z0-9._-]+", "_", subject).strip("._") or "subject"
        safe_judge = re.sub(r"[^A-Za-z0-9._-]+", "_", judge).strip("._") or "judge"
        path = directory / f"{safe_subject}__vs__{safe_judge}.xlsx"
        write_matrix_workbook(pair_rows, path)
        paths.append(path)
    return paths
