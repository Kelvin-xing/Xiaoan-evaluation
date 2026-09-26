"""Offline scenario diagnostics; never calls a provider or changes legacy scores."""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import replace
import hashlib
import json
import math
from pathlib import Path
from statistics import mean, median
from typing import Mapping

AXES = ('task_family', 'task', 'topic', 'constraints', 'dialogue', 'risk')


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def _map(value):
    return value if isinstance(value, Mapping) else {}


def _rows(value):
    return value if isinstance(value, (list, tuple)) else []


def _index(path, version, name, key):
    if not path:
        return {}
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    if data.get('schema_version') != version or not isinstance(data.get(name), list):
        raise ValueError(f'expected {version} {name} array')
    result = {}
    for row in data[name]:
        identity = tuple(row.get(field) for field in key)
        if None in identity or identity in result:
            raise ValueError(f'missing/duplicate {name} identity')
        result[identity] = row
    return result


def binding(case_id, turn, subject_id, answer, snapshot):
    return {'case_id': case_id, 'turn': turn, 'subject_id': subject_id,
            'answer_sha256': hashlib.sha256(answer.encode()).hexdigest(),
            'context_sha256': digest(snapshot) if snapshot else None}


def _turns(records, cases=(), subject_id=None):
    """Normalize case records or matrix cells without conflating Judge panels."""
    case_map = {getattr(c, 'id', ''): c for c in cases}
    for record in records:
        case_id = record.get('case_id')
        if 'subject' in record and 'answer' in record:
            case = case_map.get(case_id)
            history = [{'turn': t.turn, 'user': t.user} for t in getattr(case, 'turns', ()) if t.turn <= record['turn']]
            answer = _map(record.get('answer'))
            attr = _map(record.get('attribution'))
            attr = {'status': attr.get('status'), **_map(attr.get('result'))}
            checks = {}
            expected = None
            if case is not None:
                from .metrics import evaluate_turn_metrics
                expected = next((getattr(t, 'expected', None) for t in case.turns if t.turn == record['turn']), None)
                if expected is not None:
                    checks = {key: {'status': value.status.value, 'score': value.score, 'reason': value.reason}
                              for key, value in evaluate_turn_metrics(expected, _map(answer.get('trace')),
                                  oracle_approved=getattr(case, 'oracle_gate_eligible', False)).items()}
            yield {'case_id': case_id, 'turn': record['turn'], 'subject_id': record['subject'].get('id'),
                   'judge_id': record.get('judge', {}).get('id', 'UNKNOWN'), 'answer': answer.get('text') or '',
                   'answer_available': answer.get('status') == 'PASS', 'trace': _map(answer.get('trace')),
                   'history': history, 'observation': {'expected': {'response_oracle': __import__('dataclasses').asdict(expected.response_oracle) if expected is not None and getattr(expected, 'response_oracle', None) is not None else {}}}, 'metrics': checks, 'attribution': attr,
                   'primary': {}, 'oracle': _map(record.get('oracle_assessment')),
                   'memory': record.get('memory_metrics', []), 'primary_eligible': record.get('primary_eligible', True),
                   'quality_focus': list(getattr(case, 'quality_focus', ()))}
            continue
        pipeline = _map(record.get('pipeline'))
        observations = {x.get('turn'): x for x in pipeline.get('observations', [])}
        traces = {x.get('turn'): x for x in pipeline.get('turn_traces', [])}
        actual_turns = _map(record.get('conversation')).get('turns', [])
        actual_by_turn = {t.get('turn', i + 1): t for i, t in enumerate(actual_turns)}
        planned = pipeline.get('planned_turns') or actual_turns
        turns = [{**p, **actual_by_turn.get(p.get('turn', i + 1), {})} for i, p in enumerate(planned)]
        history = []
        for i, row in enumerate(turns):
            turn = row.get('turn', i + 1)
            history.append({'turn': turn, 'user': row.get('user_input', '')})
            obs = observations.get(turn, {})
            primary = _map(obs.get('judge'))
            trace = _map(traces.get(turn, {}).get('trace'))
            model = subject_id or record.get('subject_id') or _map(trace.get('model')).get('id') or _map(trace.get('models')).get('response') or 'UNKNOWN'
            answer = row.get('assistant_response') or ''
            yield {'case_id': case_id, 'turn': turn, 'subject_id': model, 'judge_id': 'primary',
                   'answer': answer, 'answer_available': bool(answer), 'trace': trace,
                   'history': list(history), 'observation': obs,
                   'metrics': pipeline.get('turns', [])[i] if i < len(pipeline.get('turns', [])) else {},
                   'attribution': _map(obs.get('attribution')), 'primary': primary,
                   'oracle': _map(primary.get('oracle_assessment')), 'memory': {},
                   'primary_eligible': bool(answer and primary and obs.get('oracle_approved')),
                   'quality_focus': _map(record.get('cohorts')).get('quality_focus', [])}


def _claim_facts(attribution, answer):
    if attribution.get('status') != 'AVAILABLE':
        return []
    result = []
    seen = set()
    for claim in _rows(attribution.get('claims')):
        span = _map(claim.get('answer_span'))
        a, b = span.get('start'), span.get('end')
        if (claim.get('claim_id') in seen or type(a) is not int or type(b) is not int
                or not 0 <= a < b <= len(answer) or answer[a:b] != span.get('text')):
            raise ValueError('attribution claims do not match answer or contain duplicate IDs')
        seen.add(claim.get('claim_id'))
        relations = _rows(claim.get('relations'))
        names = {r.get('relation') for r in relations}
        category = ('CONTRADICTS' if 'CONTRADICTS' in names else 'ENTAILS' if 'ENTAILS' in names
                    else 'PARTIAL_ONLY' if 'PARTIAL' in names else 'NOT_SUPPORTED')
        result.append({'claim_id': claim.get('claim_id'), 'kind': claim.get('kind', 'UNKNOWN'),
                       'text': span.get('text'), 'support_category': category,
                       'unsupported_category': claim.get('unsupported_category'),
                       'layers': sorted({r.get('layer') for r in relations if r.get('layer')}),
                       'evidence_refs': [r.get('evidence_ref') for r in relations if r.get('evidence_ref')],
                       'relations': relations, 'provenance': 'UPSTREAM_VALIDATED_ATTRIBUTION_WITH_LOCAL_ANSWER_SPAN_CHECK', 'demand_relevance': 'UNAVAILABLE_NOT_ANNOTATED'})
    return result


def _ar(entry, bound):
    empty = {'status': 'UNAVAILABLE', 'score': None, 'reason': 'MISSING_BOUND_REVERSE_QUESTIONS_AND_EMBEDDINGS'}
    if not entry:
        return empty
    if not bound['context_sha256'] or bound['subject_id'] == 'UNKNOWN' or entry.get('binding') != bound:
        return {**empty, 'reason': 'BINDING_MISMATCH_OR_CONTEXT_MISSING'}
    ar = _map(entry.get('answer_relevancy'))
    sims = ar.get('similarities')
    n = ar.get('n')
    if ar.get('status') != 'AVAILABLE':
        return {**empty, 'reason': ar.get('reason', 'PROVIDER_OR_SCORING_UNAVAILABLE')}
    if (type(n) is not int or n < 1 or not isinstance(sims, list) or len(sims) != n
            or not all(type(v) in (int, float) and math.isfinite(v) and -1 <= v <= 1 for v in sims)
            or not ar.get('embedding_model') or not ar.get('embedding_revision')
            or not ar.get('generator_model') or not ar.get('prompt_version')
            or ar.get('query_mode') != 'raw_current_user'
            or not isinstance(ar.get('questions'), list) or len(ar['questions']) != n
            or not all(isinstance(q, str) and q.strip() for q in ar['questions'])):
        return {**empty, 'reason': 'INVALID_RELEVANCY_CONTRACT'}
    return {**ar, 'score': mean(sims), 'reason': 'RAW_CURRENT_QUERY_MEAN_COSINE'}


def _oracle(oracle):
    items = [x for x in _rows(oracle.get('items')) if x.get('kind') == 'required'] if oracle.get('status') == 'AVAILABLE' else []
    known = [x for x in items if x.get('verdict') in ('SATISFIED', 'VIOLATED')]
    count = sum(x.get('verdict') == 'SATISFIED' for x in known)
    return {'status': 'AVAILABLE' if known else 'UNAVAILABLE', 'satisfied': count, 'evaluated': len(known),
            'uncertain': len(items) - len(known), 'total': len(items), 'score': count / len(known) if known else None,
            'items': items, 'interpretation': 'Judge-assessed required oracle coverage, not confirmed defect count'}


def build_analysis(records, *, annotations_path=None, analysis_path=None, cases=(), subject_id=None):
    annotations = _index(annotations_path, 'scenario-annotations/v1', 'annotations', ('case_id', 'turn'))
    sidecars = _index(analysis_path, 'answer-analysis/v1', 'results', ('case_id', 'turn', 'subject_id'))
    turns, claims = [], []
    for raw in _turns(records, cases, subject_id):
        ident = {k: raw[k] for k in ('case_id', 'turn', 'subject_id', 'judge_id')}
        snapshot = _map(raw['trace'].get('effective_context_snapshot'))
        bound = binding(raw['case_id'], raw['turn'], raw['subject_id'], raw['answer'], snapshot)
        annotation = annotations.get((raw['case_id'], raw['turn']), {})
        matches = bool(annotation and raw['history'] and annotation.get('user_history_sha256') == digest(raw['history']))
        tag_status = annotation.get('annotation_status', 'PROVISIONAL') if matches else 'UNLABELED'
        tags = {axis: annotation.get(axis, []) if matches else [] for axis in AXES}
        if any(not isinstance(v, list) or any(not isinstance(s, str) or not s for s in v) for v in tags.values()):
            raise ValueError('scenario tags must be arrays of nonempty strings')
        route = _map(raw['trace'].get('route'))
        rid = route.get('id', route.get('route_id', route.get('capsule_id')))
        branch = route.get('branch') or ('baseline' if rid == 'baseline' else 'crisis_sop' if rid == 'crisis_sop' else 'capsule' if rid else 'UNKNOWN')
        primary = {}
        for c in _rows(raw['primary'].get('faithfulness_claims')):
            if isinstance(c.get('supported'), bool) and isinstance(c.get('claim'), str):
                primary[c['claim']] = primary.get(c['claim'], True) and c['supported']
        fact_error = None
        try:
            facts = _claim_facts(raw['attribution'], raw['answer'])
        except ValueError as exc:
            facts, fact_error = [], str(exc)
        counts = Counter(c['support_category'] for c in facts)
        strict = {'status': 'AVAILABLE_PROXY' if facts else 'UNAVAILABLE', 'entailed': counts['ENTAILS'], 'total': len(facts),
                  'score': counts['ENTAILS'] / len(facts) if facts else None, 'support_counts': dict(counts), 'reason': fact_error}
        ar = _ar(sidecars.get((raw['case_id'], raw['turn'], raw['subject_id'])), bound)
        if not raw['answer_available']:
            ar = {'status': 'UNAVAILABLE', 'score': None, 'reason': 'ANSWER_UNAVAILABLE'}
        coverage = _oracle(raw['oracle'])
        checks = {name: value for name, value in raw['metrics'].items()
                  if any(token in name for token in ('route', 'ground', 'source', 'wiki', 'capsule', 'history', 'memory', 'injection', 'safety'))}
        hypotheses = []
        for name, check in checks.items():
            if str(_map(check).get('status')).upper() in ('FAIL', 'ERROR'):
                hypotheses.append({'type': 'HYPOTHESIS', 'stage': 'routing' if 'route' in name else 'context_or_observation',
                                   'evidence': name, 'next_check': 'Compare expected requirement, trace and actual injected content; failure is not causal proof.'})
        if coverage['evaluated'] > coverage['satisfied']:
            hypotheses.append({'type': 'HYPOTHESIS', 'stage': 'composer_or_oracle', 'evidence': 'required oracle violation',
                               'next_check': 'Check requirement applicability and whether needed content was actually injected before changing prompt.'})
        if counts['PARTIAL_ONLY'] or counts['NOT_SUPPORTED'] or counts['CONTRADICTS']:
            hypotheses.append({'type': 'HYPOTHESIS', 'stage': 'generation_knowledge_or_judge', 'evidence': 'claim support categories',
                               'next_check': 'Review joint evidence, claim atomization and policy priority; unsupported does not equal hallucination.'})
        turn = {**ident, 'user': raw['history'][-1]['user'] if raw['history'] else None, 'answer': raw['answer'], 'oracle_items': raw['oracle'].get('items', []), 'oracle_status': raw['observation'].get('expected', {}).get('reference_oracle', {}).get('status', 'UNKNOWN'), 'binding': bound, 'annotation_status': tag_status,
                'annotation_reason': 'MATCHED_USER_PREFIX' if matches else 'MISSING_OR_STALE_USER_PREFIX',
                'tags': tags, 'need_summary': annotation.get('need_summary') if matches else None,
                'quality_focus': raw['quality_focus'], 'actual_branch': branch, 'actual_route': rid,
                'answer_status': 'AVAILABLE' if raw['answer_available'] else 'UNAVAILABLE',
                'primary_eligible': raw['primary_eligible'],
                'primary_binary': {'supported': sum(primary.values()), 'total': len(primary), 'score': sum(primary.values()) / len(primary) if primary else None,
                                   'status': 'AVAILABLE_LEGACY_JUDGE' if primary else 'UNAVAILABLE'},
                'strict_entailment_proxy': strict, 'answer_relevancy': ar, 'required_coverage': coverage,
                'trace_checks': checks, 'observed_expected': raw['observation'].get('expected', {}),
                'observed_actual': raw['observation'].get('actual', {}), 'memory': raw['memory'],
                'context_status': 'AVAILABLE' if snapshot else 'UNAVAILABLE_TRACE_MISSING',
                'hypotheses': hypotheses,
                'claim_demand_relevance': 'UNAVAILABLE_NOT_ANNOTATED',
                'irrelevant_redundant_proportion': 'UNAVAILABLE_NOT_ANNOTATED'}
        turns.append(turn)
        claims.extend({**ident, **fact} for fact in facts)
    groups = _groups(turns)
    unique = {}
    for row in turns:
        key = (row['case_id'], row['turn'], row['subject_id'], row['binding']['answer_sha256'], row['binding']['context_sha256'])
        unique.setdefault(key, row)
    answer_groups = []
    for group in _groups([{**r, 'judge_id': 'ANSWER_LEVEL'} for r in unique.values()]):
        # AR is answer-level. Never average repeated Judge cells as independent answers.
        answer_groups.append({k: v for k, v in group.items() if k not in ('primary_binary', 'strict_entailment_proxy', 'required_coverage', 'primary_eligible_n')})
    claim_groups = []
    for axis in ('kind', 'layers', 'support_category', 'unsupported_category'):
        grouped = defaultdict(list)
        for c in claims:
            labels = c['layers'] or ['NO_EVIDENCE_LAYER'] if axis == 'layers' else [c.get(axis) or 'NONE']
            for label in labels:
                grouped[(c['subject_id'], c['judge_id'], label)].append(c)
        for (subject, judge, label), cs in sorted(grouped.items()):
            claim_groups.append({'axis': axis, 'label': label, 'subject_id': subject, 'judge_id': judge,
                                 'n_claims': len(cs), 'n_turns': len({(c['case_id'], c['turn']) for c in cs}),
                                 'n_cases': len({c['case_id'] for c in cs}),
                                 'support_counts': dict(Counter(
                                     ('CONTRADICTS' if 'CONTRADICTS' in {r.get('relation') for r in c['relations'] if r.get('layer') == label}
                                      else 'ENTAILS' if 'ENTAILS' in {r.get('relation') for r in c['relations'] if r.get('layer') == label}
                                      else 'PARTIAL_ONLY' if 'PARTIAL' in {r.get('relation') for r in c['relations'] if r.get('layer') == label}
                                      else 'NOT_SUPPORTED') if axis == 'layers' else c['support_category'] for c in cs)),
                                 'interpretation': 'Per-layer relation support; layers overlap, not additive.' if axis == 'layers' else 'Claim overall support category.'})
    return {'schema_version': 'scenario-analysis/v1', 'turns': turns, 'claims': claims, 'groups': groups, 'claim_groups': claim_groups, 'answer_groups': answer_groups,
            'interpretation': 'PROVISIONAL descriptive diagnosis; multi-label groups overlap. No causal or safety conclusion from cosine/support alone. Legacy 0.5 PARTIAL-weighted attribution remains unchanged.'}


def _groups(turns):
    buckets = defaultdict(list)
    for row in turns:
        for axis in ('all', 'actual_branch', *AXES):
            labels = ['ALL'] if axis == 'all' else [row['actual_branch']] if axis == 'actual_branch' else row['tags'][axis] or ['UNLABELED' if row['annotation_status'] == 'UNLABELED' else 'NONE_OR_NOT_TAGGED']
            for label in labels:
                buckets[(row['subject_id'], row['judge_id'], axis, label)].append(row)
    result = []
    for (subject, judge, axis, label), rows in sorted(buckets.items()):
        item = {'subject_id': subject, 'judge_id': judge, 'axis': axis, 'label': label,
                'n_cases': len({r['case_id'] for r in rows}), 'n_turns': len(rows),
                'answered_n': sum(r['answer_status'] == 'AVAILABLE' for r in rows),
                'annotated_n': sum(r['annotation_status'] != 'UNLABELED' for r in rows),
                'reviewed_n': sum(r['annotation_status'] == 'REVIEWED' for r in rows),
                'primary_eligible_n': sum(r['primary_eligible'] for r in rows)}
        for name in ('primary_binary', 'strict_entailment_proxy', 'answer_relevancy', 'required_coverage'):
            values = [r[name]['score'] for r in rows if r[name].get('score') is not None]
            item[name] = {'available_n': len(values), 'unavailable_n': len(rows) - len(values),
                          'turn_macro_mean': mean(values) if values else None, 'median': median(values) if values else None,
                          'min': min(values) if values else None, 'max': max(values) if values else None}
            if name in ('primary_binary', 'strict_entailment_proxy', 'required_coverage'):
                numkey, denkey = {'primary_binary': ('supported', 'total'), 'strict_entailment_proxy': ('entailed', 'total'), 'required_coverage': ('satisfied', 'evaluated')}[name]
                numerator = sum(r[name][numkey] for r in rows)
                denominator = sum(r[name][denkey] for r in rows)
                item[name].update(numerator=numerator, denominator=denominator, micro=numerator / denominator if denominator else None)
        result.append(item)
    return result


def render_analysis(analysis):
    def show(v):
        return f'{v:.4f}' if isinstance(v, (float, int)) else 'UNAVAILABLE'
    def esc(v):
        return str(v).replace('|', '\\|').replace('\n', ' ')
    lines = ['\n## 情境與回答診斷（描述性，非根因裁決）\n', analysis['interpretation'],
             'Faithfulness 分開呈現 legacy binary 與 strict ENTAILS proxy；PARTIAL 不計完整支持。AR 是 N 個反向問題平均餘弦，非正確率／完整性。未提供匹配的本地 AR 結果時為 UNAVAILABLE；本分析不呼叫模型。',
             '缺少 AR 時需先對本次凍結回答生成反向問題及 embeddings，再用同一 case/turn/subject/answer/context 綁定匯入。情境標籤只依使用者前綴；PROVISIONAL 未人工核准。',
             '| Subject / Judge | 分組 | 案例 / 輪 | 已答 / 標註 / 已審 | 主Judge micro | Strict micro | AR均值（有效n） | 必要需求micro |',
             '| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |']
    for r in analysis['groups']:
        ar = r['answer_relevancy']
        lines.append(f"| {esc(r['subject_id'])} / {esc(r['judge_id'])} | {esc(r['axis'])}:{esc(r['label'])} | {r['n_cases']} / {r['n_turns']} | {r['answered_n']} / {r['annotated_n']} / {r['reviewed_n']} | {show(r['primary_binary']['micro'])} | {show(r['strict_entailment_proxy']['micro'])} | {show(ar['turn_macro_mean'])} ({ar['available_n']}) | {show(r['required_coverage']['micro'])} |")
    lines += ['', '### Answer Relevancy（每個 subject 回答只計一次）', '', '| Subject | 分組 | 案例 / 輪 | AR 有效 / 缺失 | 均值 / 中位數 |', '| --- | --- | ---: | ---: | ---: |']
    for r in analysis['answer_groups']:
        ar = r['answer_relevancy']
        lines.append(f"| {esc(r['subject_id'])} | {esc(r['axis'])}:{esc(r['label'])} | {r['n_cases']} / {r['n_turns']} | {ar['available_n']} / {ar['unavailable_n']} | {show(ar['turn_macro_mean'])} / {show(ar['median'])} |")
    lines += ['', '### 逐輪診斷索引', '', '| Subject / Judge | 輪次 | 標註 | Route | AR狀態 | 證據／假說 |', '| --- | --- | --- | --- | --- | --- |']
    for r in analysis['turns']:
        lines.append(f"| {esc(r['subject_id'])} / {esc(r['judge_id'])} | {r['case_id']} T{r['turn']} | {r['annotation_status']} | {esc(r['actual_route'])} | {esc(r['answer_relevancy']['reason'])} | {esc('; '.join(h['stage'] + ': ' + h['evidence'] for h in r['hypotheses']) or 'NO_DETERMINATION')} |")
    lines += ['', '聲明 kind／evidence layer／support 細項、每輪 route/injection/history 檢查與必要需求證據見工作簿 Claims／Evidence／Requirements／Groups；情境摘要見既有輪次表；多標籤不可相加。Claim 與需求的語義相關性、冗餘比例尚未標註，不從 AR 推算。', '']
    return '\n'.join(lines)


def attach_to_model(model, analysis):
    """Use existing Metrics columns; preserve workbook round-trip contract."""
    from .workbook import HEADERS
    rows = [r for r in model.metrics if not str(r.get('metric_id', '')).startswith(('scenario.', 'scenario_chunk.'))]
    for section in ('groups', 'answer_groups', 'turns', 'claims', 'claim_groups'):
        for i, value in enumerate(analysis[section]):
            encoded = json.dumps(value, ensure_ascii=False, sort_keys=True)
            parts = [encoded[j:j+14000] for j in range(0, len(encoded), 14000)]
            for part, text in enumerate(parts):
                rows.append({**dict.fromkeys(HEADERS['03_Metrics']), 'row_key': f'scenario.{section}.{i}.{part}', 'comparison_key': f'scenario.{section}.{i}',
                         'component_run_id': value.get('subject_id'), 'case_id': value.get('case_id'), 'turn': value.get('turn'),
                         'metric_id': f'scenario.{section}' if len(parts)==1 else f'scenario_chunk.{section}', 'raw_score': part if len(parts)>1 else None,
                         'dimension': value.get('axis', value.get('kind', 'diagnostic')),
                         'status': 'PROVISIONAL', 'reason': text,
                         'score_source': 'offline_scenario_analysis/v1'})
    return replace(model, metrics=tuple(rows))


def append_matrix_sheets(path, analysis):
    from openpyxl import load_workbook
    from .detail_tables import project, write_table, HEADERS, ID, TURN_COLUMNS
    book = load_workbook(path)
    turns, tables = project(analysis)
    # Matrix has an Answers sheet; enrich it rather than duplicating answers.
    name = 'All_Answers' if 'All_Answers' in book else '02_Turns'
    if name in book:
        sheet=book[name]
        headers=[c.value for c in sheet[1]]
        existing=[dict(zip(headers,r)) for r in sheet.iter_rows(min_row=2,values_only=True)]
        for row in existing:
            match=next((t for t in turns if t['case_id']==row.get('case_id') and t['turn']==row.get('turn') and t['subject_id']==row.get('subject_id',row.get('subject'))),None)
            if match:row.update({k:match.get(k) for k in TURN_COLUMNS if k not in ('judge_id','primary_faithfulness','strict_faithfulness','required_coverage','diagnostic_next_check')})
        added=[k for k in TURN_COLUMNS if k not in headers and k not in ('judge_id','primary_faithfulness','strict_faithfulness','required_coverage','diagnostic_next_check')]
        write_table(book,name,headers+added,existing)
    else:
        write_table(book,name,tuple(dict.fromkeys((*ID,*TURN_COLUMNS))),turns)
    for name,rows in tables.items():write_table(book,name,HEADERS[name],rows)
    book.save(path)
