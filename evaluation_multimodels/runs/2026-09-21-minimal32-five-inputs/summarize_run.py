"""Offline summary of this authorized five-case trial; no network access."""
from pathlib import Path
from collections import Counter, defaultdict
import json
from statistics import mean
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent
OUTPUT=ROOT/'2026-09-21-minimal32-five-self-included'
CHECKPOINT=ROOT/'.matrix-audit'/OUTPUT.name/'matrix-checkpoint.jsonl'
answers={};judgements={};cells={};history=Counter()
for line in CHECKPOINT.open():
 r=json.loads(line);history[r.get('event')]+=1
 if r.get('event')=='answer':answers[r['answer_id']]=r
 if r.get('event')=='judgement':judgements[r['answer_id'],r['judge']['model']]=r
 if r.get('event')=='cell':cells[r['row']['answer_id'],r['row']['judge']['model']]=r['row']
subject_summary=[]
for provider in ('claude','gpt','gemini','qwen','kimi'):
 ar=[r['answer'] for r in answers.values() if r['subject']['provider']==provider]
 times=[a['elapsed_ms'] for a in ar if isinstance(a.get('elapsed_ms'),(int,float)) and a['status']=='PASS']
 subject_summary.append({'provider':provider,'planned_turns':17,'answers':len(ar),'successful':sum(a['status']=='PASS' for a in ar),'mean_response_ms':mean(times) if times else None,'input_usage_available_n':sum(isinstance(a.get('input_tokens'),int) for a in ar)})
judge_summary=[]
for provider in ('claude','gpt','gemini','qwen','kimi'):
 cr=[r for r in cells.values() if r['judge']['provider']==provider]
 jr=[r['judgement'] for r in judgements.values() if r['judge']['provider']==provider]
 judge_summary.append({'provider':provider,'planned_cells':85,'valid_cells':sum(r['status']=='PASS' for r in cr),'unavailable_cells':sum(r['status']!='PASS' for r in cr),'latest_judgement_records_including_local_skips':len(jr),'reported_input_tokens':sum(r['input_tokens'] for r in jr if isinstance(r.get('input_tokens'),int)),'reported_output_tokens':sum(r['output_tokens'] for r in jr if isinstance(r.get('output_tokens'),int)),'usage_available_n':sum(isinstance(r.get('input_tokens'),int) and isinstance(r.get('output_tokens'),int) for r in jr),'error_types':dict(Counter(r.get('error_type') or 'UNKNOWN' for r in jr if r.get('status')!='PASS'))})
self_rows=[r for r in cells.values() if r.get('self_judging')]
summary={'cases':['TC-01','TC-05','TC-17','TC-29','TC-62'],'planned_turns':17,'subject_summary':subject_summary,'judge_summary':judge_summary,'self_judging':{'cells':len(self_rows),'eligible':sum(r.get('primary_eligible') is True for r in self_rows),'valid':sum(r.get('status')=='PASS' for r in self_rows)},'checkpoint_events':dict(history),'billing':'UNAVAILABLE: subject trace lacks token usage; interruption/retry charges and actual relay prices not complete. Judge reported tokens are partial usage, not total billed cost.'}
if (OUTPUT/'results.xlsx').exists():
 book=load_workbook(OUTPUT/'results.xlsx',read_only=True,data_only=True)
 summary['matrix']=list(book['Matrix'].values);summary['sheets']=book.sheetnames;book.close()
p=Path(__file__).parent/'run-summary.json';p.write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))
