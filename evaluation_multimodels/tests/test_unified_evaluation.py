"""Behavioral contracts for shared inventories, missingness and release gates."""
from copy import deepcopy
import json

import pytest

from xiaoan_eval_core import VERSION
from xiaoan_eval_core.contracts import (digest, extraction_request, validate_inventory,
                                       assessment_request, validate_assessment)
from xiaoan_eval_core.runtime import evaluate
from xiaoan_eval_core.ablation import run as ablate
from xiaoan_eval_core.scoring import inventory_audit


def ident(name):
    return dict(id=name, provider="synthetic", model=name, family=name, prompt_version="test/v1")


def spec():
    return {"contract": VERSION, "schema_version": "evaluation-methods/v3", "extractor": ident("extractor"),
            "judges": [ident("judge-a"), ident("judge-b")], "planned_turns": {"case-1": [1]}, "planned_subjects": ["subject"],
            "rows": [{"answer_id": "a1", "case_id": "case-1", "turn": 1, "subject_id": "subject",
                      "subject_provider": "synthetic", "subject_model": "subject", "status": "AVAILABLE",
                      "question": "Can you help?", "history": [], "answer": "A. B. C.",
                      "context_capture": "EXPOSED", "context_version": "v1",
                      "context": [{"ref": "a", "layer": "SOURCE", "content": "A."},
                                  {"ref": "b", "layer": "CAPSULE", "content": "B."}],
                      "reference_facts": [], "oracle_status": "approved",
                      "requirements": [{"id": "safety", "kind": "safety", "text": "No harmful advice", "critical": True},
                                       {"id": "task", "kind": "task", "text": "Answer the request", "critical": False}]}]}


def provider(request):
    if request["task"] == "generate_ablation":
        return {"binding": request["binding"], "control_hash": request["control_hash"],
                "actual_context": request["context"], "answer": "A. B. C."}
    if request["task"] == "extract_claims":
        return {"binding": request["binding"], "claims": [
            {"id": str(i), "kind": kind, "proposition": f"Proposition {i}", "conditions": [],
             "answer_span": {"start": i * 3, "end": i * 3 + 2, "text": request["answer"][i*3:i*3+2]}}
            for i, kind in enumerate(("FACTUAL", "RECOMMENDATION", "SUPPORTIVE"))]}
    def label(verdict, ref=None):
        return {"verdict": verdict, "reason": "synthetic evidence", "evidence": [
            {"ref": ref, "start": 0, "end": 2, "text": ref.upper() + "."}] if ref else []}
    has_capsule = any(e["ref"] == "b" for e in request["context"])
    return {"binding": request["binding"], "claims": [
        {"id": "0", "faithfulness": label("ENTAILED", "a"), "correctness": label("UNKNOWN")},
        {"id": "1", "faithfulness": label("ENTAILED", "b") if has_capsule else label("UNSUPPORTED"),
         "correctness": label("NOT_APPLICABLE")},
        {"id": "2", "faithfulness": label("NOT_APPLICABLE"), "correctness": label("NOT_APPLICABLE")}],
        "requirements": [{"id": r["id"], "verdict": "SATISFIED", "reason": "synthetic fixture",
                          "answer_spans": [{"start": 0, "end": 2, "text": "A."}]} for r in request["requirements"]]}


def test_extract_once_judges_share_inventory_resume_successful_calls(tmp_path):
    calls = []
    def record(request):
        calls.append(request)
        return provider(request)
    result = evaluate(spec(), record, checkpoint_dir=tmp_path)
    assert [r["task"] for r in calls] == ["extract_claims", "assess_claims", "assess_claims"]
    assert calls[1]["inventory"] == calls[2]["inventory"]
    assert result["cells"][0]["metrics"]["faithfulness"]["eligible_n"] == 2
    assert result["cells"][0]["metrics"]["correctness"]["known_coverage"] == 0
    assert result["cells"][0]["metrics"]["correctness"]["known_support_rate"] is None
    assert result["summary"]["conversations"][0]["release_gate"] == "PASS"
    assert len(result["inventories"]) == 1
    resumed = evaluate(spec(), record, checkpoint_dir=tmp_path)
    assert resumed["cells"] == result["cells"]
    assert all(r["incremental_usage"]["input_tokens"] == 0 for r in resumed["receipts"])
    assert len(calls) == 3


def test_extraction_is_a_batch_barrier_before_any_assessment(tmp_path):
    s = spec()
    second = deepcopy(s["rows"][0])
    second.update(answer_id="a2", turn=2, answer="A. B. D.")
    s["planned_turns"]["case-1"] = [1, 2]
    s["rows"].append(second)
    calls = []

    def record(request):
        calls.append((request["task"], request.get("answer")))
        return provider(request)

    result = evaluate(s, record, checkpoint_dir=tmp_path)
    tasks = [task for task, _ in calls]
    first_assessment = tasks.index("assess_claims")
    assert tasks[:first_assessment] == ["extract_claims", "extract_claims"]
    assert result["stages"]["extraction"]["available_n"] == 2
    assert result["stages"]["assessment"]["planned_n"] == 4


def test_judge_timeout_retained_and_retry_does_not_repeat_extraction(tmp_path):
    calls = []
    def flaky(r):
        calls.append(r["task"])
        if r["task"] == "assess_claims" and r["identity"]["id"] == "judge-b":
            raise TimeoutError("secret-key-must-not-appear")
        return provider(r)
    result = evaluate(spec(), flaky, checkpoint_dir=tmp_path)
    assert result["summary"]["unavailable_cells"] == 1
    assert "secret-key" not in json.dumps(result)
    calls.clear()
    def recovered(r):
        calls.append(r["task"])
        return provider(r)
    result = evaluate(spec(), recovered, checkpoint_dir=tmp_path)
    assert calls == ["assess_claims"]
    assert result["summary"]["available_cells"] == 2


def test_unknown_and_partial_share_same_substantive_denominator():
    def changed(r):
        p = provider(r)
        if r["task"] == "assess_claims":
            p["claims"][0]["faithfulness"]["verdict"] = "PARTIAL"
            p["claims"][1]["faithfulness"] = {"verdict": "UNKNOWN", "evidence": [], "reason": "uncertain"}
        return p
    metrics = evaluate(spec(), changed)["cells"][0]["metrics"]["faithfulness"]
    assert metrics["strict_rate"] == 0
    assert "partial_weighted_diagnostic" not in metrics
    assert metrics["known_coverage"] == .5
    assert metrics["not_applicable_n"] == 1


@pytest.mark.parametrize("mutation", ["omitted", "duplicate", "span", "binding", "wrong_domain", "truth", "kind"])
def test_invalid_semantics_never_become_quality_zero(mutation):
    def invalid(r):
        p = provider(r)
        if r["task"] == "assess_claims":
            if mutation == "omitted": p["claims"].pop()
            if mutation == "duplicate": p["claims"].append(p["claims"][0])
            if mutation == "span": p["claims"][0]["faithfulness"]["evidence"][0]["text"] = "X."
            if mutation == "binding": p["binding"] = "stale"
            if mutation == "wrong_domain": p["claims"][0]["faithfulness"]["evidence"][0]["ref"] = "truth-ref"
            if mutation == "truth": p["claims"][0]["correctness"] = deepcopy(p["claims"][0]["faithfulness"])
            if mutation == "kind": p["claims"][2]["faithfulness"] = deepcopy(p["claims"][0]["faithfulness"])
        return p
    result = evaluate(spec(), invalid)
    expected = "UNAVAILABLE" if mutation in {"omitted", "duplicate", "binding"} else "PARTIAL"
    assert all(c["status"] == expected for c in result["cells"])


def test_joint_evidence_requires_explicit_semantic_verdict():
    s = spec()
    def joint(r):
        p = provider(r)
        if r["task"] == "assess_claims":
            p["claims"][0]["faithfulness"]["evidence"].append({"ref": "b", "start": 0, "end": 2, "text": "B."})
            p["claims"][0]["faithfulness"]["reason"] = "Both sources jointly support the qualified claim"
        return p
    cell = evaluate(s, joint)["cells"][0]
    assert cell["metrics"]["faithfulness"]["strict_rate"] == 1
    assert cell["metrics"]["layer_attribution"]["claims"][0]["layers"] == ["CAPSULE", "SOURCE"]


def test_persistent_constraint_and_known_safety_failure_survive_missing_turn():
    s = spec()
    s["planned_turns"]["case-1"] = [1, 2, 3]
    second = deepcopy(s["rows"][0]); second.update(turn=2, answer_id="a2")
    third = deepcopy(second); third.update(turn=3, answer_id="a3", status="UNAVAILABLE", answer=None)
    s["rows"] += [second, third]
    s["conversation_constraints"] = [{"case_id": "case-1", "id": "no-call", "text": "Do not recommend calls",
                                       "start_turn": 2, "end_turn": 2, "status": "approved", "critical": True}]
    def violated(r):
        p = provider(r)
        if r["task"] == "assess_claims":
            for req in p["requirements"]:
                if req["id"] == "no-call": req.update(verdict="VIOLATED", answer_spans=[])
        return p
    result = evaluate(s, violated)
    assert all(c["release_gate"] == "FAIL" for c in result["summary"]["conversations"])
    assert result["summary"]["unavailable_cells"] == 2
    assert "no-call" not in [r["id"] for r in result["cells"][0]["assessment"]["requirements"]]


def test_gate_is_independent_of_aggregate_approval_and_absent_critical_is_na():
    s = spec(); s["rows"][0]["oracle_status"] = "provisional"
    assert evaluate(s, provider)["cells"][0]["requirements"]["release_gate"] == "PASS"
    s = spec(); s["rows"][0]["requirements"].pop(0)
    assert evaluate(s, provider)["cells"][0]["requirements"]["release_gate"] == "NOT_APPLICABLE"


def test_same_model_judges_are_included_without_flags():
    s = spec(); s["judges"][0]["model"] = "subject"; s["judges"][0]["provider"] = "other-alias"
    result = evaluate(s, provider)
    assert result["summary"]["available_cells"] == 2
    assert all('primary_eligible' not in c and 'same_model_family' not in c for c in result['cells'])


def test_complete_planned_denominator_is_required():
    s = spec(); s["planned_turns"]["case-1"].append(2)
    with pytest.raises(ValueError, match="every planned turn"):
        evaluate(s, provider)


def test_empty_inventory_is_not_perfect_faithfulness():
    def empty(r):
        p = provider(r); p["claims"] = []; return p
    metrics = evaluate(spec(), empty)["cells"][0]["metrics"]
    assert metrics["faithfulness"]["strict_rate"] is None


def test_extraction_audit_independent_binding_and_missing_gold():
    s = spec(); request = extraction_request(s["rows"][0], s["extractor"])
    inventory = validate_inventory(provider(request), request)
    gold = [{"id": "g1"}, {"id": "g2"}]
    audit = {"status": "approved", "reviewer_id": "human", "inventory_id": inventory["inventory_id"],
             "human_inventory_hash": digest(gold), "matches": [{"human_id": "g1", "claim_id": "0"}]}
    assert inventory_audit(inventory, audit, gold)["recall"] == .5
    audit["matches"][0]["claim_id"] = "invented"
    with pytest.raises(ValueError): inventory_audit(inventory, audit, gold)


def ablation_spec():
    s = spec()
    s["controls"] = {"mode": "COMPOSER_FIXED_CONTEXT", "model": "subject", "provider": "synthetic",
                     "prompt_version": "p1", "knowledge_version": "k1", "route_id": "fixed",
                     "system_prompt": "Synthetic subject", "seed": 1, "temperature": 0, "max_output_tokens": 100}
    return s


def test_ablation_only_removes_capsules_and_retains_safety_regressions(tmp_path):
    requests = []
    def subject(r):
        requests.append(r)
        return provider(r)
    def unsafe(r):
        p = provider(r)
        if r["task"] == "assess_claims" and not any(e["layer"] == "CAPSULE" for e in r["context"]):
            p["requirements"][0]["verdict"] = "VIOLATED"
        return p
    result = ablate(ablation_spec(), unsafe, checkpoint_dir=tmp_path, subject_provider=subject)
    assert len(requests) == 2
    assert requests[0]["controls"] == requests[1]["controls"]
    assert requests[1]["context"] == [u for u in requests[0]["context"] if u["layer"] != "CAPSULE"]
    assert all(p["status"] == "FAIL" for p in result["pairs"])
    assert all(p["diagnostic_deltas"]["faithfulness"] == .5 for p in result["pairs"])


def test_ablation_rejects_drift_and_never_manufactures_missing_effect():
    def drift(r):
        p = provider(r); p["control_hash"] = "changed"; return p
    result = ablate(ablation_spec(), provider, subject_provider=drift)
    assert all(r["status"] == "UNAVAILABLE" for r in result["receipts"])
    assert all(p["status"] == "UNAVAILABLE" and "diagnostic_deltas" not in p for p in result["pairs"])


def test_ablation_known_failure_not_hidden_by_other_arm_timeout():
    def subject(r):
        if not any(u["layer"] == "CAPSULE" for u in r["context"]):
            raise TimeoutError()
        return provider(r)
    def unsafe(r):
        p = provider(r)
        if r["task"] == "assess_claims": p["requirements"][0]["verdict"] = "VIOLATED"
        return p
    result = ablate(ablation_spec(), unsafe, subject_provider=subject)
    assert all(p["status"] == "FAIL" and p["with_gate"] == "FAIL"
               and p["without_gate"] == "UNAVAILABLE" and "diagnostic_deltas" not in p for p in result["pairs"])


def test_old_cli_names_are_removed(tmp_path):
    from xiaoan_eval.cli import main
    for name in ('unified', 'staged'):
        with pytest.raises(SystemExit) as error:
            main(['measure', name, str(tmp_path/'input.json'), '--output', str(tmp_path/'out')])
        assert error.value.code == 2


def test_untrusted_metadata_not_sent_to_provider():
    s = spec()
    s["rows"][0]["context"][0]["private_trace"] = "do-not-send"
    s["rows"][0]["trace"] = "do-not-send"
    s["judges"][0]["api_key"] = "do-not-send"
    s["rows"][0]["history"] = [{"role": "user", "content": "Previous", "private_trace": "do-not-send"}]
    requests = []
    def record(r): requests.append(r); return provider(r)
    evaluate(s, record)
    assert "do-not-send" not in json.dumps(requests)


def test_atomic_checkpoint_failure_leaves_no_poisoned_final_file(tmp_path, monkeypatch):
    from xiaoan_eval_core.runtime import ResponseStore
    import xiaoan_eval_core.runtime as runtime
    s = spec(); request = extraction_request(s["rows"][0], s["extractor"])
    link = runtime.os.link
    def interrupted(*args): raise OSError("interrupted publication")
    monkeypatch.setattr(runtime.os, "link", interrupted)
    with pytest.raises(OSError): ResponseStore(tmp_path).call(request, provider, validate_inventory)
    assert not list(tmp_path.glob("*.json"))
    assert not list(tmp_path.glob("*.tmp"))
    monkeypatch.setattr(runtime.os, "link", link)
    inventory = ResponseStore(tmp_path).call(request, provider, validate_inventory)
    assert len(inventory["claims"]) == 3


def test_report_contains_bound_source_requests_for_independent_revalidation():
    result = evaluate(spec(), provider)
    for cell in result["cells"]:
        request = result["requests"][cell["assessment_request_id"]]
        assert digest(request) == cell["assessment_request_id"]
        assert validate_assessment(cell["assessment"], request)["claims"] == cell["assessment"]["claims"]
        assert request["answer"] == "A. B. C."
        assert cell["requirements"]["items"][0]["text"] == "No harmful advice"


def test_ablation_strips_nested_metadata_and_expands_constraints_only_once():
    s = ablation_spec()
    s["rows"][0]["context"][0]["private_trace"] = "do-not-send"
    s["rows"][0]["history"] = [{"role": "user", "content": "Previous", "provider_config": "do-not-send"}]
    s["conversation_constraints"] = [{"case_id": "case-1", "id": "no-call", "text": "No calls",
                                       "start_turn": 1, "status": "approved", "critical": False}]
    requests = []
    def record(r): requests.append(r); return provider(r)
    result = ablate(s, record, subject_provider=record)
    assert "do-not-send" not in json.dumps(requests)
    assert all(c["status"] == "AVAILABLE" for c in result["evaluation"]["cells"])


def test_provider_rejection_prevents_successful_assessments():
    calls = []
    def guarded(request):
        raise PermissionError('egress rejected')
    result = evaluate(spec(), guarded)
    assert all(cell['status'] == 'UNAVAILABLE' for cell in result['cells'])
    assert all(len(receipt['attempts']) == 1 for receipt in result['receipts'])


def test_saved_matrix_answer_reused_across_judges_but_conflicts_rejected():
    from xiaoan_eval.unified import attach_records
    record = {"case_id": "case-1", "turn": 1, "subject": {"id": "subject"},
              "answer": {"status": "PASS", "text": "Original answer", "trace": {}}}
    prepared = attach_records(spec(), [record, deepcopy(record)])
    assert prepared["rows"][0]["answer"] == "Original answer"
    assert "context_capture" not in prepared["rows"][0]
    other = deepcopy(record); other["answer"]["text"] = "Changed answer"
    with pytest.raises(ValueError, match="conflicting immutable"):
        attach_records(spec(), [record, other])


def test_saved_run_missing_turn_preserved_and_question_binding_checked():
    from xiaoan_eval.unified import attach_records
    prepared = attach_records(spec(), [])
    assert prepared["rows"][0]["status"] == "UNAVAILABLE"
    record = {"case_id": "case-1", "conversation": {"turns": [
        {"turn": 1, "user_input": "Different question", "assistant_response": "Original"}]}}
    with pytest.raises(ValueError, match="question"):
        attach_records(spec(), [record])


def test_workbook_escapes_formula_text_and_keeps_full_json_provenance(tmp_path):
    from xiaoan_eval_core.reporting import workbook
    from openpyxl import load_workbook
    result = evaluate(spec(), provider)
    result["cells"][0]["requirements"]["items"][0]["reason"] = "=1+1"
    path = tmp_path / "report.xlsx"
    workbook(result, path)
    book = load_workbook(path, data_only=False)
    sheet = book["Requirements"]
    headers = [c.value for c in sheet[1]]
    cell = sheet.cell(2, headers.index("reason") + 1)
    assert cell.data_type == "s" and cell.value == "'=1+1"
    book.close()


def test_builtin_egress_validator_accepts_canonical_and_rejects_extra_fields():
    from xiaoan_eval.unified import validate_evaluation_request
    result = evaluate(spec(), provider)
    for request in result["requests"].values():
        assert validate_evaluation_request(request)
        extra = deepcopy(request); extra["private_trace"] = "not allowed"
        assert not validate_evaluation_request(extra)
    assert not validate_evaluation_request({"task": "generate_ablation"})
