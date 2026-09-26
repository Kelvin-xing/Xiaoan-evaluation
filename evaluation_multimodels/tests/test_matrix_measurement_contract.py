import pytest
from openpyxl import load_workbook

from xiaoan_eval.multimodel import weighted_score, render_matrix_report, write_matrix_workbook
from xiaoan_eval.rules import load_rating_rule


def rows():
    return [dict(answer_id=f"a{i}", case_id="TC-01", turn=i+1,
        subject={"id": "s"}, judge={"id": "j"}, scores={"synthetic": value},
        weighted_score=value, status="PASS", primary_eligible=True, self_judging=False,
        answer={"text": "synthetic"}, judgement={}, attribution={"status": "NOT_RUN"})
        for i, value in enumerate([0, 0, 3])]


def test_matrix_focus_matches_rating_contract():
    rule = load_rating_rule("ratings rule.yml")
    scores = {m.name: 3 if m.name == "行动赋权" else 0 for m in rule.modules}
    assert weighted_score(scores, rule, ["行动赋权"]) == pytest.approx(0.743119266055046)


def test_matrix_renderers_share_weighted_total_and_dimension_summary(tmp_path):
    data = rows()
    report = render_matrix_report(data)
    path = tmp_path / "matrix.xlsx"
    write_matrix_workbook(data, path)
    wb = load_workbook(path, data_only=True)
    assert wb.sheetnames[:5] == ["Overview", "Score Summary", "Routing Summary", "Coverage & Usage", "Matrix"]
    assert [cell.value for cell in wb["Overview"][1]] == ["metric", "value", "status", "note"]
    assert [cell.value for cell in wb["Score Summary"][1]][:4] == ["subject_id", "judge_id", "rubric_average", "faithfulness_average"]
    assert [cell.value for cell in wb["Coverage & Usage"][1]] == ["section", "metric", "value", "status", "note"]
    assert [cell.value for cell in wb["Measurement_Contract"][1]] == ["section", "field", "value"]
    assert "raw_scores_json" not in [cell.value for cell in wb["Dimension_Statistics"][1]]
    assert wb["Matrix"]["B2"].value == 1
    assert "| s | 1.0000 |" in report
    assert wb["Dimension_By_Judge"]["C2"].value == 0
    wb.close()


def test_unavailable_scores_do_not_leak_into_dimension_reports(tmp_path):
    data = rows()
    data[2]['status'] = 'UNAVAILABLE'
    data[2]['scores']['synthetic'] = 99
    report = render_matrix_report(data)
    path = tmp_path / 'partial.xlsx'
    write_matrix_workbook(data, path)
    wb = load_workbook(path, data_only=True)
    assert '| s | UNAVAILABLE |' in report
    assert '| s | 0.0000 |' in report
    assert wb['Dimension_By_Judge']['C2'].value == 0
    wb.close()


def test_matrix_macro_averages_complete_cases_and_excludes_partial_episodes():
    from xiaoan_eval.multimodel import matrix_pair_summary
    data = rows()
    for row in data:
        row['expected_turns'] = [1, 2, 3]
    data.append({**data[0], 'case_id': 'TC-02', 'expected_turns': [1], 'weighted_score': 3})
    assert matrix_pair_summary(data)['value'] == 2
    data[1]['status'] = 'UNAVAILABLE'
    result = matrix_pair_summary(data)
    assert result['value'] == 3
    assert result['eligible_cases'] == 1
