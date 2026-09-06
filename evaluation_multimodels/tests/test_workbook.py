from dataclasses import replace
from pathlib import Path

import pytest

from xiaoan_eval.report_model import build_report_model
from xiaoan_eval.workbook import SHEETS, read_workbook, write_workbook


def _record(response="answer"):
    return {
        "case_id": "TC-01",
        "status": "FAIL",
        "safety": {"hard_gate_passed": True, "red_lines": []},
        "pipeline": {
            "turns": [{"route_accuracy": {"status": "fail", "score": 0.0, "reason": "wrong route"}}],
            "turn_traces": [{
                "turn": 1,
                "response_sha256": "response-hash",
                "trace": {
                    "route": {"id": "baseline"},
                    "safety": {"level": "baseline"},
                    "ground": {"resolved_refs": ["source:1"]},
                    "timings": {"total_ms": 12, "ttft_ms": 3},
                },
            }],
        },
        "conversation": {"conversation_id": "c1", "turns": [{"turn": 1, "user_input": "=1+1", "assistant_response": response}]},
        "quality": {"weighted_total": 1.5},
        "performance": {"total_ms": 12},
        "review": {"status": "not_requested", "per_turn": ["not_requested"]},
        "failure": {"primary_stage": "router"},
        "cohorts": {"risk": ["critical"]},
        "evidence_refs": ["TC-01:T1:route"],
    }


def test_workbook_has_fixed_human_readable_sheets_and_round_trips(tmp_path: Path):
    model = build_report_model([_record()], manifest={"rating_rule_hash": "rule-1"})
    path = tmp_path / "results.xlsx"

    write_workbook(model, path)
    facts = read_workbook(path)

    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    assert tuple(workbook.sheetnames) == SHEETS
    assert workbook.properties.title == "小安評估結果"
    overview = workbook["00_Overview"]
    assert [cell.value for cell in overview[2]] == [
        "結果", "交付物狀態", "已完成", "已完成", "此工作簿目前的生命週期狀態。",
    ]
    assert overview["A1"].comment.text == "區段（section）"
    assert workbook["01_Cases"]["D1"].comment.text.startswith("狀態（status）")
    header_comments = [cell.comment.text for sheet in workbook.worksheets for cell in sheet[1]]
    assert all(any("\u4e00" <= character <= "\u9fff" for character in text) for text in header_comments)
    dictionary_descriptions = [
        row[3].value for row in workbook["09_Data_Dictionary"].iter_rows(min_row=2)
    ]
    assert all(any("\u4e00" <= character <= "\u9fff" for character in str(value)) for value in dictionary_descriptions)
    assert facts.generation_id == model.generation_id
    assert facts.artifact_state == "FINAL"
    assert facts.cases[0]["case_id"] == "TC-01"
    assert facts.metrics[0]["raw_score"] == 0
    assert facts.metrics[0]["status"] == "FAIL"
    assert facts.metrics[0]["reason"] == "路由選擇錯誤"
    assert facts.manifest["rating_rule_hash"] == "rule-1"
    user_row = next(row for row in facts.text_content if row["role"] == "user_input")
    assert user_row["content"] == "=1+1"
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    assert "10_Text_Content" not in workbook.sheetnames


def test_overview_localizes_optional_evaluation_states(tmp_path: Path):
    empty_path = tmp_path / "empty.xlsx"
    write_workbook(build_report_model([]), empty_path)
    from openpyxl import load_workbook
    empty_rows = {
        row[1].value: (row[2].value, row[3].value)
        for row in load_workbook(empty_path)["00_Overview"].iter_rows(min_row=2)
    }

    assert empty_rows["整體分數"] == (None, "不可用")
    assert empty_rows["比較結果"] == ("未執行", "未執行")
    assert empty_rows["實驗結果"] == ("未執行", "未執行")
    assert empty_rows["分類"] == ("未測量", "未執行")

    attached_path = tmp_path / "attached.xlsx"
    model = build_report_model(
        [_record()],
        baseline=[{"domain": "quality"}],
        experiments=[{"experiment_id": "EXP-1"}],
        stability={"classification": "DESCRIPTIVE_ONLY"},
    )
    write_workbook(model, attached_path)
    attached_rows = {
        row[1].value: row[2].value
        for row in load_workbook(attached_path)["00_Overview"].iter_rows(min_row=2)
    }

    assert attached_rows["比較結果"] == "已附上"
    assert attached_rows["實驗結果"] == "已附上"
    assert attached_rows["分類"] == "僅描述性結果"


def test_long_text_is_stored_in_ordered_chunks(tmp_path: Path):
    response = "長" * 65000
    model = build_report_model([_record(response)])
    path = tmp_path / "results.xlsx"

    write_workbook(model, path)
    facts = read_workbook(path)

    chunks = sorted((row for row in facts.text_content if row["role"] == "assistant_response"), key=lambda row: row["chunk_index"])
    assert len(chunks) == 3
    assert "".join(row["content"] for row in chunks) == response
    assert {row["chunk_count"] for row in chunks} == {3}


def test_workbook_normalizes_floats_to_excel_precision_before_round_trip(tmp_path: Path):
    model = build_report_model([_record()])
    value = 0.19047619047619055
    metrics = tuple(
        {**row, "raw_score": value, "normalized_score": value}
        if index == 0 else row
        for index, row in enumerate(model.metrics)
    )
    model = replace(model, metrics=metrics)
    path = tmp_path / "results.xlsx"

    write_workbook(model, path)
    facts = read_workbook(path)

    assert facts.metrics[0]["raw_score"] == float(format(value, ".15g"))


def test_import_rejects_formula_cells(tmp_path: Path):
    model = build_report_model([_record()])
    path = tmp_path / "results.xlsx"
    write_workbook(model, path)
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    workbook["01_Cases"]["B2"] = "=1+1"
    workbook.save(path)

    with pytest.raises(ValueError, match="formulas are not allowed"):
        read_workbook(path)


def test_pending_review_is_a_lifecycle_state_not_a_score():
    record = _record()
    record["status"] = "NEEDS_REVIEW"
    record["review"] = {"status": "NEEDS_REVIEW", "per_turn": ["NEEDS_REVIEW"], "judge_audit": []}

    model = build_report_model([record])

    assert model.artifact_state == "PENDING_REVIEW"
    assert model.human_review[0]["status"] == "PENDING_REVIEW"
