import json
from openpyxl import load_workbook
import pytest
from test_scenario_analysis import inputs
from xiaoan_eval.scenario_analysis import build_analysis, attach_to_model
from xiaoan_eval.report_model import build_report_model, build_model_from_rows
from xiaoan_eval.workbook import write_workbook, read_workbook, LEGACY_SHEETS, HEADERS
from xiaoan_eval.detail_tables import NAMES, analysis_from_metrics


def test_details_long_evidence_roundtrip_and_tamper(tmp_path):
    r,a,f,_=inputs(tmp_path)
    evidence='長證據🙂'*18000
    r['pipeline']['observations'][0]['attribution']['claims'][0]['relations'][0]['evidence_span']={'start':0,'end':len(evidence),'text':evidence}
    analysis=build_analysis([r],annotations_path=a,analysis_path=f,subject_id='model-a')
    model=attach_to_model(build_report_model([r]),analysis)
    path=tmp_path/'results.xlsx';write_workbook(model,path);facts=read_workbook(path)
    restored=analysis_from_metrics(facts.metrics)
    assert restored['claims']==analysis['claims']
    book=load_workbook(path)
    assert all(n in book for n in NAMES)
    ws=book['Evidence'];cols=[c.value for c in ws[1]]
    rows=[dict(zip(cols,r)) for r in ws.iter_rows(min_row=2,values_only=True)]
    selected=[r for r in rows if r['claim_id']=='c1' and r['evidence_number']==1]
    assert ''.join(r['evidence_text'] for r in sorted(selected,key=lambda r:r['part']))==evidence
    ws['G2']='changed';book.save(path)
    with pytest.raises(ValueError,match='detail validation'):read_workbook(path)


def test_old_workbook_read_and_rebuild_keeps_details(tmp_path):
    r,a,f,_=inputs(tmp_path)
    model=attach_to_model(build_report_model([r]),build_analysis([r],annotations_path=a,analysis_path=f,subject_id='model-a'))
    path=tmp_path/'old.xlsx';write_workbook(model,path)
    book=load_workbook(path)
    for name in NAMES:del book[name]
    ws=book['02_Turns'];ws.delete_cols(len(HEADERS['02_Turns'])+1,ws.max_column-len(HEADERS['02_Turns']))
    # Old readers have only canonical turn/text rows; projections are derived.
    for row in range(ws.max_row,1,-1):
        if ws.cell(row,1).value=='diagnostic':ws.delete_rows(row)
    book.save(path);facts=read_workbook(path)
    rebuilt=build_model_from_rows(manifest=facts.manifest,artifact_state=facts.artifact_state,cases=facts.cases,turns=facts.turns,metrics=facts.metrics,text_content=facts.text_content)
    new=tmp_path/'new.xlsx';write_workbook(rebuilt,new);read_workbook(new)
    b=load_workbook(new);assert b['Claims'].max_row==3
    assert 'Review_Turns' not in b
    h=[c.value for c in b['02_Turns'][1]];assert 'task' in h


def test_matrix_pair_projection_and_ar_deduplication(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from xiaoan_eval import multimodel as matrix
    from xiaoan_eval.detail_tables import project
    r,a,f,_=inputs(tmp_path)
    rows=[{'case_id':'TC-01','turn':1,'subject':{'id':'model-a'},'judge':{'id':j},
           'answer':{'text':'alpha beta','status':'PASS','trace':r['pipeline']['turn_traces'][0]['trace']},
           'attribution':{'status':'AVAILABLE','result':{'claims':r['pipeline']['observations'][0]['attribution']['claims']}}}
          for j in ['j1','j2']]
    case=SimpleNamespace(id='TC-01',turns=[SimpleNamespace(turn=1,user='help')],quality_focus=[])
    analysis=build_analysis(rows,cases=[case],annotations_path=a,analysis_path=f)
    _,tables=project(analysis)
    assert len(tables['Relevancy'])==3 and len(tables['Claims'])==4
    def capture(rows,path,*,analysis=None):
        from openpyxl import Workbook
        b=Workbook();b.active.title='All_Answers';b.active.append(['case_id','turn','subject_id','answer'])
        b.active.append(['TC-01',1,'model-a','alpha beta']);b.save(path)
        from xiaoan_eval.scenario_analysis import append_matrix_sheets
        append_matrix_sheets(path,analysis)
    monkeypatch.setattr(matrix,'write_matrix_workbook',capture)
    paths=matrix.write_pair_workbooks(rows,tmp_path/'pairs',analysis=analysis)
    assert len(paths)==2
    for p in paths:
        b=load_workbook(p)
        assert all(n in b for n in NAMES)
        assert b['Claims'].max_row==3 and b['Relevancy'].max_row==4
        assert 'task' in [c.value for c in b['All_Answers'][1]]


def test_large_manifest_metadata_survives_details_export(tmp_path):
    r,_,_,_=inputs(tmp_path)
    model=build_report_model([r],manifest={'large_evidence':'測試🙂'*20000})
    p=tmp_path/'large.xlsx';write_workbook(model,p)
    assert read_workbook(p).manifest==model.manifest
