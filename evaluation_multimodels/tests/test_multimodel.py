import json
import fcntl
from pathlib import Path
import threading
import time
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from xiaoan_eval import cli
from xiaoan_eval.multimodel import (
    JudgeEvidenceError,
    ModelSpec,
    build_judge_evidence,
    default_judge_specs,
    default_subject_specs,
    render_matrix_report,
    run_matrix,
    write_matrix_workbook,
    write_pair_workbooks,
)
from xiaoan_eval.rules import load_rating_rule
from tools.compare_matrix_runs import compare
from xiaoan_eval.cases import (
    OracleProvenance,
    TestCase as EvaluationCase,
    TestTurn as EvaluationTurn,
    load_cases,
)


def _judge_json(rule, score: int, *, triggered: str | None = None) -> str:
    return json.dumps({
        "red_lines": [
            {"id": item.id, "triggered": item.id == triggered, "evidence": ["unsafe span"] if item.id == triggered else []}
            for item in rule.red_lines
        ],
        "dimensions": {module.name: score for module in rule.modules},
    }, ensure_ascii=False)


def test_default_matrix_has_ten_subjects_and_five_judges(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    subjects = default_subject_specs()
    judges = default_judge_specs()
    assert [spec.model for spec in subjects] == [
        "claude-opus-5",
        "claude-sonnet-5",
        "gpt-5.6-sol",
        "o4-mini-2025-04-16",
        "gemini-3-pro-preview-thinking",
        "gemini-3.8-flash",
        "qwen3.8-max",
        "qwen3.7-max",
        "kimi-k3",
        "kimi-k2.6",
    ]
    assert [spec.model for spec in judges] == [
        "claude-opus-5",
        "gpt-5.6-sol",
        "gemini-3-pro-preview-thinking",
        "qwen3.8-max",
        "kimi-k3",
    ]

    def subject(_spec, prompt):
        return {"text": "A" + prompt, "usage": {"prompt_tokens": 2, "completion_tokens": 3}}

    def judge(_spec, _prompt):
        return {"text": _judge_json(rule, 3), "usage": {"input_tokens": 4, "output_tokens": 5}}

    rows = run_matrix([{"id": "TC-01", "turns": [{"turn": 1, "user": "hello"}]}], subjects, judges, subject_transport=subject, judge_transport=judge, rating_rule=rule)
    assert len(subjects) == 10
    assert len(rows) == 50
    assert rows[0]["answer"]["first_char"] == "A"
    assert rows[0]["answer"]["total_tokens"] == 5
    assert "XiaoAn subject" in render_matrix_report(rows)
    write_matrix_workbook(rows, tmp_path / "results.xlsx")
    assert (tmp_path / "results.xlsx").exists()
    workbook = load_workbook(tmp_path / "results.xlsx", read_only=True)
    assert "Dimension_By_Judge" in workbook.sheetnames
    assert "Judge_Agreement" in workbook.sheetnames
    assert "Measurement_Contract" in workbook.sheetnames
    assert "dimension:基础能力" in next(workbook["All_Judgements"].iter_rows(values_only=True))
    pair_paths = write_pair_workbooks(rows, tmp_path / "pairs")
    assert len(pair_paths) == 50
    assert all(path.exists() for path in pair_paths)


def test_matrix_accepts_loaded_testcase_dataclasses() -> None:
    rule = load_rating_rule("ratings rule.yml")
    loaded = load_cases("test-cases/TC-04.yaml", rule)
    case = loaded[0].case
    assert case is not None

    subjects = (ModelSpec("gpt", "subject", "latest", "medium"),)
    judges = (ModelSpec("gpt", "judge", "judge", "medium"),)

    def subject(_spec, prompt):
        return {"text": f"answer:{prompt}", "usage": {"prompt_tokens": 1, "completion_tokens": 2}}

    def judge(_spec, prompt):
        assert '"case_id": "TC-04"' in prompt
        assert '"turn": 1' in prompt
        return {"text": _judge_json(rule, 2)}

    rows = run_matrix((case,), subjects, judges, subject_transport=subject, judge_transport=judge, rating_rule=rule)
    assert rows
    assert rows[0]["case_id"] == "TC-04"
    assert rows[0]["turn"] == 1


def test_matrix_accepts_testcase_and_testturn_dataclasses_directly() -> None:
    rule = load_rating_rule("ratings rule.yml")
    case = EvaluationCase(
        schema_version="2.0",
        id="TC-DATACLASS",
        test_objective="dataclass compatibility",
        quality_focus=(),
        turns=(EvaluationTurn(turn=7, user="hello dataclass"),),
        oracle_provenance=OracleProvenance(source="test", status="reviewed"),
    )
    subjects = (ModelSpec("gpt", "subject", "latest", "medium"),)
    judges = (ModelSpec("gpt", "judge", "judge", "medium"),)
    rows = run_matrix(
        (case,), subjects, judges,
        subject_transport=lambda _spec, prompt: {"text": prompt},
        judge_transport=lambda _spec, _prompt: {"text": _judge_json(rule, 1)},
        rating_rule=rule,
    )
    assert rows[0]["case_id"] == "TC-DATACLASS"
    assert rows[0]["turn"] == 7
    assert rows[0]["answer"]["text"] == "hello dataclass"


def _trace_with_large_duplicate_context() -> dict:
    ground_item = {
        "ref": "personal-safety-protection-order",
        "title": "Protection order",
        "content": "authoritative ground text",
    }
    return {
        "route": {
            "capsule_id": "N5p",
            "confidence": 0.93,
            "fallback_reason": "",
            "raw_candidate": "must not reach judge",
        },
        "safety": {"level": "normal", "reason": "no immediate danger"},
        "capsule": {"id": "N5p", "version": "2", "injected_units": ["duplicate"]},
        "ground": {
            "loaded": True,
            "resolved_ground": ["personal-safety-protection-order"],
            "resolved_items": [ground_item],
        },
        "state": {"active_capsule_id": "N5p", "ttl_turns": 2},
        "output_guard": {"passed": True, "warnings": []},
        "effective_context_snapshot": {
            "schema_version": "effective-context-snapshot/v1",
            "snapshot_id": "snapshot-1",
            "turn": 1,
            "context_kind": "ORDINARY_CAPSULE",
            "invocations": {
                "router": {
                    "status": "INVOKED",
                    "provider_request": {"messages": [{"content": "R" * 100_000}]},
                    "context_units": [{"layer": "PROMPT", "content": "R" * 50_000}],
                },
                "composer": {
                    "status": "INVOKED",
                    "provider_request": {"messages": [{"content": "C" * 100_000}]},
                    "context_units": [
                        {
                            "ref": "composer:prompt/main-agent/instructions",
                            "layer": "PROMPT",
                            "field_path": "instructions",
                            "content": "P" * 50_000,
                        },
                        {
                            "ref": "composer:capsule/N5p/response_policy/response_policy",
                            "layer": "CAPSULE",
                            "entity_id": "N5p",
                            "field_path": "response_policy",
                            "item_id": "response_policy",
                            "content": {"rules": [{"id": "keep-safe", "effect": "support agency"}]},
                        },
                        {
                            "ref": "composer:knowledge/wiki/nodes/personal-safety-protection-order",
                            "layer": "WIKI",
                            "entity_id": "personal-safety-protection-order",
                            "field_path": "node",
                            "item_id": "ground:0",
                            "content": ground_item,
                        },
                    ],
                },
            },
        },
        "timings": {"total_ms": 999},
        "tokens": {"input": 99999},
    }


def test_judge_evidence_is_compact_deterministic_and_keeps_semantic_support() -> None:
    trace = _trace_with_large_duplicate_context()

    evidence = build_judge_evidence(trace)
    encoded = json.dumps(evidence, ensure_ascii=False, sort_keys=True)
    full = json.dumps(trace, ensure_ascii=False, sort_keys=True)

    assert evidence == build_judge_evidence(trace)
    assert evidence["schema_version"] == "judge-evidence/v1"
    assert evidence["trace_sha256"].startswith("sha256:")
    assert evidence["evidence_integrity"] == {"status": "COMPLETE", "unresolved_refs": []}
    assert evidence["route"]["capsule_id"] == "N5p"
    assert evidence["context"]["units"][0]["field_path"] == "response_policy"
    assert evidence["context"]["units"][1]["content"]["content"] == "authoritative ground text"
    assert "provider_request" not in encoded
    assert "raw_candidate" not in encoded
    assert "timings" not in encoded
    assert "tokens" not in encoded
    assert len(encoded) < len(full) * 0.25


def test_judge_evidence_fails_closed_when_ground_ref_has_no_evidence_text() -> None:
    trace = _trace_with_large_duplicate_context()
    trace["ground"]["resolved_items"] = []
    trace["effective_context_snapshot"]["invocations"]["composer"]["context_units"] = []

    with pytest.raises(JudgeEvidenceError, match="unresolved ground refs"):
        build_judge_evidence(trace)


def test_judge_evidence_rejects_matching_ref_with_empty_body() -> None:
    trace = _trace_with_large_duplicate_context()
    trace["ground"]["resolved_items"] = []
    ground_unit = trace["effective_context_snapshot"]["invocations"]["composer"]["context_units"][2]
    ground_unit["content"] = {"ref": "personal-safety-protection-order", "text": ""}

    with pytest.raises(JudgeEvidenceError, match="unresolved ground refs"):
        build_judge_evidence(trace)


def test_judge_evidence_rejects_non_array_ground_refs() -> None:
    trace = _trace_with_large_duplicate_context()
    trace["ground"]["resolved_ground"] = "personal-safety-protection-order"

    with pytest.raises(JudgeEvidenceError, match="must be an array"):
        build_judge_evidence(trace)


def test_incomplete_judge_scores_are_unavailable() -> None:
    rule = load_rating_rule("ratings rule.yml")
    partial_scores = {rule.modules[0].name: 3}
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda _spec, _prompt: {"text": "answer"},
        judge_transport=lambda _spec, _prompt: {"text": json.dumps({"dimensions": partial_scores}, ensure_ascii=False)},
        rating_rule=rule,
    )
    assert rows[0]["status"] == "UNAVAILABLE"
    assert rows[0]["scores"] == {}
    assert rows[0]["weighted_score"] is None


def test_matrix_red_line_contract_zeroes_every_quality_dimension() -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "unsafe answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3, triggered=rule.red_lines[0].id)},
        rating_rule=rule,
    )
    assert rows[0]["status"] == "PASS"
    assert rows[0]["triggered_red_lines"] == (rule.red_lines[0].id,)
    assert set(rows[0]["scores"].values()) == {0}
    assert rows[0]["weighted_score"] == 0


def test_matrix_requires_every_red_line_id_exactly_once() -> None:
    rule = load_rating_rule("ratings rule.yml")
    invalid = json.loads(_judge_json(rule, 2))
    invalid["red_lines"] = invalid["red_lines"][:-1]
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": json.dumps(invalid)},
        rating_rule=rule,
    )
    assert rows[0]["status"] == "UNAVAILABLE"
    assert rows[0]["weighted_score"] is None


def test_matrix_rejects_non_string_red_line_evidence() -> None:
    rule = load_rating_rule("ratings rule.yml")
    invalid = json.loads(_judge_json(rule, 2))
    invalid["red_lines"][0]["evidence"] = [123]
    rows = run_matrix(
        ({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": json.dumps(invalid)}, rating_rule=rule,
    )
    assert rows[0]["status"] == "UNAVAILABLE"
    assert rows[0]["scores"] == {}


def test_matrix_requires_non_empty_evidence_for_triggered_red_line() -> None:
    rule = load_rating_rule("ratings rule.yml")
    invalid = json.loads(_judge_json(rule, 2))
    invalid["red_lines"][0].update({"triggered": True, "evidence": []})
    rows = run_matrix(
        ({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": json.dumps(invalid)}, rating_rule=rule,
    )
    assert rows[0]["status"] == "UNAVAILABLE"
    assert rows[0]["weighted_score"] is None


def test_matrix_marks_self_judging_outside_primary_denominator() -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "same", "latest", "medium"),),
        (ModelSpec("gpt", "same", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)},
        rating_rule=rule,
    )
    assert rows[0]["self_judging"] is True
    assert rows[0]["primary_eligible"] is False
    assert "SELF_ISOLATED" in render_matrix_report(rows)


def test_workbook_main_matrix_isolates_self_score_in_separate_sheet(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "same", "latest", "medium"),),
        (ModelSpec("gpt", "same", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3)}, rating_rule=rule,
    )
    path = tmp_path / "results.xlsx"
    write_matrix_workbook(rows, path)
    workbook = load_workbook(path, read_only=True)
    assert workbook["Matrix"].cell(2, 2).value == "SELF_ISOLATED"
    isolated = list(workbook["Self_Judging_Isolated"].iter_rows(values_only=True))
    assert isolated[1][5] == 3


def test_first_character_latency_is_only_copied_from_explicit_telemetry() -> None:
    rule = load_rating_rule("ratings rule.yml")
    subjects = (
        ModelSpec("gpt", "with-ttfc", "latest", "medium"),
        ModelSpec("claude", "without-ttfc", "latest", "medium"),
    )
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},), subjects,
        (ModelSpec("qwen", "judge", "judge", "medium"),),
        subject_transport=lambda spec, _prompt: ({"text": "answer", "_xiaoan_first_character_ms": 12.5} if spec.model == "with-ttfc" else {"text": "answer"}),
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)}, rating_rule=rule,
    )
    assert rows[0]["answer"]["first_character_ms"] == 12.5
    assert rows[1]["answer"]["first_character_ms"] is None


def test_matrix_optional_attribution_pass_is_bound_to_frozen_answer_snapshot() -> None:
    rule = load_rating_rule("ratings rule.yml")
    trace = {
        "effective_context_snapshot": {
            "schema_version": "effective-context-snapshot/v1", "snapshot_id": "snap-1",
            "turn": 1, "context_kind": "BASELINE",
            "router": {"status": "NOT_APPLICABLE", "units": []},
            "composer": {"status": "INVOKED", "units": []},
        }
    }
    seen = []

    def attribution(request):
        seen.append(request["assistant_answer"])
        return json.dumps({
            "contract_version": "attribution/v1",
            "claims": [{
                "claim_id": "c1", "kind": "SUPPORTIVE",
                "answer_span": {"start": 0, "end": 6, "text": "answer"},
                "relations": [{"relation": "UNSUPPORTED", "evidence_ref": None, "evidence_span": None}],
                "unsupported_category": "NON_FACTUAL_SUPPORTIVE", "uncertainty": "LOW",
            }],
            "policies": [], "abstention": {"status": "ANSWERED", "reason": None},
        })

    rows = run_matrix(
        ({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer", "chatflow_debug": trace},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)},
        rating_rule=rule, attribution_provider=attribution,
    )
    assert seen == ["answer"]
    assert rows[0]["attribution"]["status"] == "AVAILABLE"
    assert rows[0]["attribution"]["result"]["claims"][0]["claim_id"] == "c1"


def test_invalid_attribution_is_unavailable_without_becoming_quality_zero() -> None:
    rule = load_rating_rule("ratings rule.yml")
    trace = {"effective_context_snapshot": {"schema_version": "effective-context-snapshot/v1", "snapshot_id": "snap", "turn": 1, "context_kind": "BASELINE", "router": {"status": "NOT_APPLICABLE", "units": []}, "composer": {"status": "INVOKED", "units": []}}}
    rows = run_matrix(
        ({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer", "chatflow_debug": trace},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)},
        rating_rule=rule, attribution_provider=lambda _request: "not-json",
    )
    assert rows[0]["status"] == "PASS"
    assert rows[0]["weighted_score"] == pytest.approx(2)
    assert rows[0]["attribution"]["status"] == "UNAVAILABLE"


def test_matrix_collects_memory_lifecycle_from_structured_trace() -> None:
    rule = load_rating_rule("ratings rule.yml")
    case = {
        "id": "TC-memory",
        "memory_checkpoints": [{"after_turn": 1, "facts": ["has-child"], "usage": "use it"}],
        "turns": ({"turn": 1, "user": "hello"},),
    }
    rows = run_matrix(
        (case,), (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer", "chatflow_debug": {"state": {"memory_facts": ["has-child"], "memory_used": True, "memory_isolated": True, "memory_stale_or_unsafe": False}}},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)}, rating_rule=rule,
    )
    assert rows[0]["memory_metrics"][0]["status"] == "pass"
    assert rows[0]["memory_metrics"][0]["isolated"] is None


def test_memory_checkpoint_types_do_not_all_require_memory_used_true() -> None:
    rule = load_rating_rule("ratings rule.yml")
    case = {
        "id": "TC-memory-types",
        "memory_checkpoints": [
            {"after_turn": 1, "facts": ["forbidden"], "usage": "must not use", "type": "not_use"},
            {"after_turn": 1, "facts": ["tenant-a"], "usage": "isolate", "type": "isolation"},
        ],
        "turns": ({"turn": 1, "user": "hello"},),
    }
    state = {
        "memory_used": True,
        "memory_used_facts": ["allowed"],
        "memory_isolated": False,
        "memory_contamination_candidates": ["tenant-b"],
    }
    rows = run_matrix(
        (case,), (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer", "chatflow_debug": {"state": state}},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)}, rating_rule=rule,
    )
    assert [item["status"] for item in rows[0]["memory_metrics"]] == ["pass", "fail"]
    assert rows[0]["memory_metrics"][1]["contamination_candidates"] == ("tenant-b",)


def test_matrix_checkpoints_each_completed_provider_call(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "nested" / "matrix-checkpoint.jsonl"
    judge_calls = 0

    def subject(_spec, _prompt):
        return {"text": "answer"}

    def judge(_spec, _prompt):
        nonlocal judge_calls
        events = [json.loads(line) for line in checkpoint.read_text(encoding="utf-8").splitlines()]
        if judge_calls == 0:
            assert [event["event"] for event in events] == ["answer"]
        else:
            assert [event["event"] for event in events] == [
                "answer",
                "judgement",
                "cell",
            ]
        judge_calls += 1
        return {
            "text": _judge_json(rule, 2)
        }

    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (
            ModelSpec("gpt", "judge-a", "judge", "medium"),
            ModelSpec("claude", "judge-b", "judge", "medium"),
        ),
        subject_transport=subject,
        judge_transport=judge,
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )

    events = [json.loads(line) for line in checkpoint.read_text(encoding="utf-8").splitlines()]
    assert len(rows) == 2
    assert checkpoint.stat().st_mode & 0o777 == 0o600
    assert [event["event"] for event in events] == [
        "answer",
        "judgement",
        "cell",
        "judgement",
        "cell",
    ]
    assert events[0]["answer"]["text"] == "answer"
    assert events[0]["answer_id"] == rows[0]["answer_id"]
    assert events[1]["judge"]["model"] == "judge-a"
    assert events[3]["judge"]["model"] == "judge-b"
    assert "answer" not in events[2]["row"]
    assert "judgement" not in events[2]["row"]
    assert events[2]["row"]["answer_id"] == rows[0]["answer_id"]
    assert rows[0]["answer"] is rows[1]["answer"]


def test_matrix_checkpoints_judge_before_validating_response_shape(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda _spec, _prompt: {"text": "answer"},
        judge_transport=lambda _spec, _prompt: {"text": "[]"},
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )

    events = [json.loads(line) for line in checkpoint.read_text(encoding="utf-8").splitlines()]
    assert [event["event"] for event in events] == ["answer", "judgement", "cell"]
    assert events[1]["judgement"]["text"] == "[]"
    assert rows[0]["status"] == "UNAVAILABLE"
    assert events[2]["row"]["status"] == "UNAVAILABLE"


def test_matrix_refuses_to_mix_runs_in_one_checkpoint(tmp_path: Path) -> None:
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    checkpoint.write_text('{"event":"answer"}\n', encoding="utf-8")

    with pytest.raises(ValueError, match="checkpoint already exists"):
        run_matrix(
            (),
            (),
            (),
            subject_transport=lambda *_args: {},
            judge_transport=lambda *_args: {},
            rating_rule=load_rating_rule("ratings rule.yml"),
            checkpoint_path=checkpoint,
        )


def test_matrix_resume_rejects_legacy_checkpoint_without_explicit_opt_in(tmp_path: Path) -> None:
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    checkpoint.write_text('{"event":"answer"}\n', encoding="utf-8")

    with pytest.raises(ValueError, match="legacy matrix checkpoint"):
        run_matrix(
            (),
            (),
            (),
            subject_transport=lambda *_args: {},
            judge_transport=lambda *_args: {},
            rating_rule=load_rating_rule("ratings rule.yml"),
            checkpoint_path=checkpoint,
            resume=True,
        )


def test_matrix_resume_ignores_only_a_truncated_final_record(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    common = dict(
        cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("gpt", "judge", "judge", "medium"),),
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )
    expected = run_matrix(
        **common,
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3)},
    )
    with checkpoint.open("a", encoding="utf-8") as target:
        target.write('{"event":"torn"')

    resumed = run_matrix(
        **common,
        resume=True,
        subject_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("subject repeated")),
        judge_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("judge repeated")),
    )

    assert resumed == expected
    assert not checkpoint.read_text(encoding="utf-8").endswith('{"event":"torn"')


def test_matrix_cli_creates_checkpoint_directory_before_running(tmp_path: Path, monkeypatch) -> None:
    output = tmp_path / "new-run"
    args = SimpleNamespace(
        cases="unused",
        output=str(output),
        subject_transport="plugins:subject",
        judge_transport="plugins:judge",
        rating_rule="unused",
        subjects=None,
        judges=None,
    )
    monkeypatch.setattr(cli, "load_rating_rule", lambda _path: object())
    monkeypatch.setattr(
        cli,
        "load_cases",
        lambda *_args: [SimpleNamespace(case={"id": "TC-01", "turns": ()})],
    )
    monkeypatch.setattr(cli, "_load_plugin", lambda _reference: lambda *_args: {})
    monkeypatch.setattr(cli, "default_subject_specs", lambda: ())
    monkeypatch.setattr(cli, "_load_matrix_specs", lambda _path, defaults: tuple(defaults))

    def run_matrix_at_existing_path(*_args, checkpoint_path, **_kwargs):
        assert checkpoint_path == tmp_path / ".matrix-audit" / "new-run" / "matrix-checkpoint.jsonl"
        assert checkpoint_path.parent.is_dir()
        return []

    monkeypatch.setattr(cli, "run_matrix", run_matrix_at_existing_path)
    monkeypatch.setattr(cli, "write_matrix_workbook", lambda *_args: None)
    monkeypatch.setattr(cli, "write_pair_workbooks", lambda *_args: [])
    monkeypatch.setattr(cli, "render_matrix_report", lambda _rows: "")

    assert cli._matrix(args) == 1
    assert (output / "report.md").exists()


def test_workbook_forces_untrusted_text_to_string(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda _spec, _prompt: {"text": "=1+1"},
        judge_transport=lambda _spec, _prompt: {"text": _judge_json(rule, 3)},
        rating_rule=rule,
    )
    path = tmp_path / "results.xlsx"
    write_matrix_workbook(rows, path)
    workbook = load_workbook(path, data_only=False)
    answer = workbook["All_Answers"]["F2"]
    assert answer.value == "=1+1"
    assert answer.data_type == "s"


def test_matrix_runs_judges_concurrently_and_keeps_declared_order() -> None:
    rule = load_rating_rule("ratings rule.yml")
    judges = tuple(ModelSpec("gpt", f"judge-{index}", "judge", "medium") for index in range(3))
    lock = threading.Lock()
    active = 0
    maximum_active = 0

    def judge(_spec, _prompt):
        nonlocal active, maximum_active
        with lock:
            active += 1
            maximum_active = max(maximum_active, active)
        time.sleep(0.03)
        with lock:
            active -= 1
        return {"text": _judge_json(rule, 2)}

    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),), judges,
        subject_transport=lambda *_args: {"text": "answer"}, judge_transport=judge,
        rating_rule=rule, judge_concurrency=3, max_in_flight=3,
        per_provider_concurrency=3,
    )

    assert maximum_active == 3
    assert [row["judge"]["model"] for row in rows] == [spec.model for spec in judges]


def test_matrix_runs_subject_case_lanes_concurrently_but_turns_stay_sequential() -> None:
    rule = load_rating_rule("ratings rule.yml")
    lock = threading.Lock()
    active = 0
    maximum_active = 0
    seen: dict[str, list[str]] = {}

    class Subject:
        local = threading.local()

        def start_case(self, spec, case_id):
            self.local.lane = f"{spec.model}:{case_id}"
            seen[self.local.lane] = []

        def __call__(self, _spec, prompt):
            nonlocal active, maximum_active
            lane = self.local.lane
            with lock:
                active += 1
                maximum_active = max(maximum_active, active)
            time.sleep(0.03)
            seen[lane].append(prompt)
            with lock:
                active -= 1
            return {"text": f"{lane}:{prompt}"}

        def end_case(self, _spec, _case_id):
            self.local.lane = None

    cases = tuple(
        {"id": f"TC-{index}", "turns": ({"turn": 1, "user": "one"}, {"turn": 2, "user": "two"})}
        for index in range(2)
    )
    rows = run_matrix(
        cases,
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=Subject(),
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)},
        rating_rule=rule,
        subject_concurrency=2,
        max_in_flight=2,
        per_provider_concurrency=2,
    )

    assert maximum_active == 2
    assert all(prompts == ["one", "two"] for prompts in seen.values())
    assert [(row["case_id"], row["turn"]) for row in rows] == [
        ("TC-0", 1), ("TC-0", 2), ("TC-1", 1), ("TC-1", 2)
    ]


def test_matrix_finishes_all_subject_answers_before_any_judge_call() -> None:
    rule = load_rating_rule("ratings rule.yml")
    events = []
    cases = (
        {"id": "TC-1", "turns": ({"turn": 1, "user": "one"},)},
        {"id": "TC-2", "turns": ({"turn": 1, "user": "two"},)},
    )

    def subject(_spec, prompt):
        events.append(("subject", prompt))
        return {"text": f"answer:{prompt}"}

    def judge(_spec, _prompt):
        events.append(("judge", None))
        assert sum(kind == "subject" for kind, _ in events) == 2
        return {"text": _judge_json(rule, 2)}

    run_matrix(
        cases, (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=subject, judge_transport=judge, rating_rule=rule,
    )
    assert [kind for kind, _ in events] == ["subject", "subject", "judge", "judge"]


def test_matrix_limits_same_provider_calls() -> None:
    rule = load_rating_rule("ratings rule.yml")
    lock = threading.Lock()
    active = 0
    maximum_active = 0

    def judge(_spec, _prompt):
        nonlocal active, maximum_active
        with lock:
            active += 1
            maximum_active = max(maximum_active, active)
        time.sleep(0.02)
        with lock:
            active -= 1
        return {"text": _judge_json(rule, 2)}

    run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        tuple(ModelSpec("gpt", f"judge-{index}", "judge", "medium") for index in range(3)),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=judge,
        rating_rule=rule,
        judge_concurrency=3,
        max_in_flight=3,
        per_provider_concurrency=1,
    )

    assert maximum_active == 1


def test_matrix_resume_reuses_completed_lanes(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    args = dict(
        cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("gpt", "judge", "judge", "medium"),),
        rating_rule=rule, checkpoint_path=checkpoint,
    )
    first = run_matrix(
        **args,
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3)},
    )

    second = run_matrix(
        **args, resume=True,
        subject_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("subject repeated")),
        judge_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("judge repeated")),
    )

    assert second == first


def test_matrix_resume_reuses_attribution_without_duplicate_event(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    calls = 0
    trace = {"effective_context_snapshot": {"schema_version": "effective-context-snapshot/v1", "snapshot_id": "snap", "turn": 1, "context_kind": "BASELINE", "router": {"status": "NOT_APPLICABLE", "units": []}, "composer": {"status": "INVOKED", "units": []}}}

    def attribution(_request):
        nonlocal calls
        calls += 1
        return json.dumps({"contract_version": "attribution/v1", "claims": [], "policies": [], "abstention": {"status": "ANSWERED", "reason": None}})

    args = dict(
        cases=({"id": "TC", "turns": ({"turn": 1, "user": "hello"},)},),
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("claude", "judge", "judge", "medium"),),
        rating_rule=rule, checkpoint_path=checkpoint,
        attribution_provider=attribution,
    )
    first = run_matrix(
        **args,
        subject_transport=lambda *_args: {"text": "answer", "chatflow_debug": trace},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 2)},
    )
    event_count = sum(json.loads(line)["event"] == "attribution" for line in checkpoint.read_text(encoding="utf-8").splitlines())
    second = run_matrix(
        **args, resume=True,
        subject_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("subject repeated")),
        judge_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("judge repeated")),
    )
    assert second == first
    assert calls == 1
    assert sum(json.loads(line)["event"] == "attribution" for line in checkpoint.read_text(encoding="utf-8").splitlines()) == event_count


def test_matrix_resume_preserves_validated_error_class(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    kwargs = dict(
        cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": "not-json"},
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )
    first = run_matrix(**kwargs)
    resumed = run_matrix(**kwargs, resume=True)

    assert first[0]["judgement"]["error_type"] == "INVALID_JUDGE_JSON"
    assert resumed[0]["judgement"]["error_type"] == "INVALID_JUDGE_JSON"


def test_matrix_refuses_checkpoint_locked_by_another_writer(tmp_path: Path) -> None:
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    checkpoint.touch()
    with checkpoint.open("a", encoding="utf-8") as locked:
        fcntl.flock(locked.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        with pytest.raises(RuntimeError, match="already locked"):
            run_matrix(
                (), (), (), subject_transport=lambda *_args: {}, judge_transport=lambda *_args: {},
                rating_rule=load_rating_rule("ratings rule.yml"), checkpoint_path=checkpoint,
                resume=True, allow_legacy_checkpoint=True,
            )


def test_matrix_resume_reuses_judgement_written_before_cell(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    args = dict(
        cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("gpt", "judge", "judge", "medium"),),
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )
    first = run_matrix(
        **args,
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3)},
    )
    events = [json.loads(line) for line in checkpoint.read_text(encoding="utf-8").splitlines()]
    checkpoint.write_text(
        "\n".join(json.dumps(event, ensure_ascii=False) for event in events if event["event"] != "cell") + "\n",
        encoding="utf-8",
    )

    resumed = run_matrix(
        **args,
        resume=True,
        subject_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("subject repeated")),
        judge_transport=lambda *_args: (_ for _ in ()).throw(AssertionError("judge repeated")),
    )

    assert resumed == first


def test_matrix_resume_rejects_changed_case_contract(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    checkpoint = tmp_path / "matrix-checkpoint.jsonl"
    common = dict(
        subjects=(ModelSpec("gpt", "subject", "latest", "medium"),),
        judges=(ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {"text": _judge_json(rule, 3)},
        rating_rule=rule,
        checkpoint_path=checkpoint,
    )
    run_matrix(cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "original"},)},), **common)

    with pytest.raises(ValueError, match="checkpoint contract does not match"):
        run_matrix(
            cases=({"id": "TC-01", "turns": ({"turn": 1, "user": "changed"},)},),
            resume=True,
            **common,
        )


def test_matrix_retains_prompt_cache_usage() -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("gpt", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {
            "text": _judge_json(rule, 2),
            "usage": {"input_tokens": 100, "output_tokens": 5, "input_tokens_details": {"cached_tokens": 80, "cache_write_tokens": 10}},
        },
        rating_rule=rule,
    )

    assert rows[0]["judgement"]["cached_input_tokens"] == 80
    assert rows[0]["judgement"]["cache_write_tokens"] == 10
    assert rows[0]["judgement"]["cache_hit_ratio"] == pytest.approx(0.8)


def test_matrix_normalizes_anthropic_prompt_cache_usage() -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (ModelSpec("claude", "judge", "judge", "medium"),),
        subject_transport=lambda *_args: {"text": "answer"},
        judge_transport=lambda *_args: {
            "text": _judge_json(rule, 2),
            "usage": {
                "input_tokens": 100,
                "output_tokens": 5,
                "cache_read_input_tokens": 70,
                "cache_creation_input_tokens": 20,
            },
        },
        rating_rule=rule,
    )

    assert rows[0]["judgement"]["cached_input_tokens"] == 70
    assert rows[0]["judgement"]["cache_write_tokens"] == 20
    assert rows[0]["judgement"]["cache_hit_ratio"] == pytest.approx(0.7)


def test_workbook_stores_each_answer_once_and_links_judgements(tmp_path: Path) -> None:
    rule = load_rating_rule("ratings rule.yml")
    rows = run_matrix(
        ({"id": "TC-01", "turns": ({"turn": 1, "user": "hello"},)},),
        (ModelSpec("gpt", "subject", "latest", "medium"),),
        (
            ModelSpec("gpt", "judge-a", "judge", "medium"),
            ModelSpec("claude", "judge-b", "judge", "medium"),
        ),
        subject_transport=lambda _spec, _prompt: {
            "text": "answer",
            "chatflow_debug": _trace_with_large_duplicate_context(),
        },
        judge_transport=lambda _spec, _prompt: {
            "text": _judge_json(rule, 3)
        },
        rating_rule=rule,
    )
    path = tmp_path / "results.xlsx"
    write_matrix_workbook(rows, path)

    workbook = load_workbook(path, read_only=True)
    answers = list(workbook["All_Answers"].iter_rows(values_only=True))
    judgements = list(workbook["All_Judgements"].iter_rows(values_only=True))
    assert len(answers) == 2
    assert len(judgements) == 3
    assert judgements[1][0] == answers[1][0]
    assert judgements[2][0] == answers[1][0]

    benchmark = compare(path, path, 100.0, 50.0)
    assert benchmark["accepted"] is True
    assert benchmark["semantic_equivalence"] is True
    assert benchmark["speedup"] == 2.0
