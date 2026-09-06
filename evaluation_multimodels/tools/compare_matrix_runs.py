#!/usr/bin/env python3
"""Compare serial and parallel matrix workbooks without rerunning providers."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook


def _rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"unsupported matrix workbook (missing {sheet_name}): {path}")
    values = workbook[sheet_name].iter_rows(values_only=True)
    headers = tuple(str(value) for value in next(values))
    return [dict(zip(headers, row)) for row in values]


def _p95(values: Sequence[float]) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    return ordered[max(0, int(len(ordered) * 0.95 + 0.999999) - 1)]


def _snapshot(path: Path) -> dict[str, Any]:
    answers = _rows(path, "All_Answers")
    judgements = _rows(path, "All_Judgements")
    dimensions = sorted(key for key in judgements[0] if key.startswith("dimension:")) if judgements else []
    answer_facts = {
        str(row["answer_id"]): (row.get("status"), row.get("answer"), row.get("error_type")) for row in answers
    }
    judgement_facts = {
        (str(row["answer_id"]), str(row["judge_id"])):
        (row.get("status"), row.get("error_type"), row.get("weighted_score"), *(row.get(name) for name in dimensions))
        for row in judgements
    }
    queue_values = [
        float(value) for row in (*answers, *judgements)
        for value in (row.get("answer_queue_ms"), row.get("judge_queue_ms"))
        if isinstance(value, (int, float))
    ]
    return {
        "answer_facts": answer_facts,
        "judgement_facts": judgement_facts,
        "answers": len(answers),
        "judgements": len(judgements),
        "unavailable": sum(row.get("status") != "PASS" for row in judgements),
        "rate_limits": sum(row.get("error_type") == "RATE_LIMIT" for row in (*answers, *judgements)),
        "provider_5xx": sum(row.get("error_type") == "PROVIDER_5XX" for row in (*answers, *judgements)),
        "attempts": sum(int(row.get("judge_attempt_count") or 1) for row in judgements),
        "cached_input_tokens": sum(int(row.get("judge_cached_input_tokens") or 0) for row in judgements),
        "cache_write_tokens": sum(int(row.get("judge_cache_write_tokens") or 0) for row in judgements),
        "queue_p95_ms": _p95(queue_values),
    }


def compare(serial: Path, parallel: Path, serial_seconds: float, parallel_seconds: float) -> Mapping[str, Any]:
    before, after = _snapshot(serial), _snapshot(parallel)
    equivalent = before["answer_facts"] == after["answer_facts"] and before["judgement_facts"] == after["judgement_facts"]
    accepted = equivalent and after["rate_limits"] <= before["rate_limits"] and after["provider_5xx"] <= before["provider_5xx"]
    public_keys = ("answers", "judgements", "unavailable", "rate_limits", "provider_5xx", "attempts", "cached_input_tokens", "cache_write_tokens", "queue_p95_ms")
    return {
        "accepted": accepted,
        "semantic_equivalence": equivalent,
        "speedup": serial_seconds / parallel_seconds if parallel_seconds > 0 else None,
        "serial_seconds": serial_seconds,
        "parallel_seconds": parallel_seconds,
        "serial": {key: before[key] for key in public_keys},
        "parallel": {key: after[key] for key in public_keys},
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("serial", type=Path)
    parser.add_argument("parallel", type=Path)
    parser.add_argument("--serial-seconds", type=float, required=True)
    parser.add_argument("--parallel-seconds", type=float, required=True)
    args = parser.parse_args(argv)
    result = compare(args.serial, args.parallel, args.serial_seconds, args.parallel_seconds)
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
    return 0 if result["accepted"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
