"""Offline scope, loader, and reviewer-artifact checks. Does not approve content."""
from pathlib import Path
import hashlib,json,re,subprocess,sys
import yaml
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
manifest=json.loads((HERE/'review-manifest.json').read_text())
receipt=json.loads((HERE/'approval-receipt.json').read_text())
assert manifest['status']=='REVIEWED' and receipt['authorization']=='全部審核通過'
for entry in receipt['changes']:
    assert hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest()==entry['after_sha256'],entry['path']
counts={}
for project in ('evaluation','evaluation_multimodels'):
    # Separate process ensures each project's own imports are exercised.
    code='''from xiaoan_eval.cases import load_cases
from xiaoan_eval.rules import load_rating_rule
import json
rule=load_rating_rule("ratings rule.yml")
cases=[r.case for r in load_cases("test-cases/response-oracle-drafts",rule,pii_validator=lambda _:True)]
assert len(cases)==81 and all(cases)
assert sum(len(c.turns) for c in cases)==231
assert all(c.oracle_provenance.status=="reviewed" and c.maturity=="REVIEWED" and c.oracle_gate_eligible and not c.aggregate_eligible for c in cases)
assert all(c.oracle_provenance.reviewed_by=="mat" and c.oracle_provenance.reviewed_at=="2026-09-13" for c in cases)
assert all(t.expected.response_oracle and len(t.expected.response_oracle.required_claims)==2 and len(t.expected.response_oracle.forbidden_claims)>=1 for c in cases for t in c.turns)
normal=[r.case for r in load_cases("test-cases",rule,pii_validator=lambda _:True)]
assert len(normal)==74 and sum(len(c.turns) for c in normal)==217
assert all(c.oracle_gate_eligible and not c.aggregate_eligible for c in normal)
assert sum(t.expected.response_oracle is not None for c in normal for t in c.turns)==217
print(json.dumps({"draft_cases":81,"draft_turns":231,"default_cases":74,"default_turns":217,"draft_gate_eligible":81,"draft_aggregate_eligible":0}))
'''
    proc=subprocess.run([sys.executable,'-c',code],cwd=ROOT/project,text=True,capture_output=True,check=True)
    counts[project]=json.loads(proc.stdout)
    for n in range(1,82):
        p=Path('test-cases')/('proposed' if n>74 else '')/f'TC-{n:02}.yaml'
        original=yaml.safe_load((ROOT/project/p).read_text())
        draft=yaml.safe_load((ROOT/project/'test-cases/response-oracle-drafts'/p.name).read_text())
        assert len(original['turns'])==len(draft['turns'])
        for orig,new in zip(original['turns'],draft['turns']):
            orig.setdefault('expected',{}).pop('response_oracle',None)
            new['expected'].pop('response_oracle')
        for key in ('oracle_provenance','maturity'):
            original.pop(key,None);draft.pop(key,None)
        assert original==draft,(project,p)
for n in range(1,82):
    p=Path('test-cases/response-oracle-drafts')/f'TC-{n:02}.yaml'
    assert (ROOT/'evaluation'/p).read_bytes()==(ROOT/'evaluation_multimodels'/p).read_bytes()
wb=load_workbook(HERE/'response-oracle-review.xlsx');ws=wb['01_逐輪審核']
assert ws.max_row==232 and wb['02_案例概覽'].max_row==82
assert ws.freeze_panes=='E2' and len(ws.data_validations.dataValidation)==1
seen=set(); required=forbidden=0
for row in list(ws.values)[1:]:
    case,turn=row[:2];assert (case,turn) not in seen;seen.add((case,turn))
    t=yaml.safe_load((ROOT/'evaluation/test-cases/response-oracle-drafts'/f'{case}.yaml').read_text())['turns'][turn-1]
    assert t['user']==row[3]
    for key,col in [('required_claims',4),('forbidden_claims',5)]:
        assert [x.split('. ',1)[1] for x in row[col].splitlines()]==t['expected']['response_oracle'][key],(case,turn,key)
    required+=len(t['expected']['response_oracle']['required_claims']);forbidden+=len(t['expected']['response_oracle']['forbidden_claims'])
    assert row[8]=='通過' and all(x is None for x in row[9:12]) and row[12]=='mat' and row[13]=='2026-09-13'
    p=ROOT/'evaluation/test-cases'/('proposed' if int(case[3:])>74 else '')/f'{case}.yaml'
    current=yaml.safe_load(p.read_text())['turns'][turn-1]['expected']['response_oracle']
    assert current==t['expected']['response_oracle']
for sheet in wb:
    assert not any(c.data_type=='f' for row in sheet for c in row)
# validation.json is this check's output and is created below.
for p in [*HERE.glob('*.md'),*(ROOT/project/'test-cases/response-oracle-drafts/README.md' for project in counts)]:
    text=p.read_text()
    for link in re.findall(r'\]\(([^)]+)\)',text):
        if link.startswith(('http','#')):continue
        target=(p.parent/link.split('#')[0]).resolve()
        if target==HERE/'validation.json':continue
        assert target.exists(),(p,link)
result={'status':'PASS','validation_date':'2026-09-13','scope':'User approval recorded; offline structure and artifact checks; no live Judge validation',
        'projects':counts,'approved_source_files_verified':len(receipt['changes']),'mirrored_yaml_pairs':81,
        'only_oracle_and_draft_provenance_changed':True,'workbook_rows':231,'workbook_yaml_exact_match':True,
        'all_review_decisions':'APPROVED_BY_USER','required_claims':required,'forbidden_claims':forbidden,
        'markdown_file_links':'PASS','paid_model_calls':0}
(HERE/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
