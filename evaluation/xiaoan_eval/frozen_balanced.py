"""Assemble an equal-coverage matrix from independently sealed subject cohorts."""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from xiaoan_eval_core.contracts import digest
from xiaoan_eval_core.costs import build_answer_costs
from xiaoan_eval_core.results import aggregate_complete_results, seal_complete_results, validate_complete_results
from .frozen_cli import write_json
from .frozen_export import export_results_workbook


def balanced_comparison(result):
    cells = result['aggregates']['case_metrics']
    subjects = [item['id'] for item in result['plan']['subjects']]
    judges = [item['id'] for item in result['plan']['judges']]
    rows = []
    for metric in ('rubric', 'faithfulness', 'correctness', 'rubric_gate', 'requirements_gate'):
        matching = [item for item in cells if item['metric'] == metric and item['eligible']]
        by_pair = {(subject, judge): {item['case_id']: item for item in matching
                                      if item['subject_id'] == subject and item['judge_id'] == judge}
                   for subject in subjects for judge in judges}
        common = sorted(set.intersection(*(set(items) for items in by_pair.values())))
        for subject in subjects:
            for judge in judges:
                entries = [by_pair[(subject, judge)][case] for case in common]
                if metric.endswith('_gate'):
                    denominator = sum(item['gate'] != 'NOT_APPLICABLE' for item in entries)
                    value = sum(item['gate'] == 'PASS' for item in entries) / denominator if denominator else None
                else:
                    denominator = len(entries)
                    value = sum(item['value'] for item in entries) / denominator if denominator else None
                rows.append({'metric': metric, 'subject_id': subject, 'judge_id': judge,
                             'effective_cases': len(common), 'denominator': denominator,
                             'value': value, 'case_ids': common})
    return rows


def _append_comparison(path, rows):
    workbook = load_workbook(path)
    sheet = workbook.create_sheet('Balanced Comparison', 1)
    sheet.append(['指標', 'Subject', 'Judge', '共同案例數', '可計算分母', '分數／比例', '共同案例'])
    for item in rows:
        sheet.append([item['metric'], item['subject_id'], item['judge_id'], item['effective_cases'],
                      item['denominator'], item['value'], ', '.join(item['case_ids'])])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    for column, width in {'A': 22, 'B': 26, 'C': 26, 'D': 16, 'E': 16, 'F': 16, 'G': 62}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def assemble_balanced(results, paths, output, *, require_relevancy=True, require_assessment=True,
                      require_rubric=True):
    """Keep only whole cases complete across every subject and every Judge.

    Source results stay immutable. Original answer IDs and row manifest bindings
    are retained; the new manifest describes a derived, multi-manifest result.
    """
    if len(results) < 2 or len(results) != len(paths):
        raise ValueError('Provide at least two result cohorts with matching source paths')
    sources = [validate_complete_results(result) for result in results]
    first = sources[0]
    judges = first['plan']['judges']
    judge_ids = [judge['id'] for judge in judges]
    if len(set(judge_ids)) != 4:
        raise ValueError('Balanced matrix requires exactly four distinct Judges')
    subjects = {}
    covered_cases = {}
    for source in sources:
        plan = source['plan']
        for field in ('suite_id', 'subject_mode', 'extractor', 'relevancy_generator'):
            if plan.get(field) != first['plan'].get(field):
                raise ValueError(f'Incompatible source {field}')
        if plan['judges'] != judges or source['evaluation_config'] != first['evaluation_config']:
            raise ValueError('Judge identities or evaluator assets differ across cohorts')
        for subject in plan['subjects']:
            sid = subject['id']
            if sid in subjects and subjects[sid] != subject:
                raise ValueError(f'Subject identity changed across shards: {sid}')
            subjects[sid] = subject
            overlap = covered_cases.setdefault(sid, set()) & set(plan['case_ids'])
            if overlap:
                raise ValueError(f'Overlapping subject case shards: {sid} {sorted(overlap)}')
            covered_cases[sid].update(plan['case_ids'])
    subject_ids = list(subjects)
    if len(subject_ids) != 8:
        raise ValueError('Balanced matrix requires exactly eight distinct subjects')
    common_cases = set.intersection(*covered_cases.values())
    order = [case for case in first['plan']['case_ids'] if case in common_cases]
    answers = [row for source in sources for row in source['answers'] if row['case_id'] in common_cases]
    envelopes = {env['answer_id']: env for source in sources for env in source['envelopes']}
    inventories = {inv['inventory_id']: inv for source in sources for inv in source['inventories']}
    by_case = {}
    for answer in answers:
        by_case.setdefault(answer['case_id'], {}).setdefault(answer['subject_id'], []).append(answer)
    selected = []
    for case in order:
        turns = first['plan']['planned_turns'][case]
        if any(source['plan']['planned_turns'][case] != turns for source in sources
               if case in source['plan']['case_ids']):
            raise ValueError(f'Case turn plan differs: {case}')
        rows = by_case.get(case, {})
        if set(rows) != set(subject_ids):
            continue
        baseline = None
        complete = True
        for subject in subject_ids:
            lane = sorted(rows[subject], key=lambda row: row['turn'])
            if [row['turn'] for row in lane] != turns:
                complete = False
                break
            signature = [(row['turn'], row['question'], row['oracle_source']) for row in lane]
            if baseline is not None and signature != baseline:
                raise ValueError(f'Case oracle/question differs between subjects: {case}')
            baseline = signature
            for row in lane:
                env = envelopes[row['answer_id']]
                if (row['status'] != 'AVAILABLE' or
                        require_assessment and env['inventory_id'] not in inventories or
                        require_relevancy and env['relevancy'].get('status') != 'AVAILABLE'):
                    complete = False
                    break
                for branch in ('rubric', 'assessments'):
                    cells = env[branch]
                    if (len(cells) != 4 or {cell['judge_id'] for cell in cells} != set(judge_ids)
                            or (branch == 'rubric' and require_rubric or branch == 'assessments' and require_assessment)
                            and any(cell['status'] != 'AVAILABLE' for cell in cells)):
                        complete = False
                        break
        if complete:
            selected.append(case)
    if not selected:
        raise ValueError('No case has eight available subjects and four planned Judge cells')

    selected_set = set(selected)
    kept_answers = [deepcopy(row) for row in answers if row['case_id'] in selected_set]
    kept_ids = {row['answer_id'] for row in kept_answers}
    kept_envelopes = [deepcopy(env) for source in sources for env in source['envelopes'] if env['answer_id'] in kept_ids]
    kept_inventories = [deepcopy(inv) for source in sources for inv in source['inventories'] if inv['answer_id'] in kept_ids]
    kept_stages = [deepcopy(stage) for source in sources for stage in source['stages'] if stage.get('answer_id') in kept_ids]
    counts = Counter((row['subject_id'], cell['judge_id']) for row in kept_answers
                     for cell in envelopes[row['answer_id']]['rubric'])
    if len(counts) != 32 or len(set(counts.values())) != 1:
        raise ValueError('Matrix cell counts are not equal')
    manifest = {'schema_version': 'balanced-assembly/v1', 'suite_id': first['plan']['suite_id'],
                'case_ids': selected, 'source_manifest_digests': [source['manifest']['manifest_digest'] for source in sources],
                'source_generations': [source['result_generation'] for source in sources]}
    manifest['manifest_digest'] = digest(manifest)
    plan = deepcopy(first['plan'])
    plan.update(case_ids=selected, subjects=deepcopy(list(subjects.values())), planned_subjects=subject_ids,
                planned_turns={case: first['plan']['planned_turns'][case] for case in selected},
                judges=deepcopy(judges))
    assembled = {'schema_version': first['schema_version'], 'contract': first['contract'],
                 'run_ref': None, 'manifest': manifest, 'evaluation_config': deepcopy(first['evaluation_config']),
                 'plan': plan, 'answers': kept_answers, 'inventories': kept_inventories,
                 'envelopes': kept_envelopes, 'stages': kept_stages, 'artifacts': [],
                 'provenance': [{'operation': 'assemble_balanced', 'source_results': [str(Path(p).resolve()) for p in paths],
                                 'source_generations': manifest['source_generations'],
                                 'case_selection': ('whole_case_all_stages_available' if require_assessment else
                                                    'whole_case_rubric_comparable' if require_rubric else 'whole_case_answer_balanced'),
                                 'require_relevancy': require_relevancy,
                                 'require_assessment': require_assessment,
                                 'require_rubric': require_rubric,
                                 'original_answer_manifest_bindings_preserved': True}]}
    assembled['aggregates'] = aggregate_complete_results(assembled)
    assembled['aggregates']['answer_costs'] = build_answer_costs(kept_answers)
    assembled['aggregates']['balanced_comparison'] = balanced_comparison(assembled)
    seal_complete_results(assembled)
    validate_complete_results(assembled)
    out = Path(output)
    if (out / 'results.json').exists():
        raise ValueError('Output already sealed; choose a new directory')
    write_json(out / 'results.json', assembled)
    export_results_workbook(assembled, out / 'results.xlsx')
    _append_comparison(out / 'results.xlsx', assembled['aggregates']['balanced_comparison'])
    return assembled
