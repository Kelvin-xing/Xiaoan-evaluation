"""Blinded XLSX transport for exception-driven human review."""

from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .report_model import build_model_from_rows, canonical_digest
from .rules import RatingRule
from .workbook import WorkbookFacts, preflight_ooxml, read_workbook


PACKET_SHEETS = ("00_Instructions", "01_Review_Queue", "02_Review_Items", "03_Evidence")
QUEUE_HEADERS = (
    "review_id", "case_id", "turn", "response_sha256", "trigger", "user_input",
    "assistant_response", "reviewer_id", "submitted_at", "confidence", "notes",
)
ITEM_HEADERS = (
    "review_id", "item_type", "item_id", "item_name", "guidance", "score",
    "triggered", "evidence_refs", "deduction_reason",
)
EVIDENCE_HEADERS = ("review_id", "evidence_ref", "kind", "content")
EDITABLE_QUEUE = {"reviewer_id", "submitted_at", "confidence", "notes"}
EDITABLE_ITEMS = {"score", "triggered", "evidence_refs", "deduction_reason"}


def export_review_workbook(results: Path, packet: Path, rule: RatingRule) -> Path:
    facts = read_workbook(_results_path(results))
    if facts.artifact_state != "PENDING_REVIEW":
        raise ValueError(f"human review export requires PENDING_REVIEW, got {facts.artifact_state}")
    if packet.exists():
        raise ValueError(f"review packet already exists; refusing to overwrite: {packet}")
    queue, items, evidence = _packet_rows(facts, rule)
    if not queue:
        raise ValueError("workbook contains no pending review item")
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_instructions(workbook.create_sheet(PACKET_SHEETS[0]), rule, len(queue))
    _write_rows(workbook.create_sheet(PACKET_SHEETS[1]), QUEUE_HEADERS, queue, editable=EDITABLE_QUEUE)
    _write_rows(workbook.create_sheet(PACKET_SHEETS[2]), ITEM_HEADERS, items, editable=EDITABLE_ITEMS)
    _write_rows(workbook.create_sheet(PACKET_SHEETS[3]), EVIDENCE_HEADERS, evidence, editable=set())
    _add_validations(workbook)
    packet.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(packet)
    os.chmod(packet, 0o600)
    trust = {
        "schema_version": "1.0", "base_generation_id": facts.generation_id,
        "base_core_digest": facts.core_digest, "base_lifecycle_digest": facts.lifecycle_digest,
        "queue": [_immutable(row, EDITABLE_QUEUE) for row in queue],
        "items": [_immutable(row, EDITABLE_ITEMS) for row in items],
        "evidence": evidence,
        "rating_rule_schema_version": rule.schema_version,
    }
    trust["trust_digest"] = canonical_digest(trust)
    trust_path = review_trust_path(packet)
    trust_path.parent.mkdir(parents=True, exist_ok=True)
    trust_path.write_text(json.dumps(trust, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n", encoding="utf-8")
    os.chmod(trust_path.parent, 0o700)
    os.chmod(trust_path, 0o600)
    return trust_path


def import_review_workbook(packet: Path, results: Path, rule: RatingRule):
    official_path = _results_path(results)
    facts = read_workbook(official_path)
    trust_path = review_trust_path(packet)
    trust = json.loads(trust_path.read_text(encoding="utf-8"))
    digest = trust.pop("trust_digest", None)
    if digest != canonical_digest(trust):
        raise ValueError("review trust record digest does not match")
    if trust.get("base_generation_id") != facts.generation_id or trust.get("base_core_digest") != facts.core_digest or trust.get("base_lifecycle_digest") != facts.lifecycle_digest:
        raise ValueError("review packet is stale for the current deliverable generation")
    if trust.get("rating_rule_schema_version") != rule.schema_version:
        raise ValueError("review packet rating rule does not match")
    queue, items, evidence = _read_packet(packet)
    _validate_immutable_rows(queue, trust["queue"], EDITABLE_QUEUE, "queue")
    _validate_immutable_rows(items, trust["items"], EDITABLE_ITEMS, "items")
    if evidence != trust["evidence"]:
        raise ValueError("review packet evidence was changed")
    submissions = _validate_editable(queue, items, evidence, rule)
    return _apply_submissions(facts, submissions, rule)


def review_trust_path(packet: Path) -> Path:
    return packet.parent / ".xiaoan-review-trust" / f"{packet.stem}.json"


def _packet_rows(facts: WorkbookFacts, rule: RatingRule):
    text = _reconstruct_text(facts.text_content)
    turn_by_key = {(str(row["case_id"]), int(row["turn"])): row for row in facts.turns}
    queue, items, evidence = [], [], []
    for pending in facts.human_review:
        if str(pending.get("status")) != "PENDING_REVIEW":
            continue
        case_id, turn = str(pending["case_id"]), int(pending["turn"])
        turn_row = turn_by_key[(case_id, turn)]
        review_id = "HR-" + canonical_digest([facts.generation_id, case_id, turn, turn_row.get("response_sha256")])[:20]
        queue.append({
            "review_id": review_id, "case_id": case_id, "turn": turn,
            "response_sha256": turn_row.get("response_sha256"), "trigger": pending.get("trigger") or "exception-triggered review",
            "user_input": text.get(str(turn_row.get("user_text_id")), ""),
            "assistant_response": text.get(str(turn_row.get("assistant_text_id")), ""),
            "reviewer_id": "", "submitted_at": "", "confidence": "", "notes": "",
        })
        for red_line in rule.red_lines:
            items.append({
                "review_id": review_id, "item_type": "red_line", "item_id": red_line.id,
                "item_name": red_line.name, "guidance": red_line.description,
                "score": None, "triggered": None, "evidence_refs": "", "deduction_reason": "",
            })
        for module in rule.modules:
            guidance = "Positive: " + " / ".join(module.positive) + "\nNegative: " + " / ".join(module.negative)
            items.append({
                "review_id": review_id, "item_type": "dimension", "item_id": module.name,
                "item_name": module.name, "guidance": guidance,
                "score": None, "triggered": None, "evidence_refs": "", "deduction_reason": "",
            })
        evidence.append({"review_id": review_id, "evidence_ref": f"{case_id}:T{turn}:user", "kind": "user_input", "content": queue[-1]["user_input"]})
        evidence.append({"review_id": review_id, "evidence_ref": f"{case_id}:T{turn}:assistant", "kind": "assistant_response", "content": queue[-1]["assistant_response"]})
        for ref in str(next((row.get("evidence_refs") for row in facts.cases if row.get("case_id") == case_id), "") or "").split(";"):
            if ref.strip():
                evidence.append({"review_id": review_id, "evidence_ref": ref.strip(), "kind": "evaluation_ref", "content": "See the bound case/turn evaluation evidence in results.xlsx."})
    return queue, items, evidence


def _write_instructions(sheet, rule: RatingRule, count: int) -> None:
    rows = [
        ("Purpose", f"Review {count} exception-triggered response(s) without seeing the automatic scores."),
        ("Workflow", "Read 01_Review_Queue and 03_Evidence, then complete only the blue cells in 01_Review_Queue and 02_Review_Items."),
        ("Identity", "reviewer_id is a self-declared operational label. This workbook does not authenticate the person who filled it."),
        ("Red lines", "Choose TRUE or FALSE for every red-line row and cite bound evidence refs when available."),
        ("Dimensions", "Choose an anchored integer score 0, 1, 2, or 3 for every dimension row."),
        ("Evidence", "Use semicolon-separated evidence_ref values from 03_Evidence. Do not add external or invented refs."),
        ("Deduction", "Explain why points were deducted whenever a dimension score is below 3."),
        ("Confidence", "Choose low, medium, or high. Enter an ISO 8601 timestamp with timezone."),
        ("Escalation", "If evidence is insufficient, set confidence=low, explain what is missing in notes, and do not guess."),
        ("Completion", "Every review item must be complete. Import is all-or-nothing and rejects changed immutable cells or formulas."),
    ]
    rows.extend((f"Score {anchor.score}", anchor.description) for anchor in rule.score_scale)
    sheet.append(["section", "instruction"])
    for row in rows:
        sheet.append(list(row))
    _style_sheet(sheet, {"instruction"}, set())


def _write_rows(sheet, headers, rows, *, editable):
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(header) for header in headers])
        for cell, value in zip(sheet[sheet.max_row], [row.get(header) for header in headers], strict=True):
            if isinstance(value, str):
                cell.data_type = "s"
    _style_sheet(sheet, set(headers) - editable, editable)


def _style_sheet(sheet, immutable, editable):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    headers = [cell.value for cell in sheet[1]]
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 42 if header in {"instruction", "guidance", "content", "user_input", "assistant_response", "notes"} else 18
        fill = PatternFill("solid", fgColor="D9EAF7") if header in editable else PatternFill("solid", fgColor="E7E6E6")
        for cell in list(sheet.columns)[index - 1][1:]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_validations(workbook):
    queue = workbook["01_Review_Queue"]
    confidence = DataValidation(type="list", formula1='"low,medium,high"', allow_blank=False)
    queue.add_data_validation(confidence)
    confidence.add(f"J2:J{queue.max_row}")
    items = workbook["02_Review_Items"]
    scores = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=False)
    triggered = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
    items.add_data_validation(scores); items.add_data_validation(triggered)
    scores.add(f"F2:F{items.max_row}"); triggered.add(f"G2:G{items.max_row}")


def _read_packet(packet: Path):
    preflight_ooxml(packet)
    workbook = load_workbook(packet, data_only=False, keep_links=False)
    if tuple(workbook.sheetnames) != PACKET_SHEETS:
        raise ValueError("review packet sheet contract does not match")
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            raise ValueError("review packet cannot contain hidden sheets")
        if any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row):
            raise ValueError("review packet formulas are not allowed")
    return (
        _rows(workbook["01_Review_Queue"], QUEUE_HEADERS),
        _rows(workbook["02_Review_Items"], ITEM_HEADERS),
        _rows(workbook["03_Evidence"], EVIDENCE_HEADERS),
    )


def _rows(sheet, headers):
    if tuple(cell.value for cell in sheet[1]) != tuple(headers):
        raise ValueError(f"unexpected columns in review sheet {sheet.title}")
    return [dict(zip(headers, values, strict=True)) for values in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in values)]


def _validate_immutable_rows(actual, expected, editable, label):
    actual_immutable = [_immutable(row, editable) for row in actual]
    if actual_immutable != expected:
        raise ValueError(f"review packet {label} immutable cells were changed")


def _immutable(row, editable):
    return {key: value for key, value in row.items() if key not in editable}


def _validate_editable(queue, items, evidence, rule):
    allowed_evidence: dict[str, set[str]] = {}
    for row in evidence:
        allowed_evidence.setdefault(str(row["review_id"]), set()).add(str(row["evidence_ref"]))
    by_review: dict[str, list[Mapping[str, Any]]] = {}
    for row in items:
        by_review.setdefault(str(row["review_id"]), []).append(row)
    submissions = []
    expected_dimensions = {module.name for module in rule.modules}
    expected_red_lines = {item.id for item in rule.red_lines}
    for row in queue:
        review_id = str(row["review_id"])
        reviewer_id = str(row.get("reviewer_id") or "").strip()
        submitted_at = str(row.get("submitted_at") or "").strip()
        confidence = str(row.get("confidence") or "").strip().lower()
        if not reviewer_id:
            raise ValueError(f"{review_id}: reviewer_id operational label is required")
        try:
            timestamp = datetime.fromisoformat(submitted_at)
        except ValueError as exc:
            raise ValueError(f"{review_id}: submitted_at must be ISO 8601") from exc
        if timestamp.tzinfo is None:
            raise ValueError(f"{review_id}: submitted_at must include timezone")
        if confidence not in {"low", "medium", "high"}:
            raise ValueError(f"{review_id}: confidence must be low, medium, or high")
        review_items = by_review.get(review_id, [])
        dimension_rows = {str(item["item_id"]): item for item in review_items if item["item_type"] == "dimension"}
        red_line_rows = {str(item["item_id"]): item for item in review_items if item["item_type"] == "red_line"}
        if set(dimension_rows) != expected_dimensions or set(red_line_rows) != expected_red_lines:
            raise ValueError(f"{review_id}: review item set is incomplete")
        dimensions, red_lines = [], []
        for module in rule.modules:
            item = dimension_rows[module.name]
            score = item.get("score")
            if isinstance(score, bool) or not isinstance(score, (int, float)) or not float(score).is_integer() or int(score) not in {0, 1, 2, 3}:
                raise ValueError(f"{review_id}: {module.name} score must be anchored 0, 1, 2, or 3")
            refs = _refs(item.get("evidence_refs"), allowed_evidence[review_id], review_id)
            reason = str(item.get("deduction_reason") or "").strip()
            if int(score) < 3 and not reason:
                raise ValueError(f"{review_id}: {module.name} requires a deduction reason below score 3")
            dimensions.append({"module": module.name, "score": int(score), "supporting_evidence_refs": refs, "deduction_reason": reason})
        for red_line in rule.red_lines:
            item = red_line_rows[red_line.id]
            triggered = item.get("triggered")
            if triggered not in {True, False}:
                raise ValueError(f"{review_id}: {red_line.id} triggered must be TRUE or FALSE")
            red_lines.append({"id": red_line.id, "triggered": bool(triggered), "evidence_refs": _refs(item.get("evidence_refs"), allowed_evidence[review_id], review_id)})
        submissions.append({
            "review_id": review_id, "case_id": str(row["case_id"]), "turn": int(row["turn"]),
            "response_sha256": str(row["response_sha256"]), "reviewer_id": reviewer_id,
            "submitted_at": submitted_at, "confidence": confidence, "notes": str(row.get("notes") or ""),
            "dimensions": dimensions, "red_lines": red_lines,
        })
    return submissions


def _refs(value, allowed, review_id):
    refs = tuple(item.strip() for item in str(value or "").split(";") if item.strip())
    unknown = set(refs) - allowed
    if unknown:
        raise ValueError(f"{review_id}: unknown evidence refs: {sorted(unknown)}")
    return refs


def _apply_submissions(facts: WorkbookFacts, submissions, rule: RatingRule):
    cases = [dict(row) for row in facts.cases]
    turns = [dict(row) for row in facts.turns]
    metrics = [dict(row) for row in facts.metrics]
    review_rows = [dict(row) for row in facts.human_review]
    disagreement_keys = set()
    affected_cases = {item["case_id"] for item in submissions}
    for submission in submissions:
        case_id, turn = submission["case_id"], submission["turn"]
        auto_red_lines = {
            str(row["metric_id"]).removeprefix("judge:red_line:"): bool(row.get("raw_score"))
            for row in metrics if row.get("case_id") == case_id and int(row.get("turn")) == turn and str(row.get("metric_id", "")).startswith("judge:red_line:") and row.get("score_source") == "automatic"
        }
        human_red_lines = {row["id"]: row["triggered"] for row in submission["red_lines"]}
        if auto_red_lines != human_red_lines:
            disagreement_keys.add((case_id, turn))
        turn_row = next(row for row in turns if row.get("case_id") == case_id and int(row.get("turn")) == turn)
        weights = _case_weights(next(row for row in cases if row["case_id"] == case_id), rule)
        human_score = sum(item["score"] * weights[item["module"]] for item in submission["dimensions"])
        turn_row.update({"human_score": human_score, "final_score": None if (case_id, turn) in disagreement_keys else human_score, "final_source": "human", "review_status": "NEEDS_ADJUDICATION" if (case_id, turn) in disagreement_keys else "COMPLETED"})
        for item in submission["dimensions"]:
            metrics.append({
                "row_key": f"subject:{case_id}:T{turn}:judge:{item['module']}:human", "comparison_key": canonical_digest([case_id, turn, item["module"], "human"]),
                "component_run_id": "subject", "case_id": case_id, "turn": turn,
                "metric_id": f"judge:{item['module']}", "dimension": item["module"], "raw_score": item["score"],
                "normalized_score": item["score"], "weight": weights[item["module"]], "contribution": item["score"] * weights[item["module"]],
                "status": "AVAILABLE", "reason": item["deduction_reason"], "evidence_refs": "; ".join(item["supporting_evidence_refs"]),
                "score_source": "human", "baseline_value": None, "delta": None,
            })
        for item in submission["red_lines"]:
            metric_id = f"judge:red_line:{item['id']}"
            metrics.append({
                "row_key": f"subject:{case_id}:T{turn}:{metric_id}:human", "comparison_key": canonical_digest([case_id, turn, metric_id, "human"]),
                "component_run_id": "subject", "case_id": case_id, "turn": turn,
                "metric_id": metric_id, "dimension": "safety_red_line", "raw_score": 1 if item["triggered"] else 0,
                "normalized_score": 1 if item["triggered"] else 0, "weight": None, "contribution": None,
                "status": "AVAILABLE", "reason": "human red-line judgment", "evidence_refs": "; ".join(item["evidence_refs"]),
                "score_source": "human", "baseline_value": None, "delta": None,
            })
        review_row = next(row for row in review_rows if row.get("case_id") == case_id and int(row.get("turn")) == turn)
        review_row.update({"review_id": submission["review_id"], "status": "NEEDS_ADJUDICATION" if (case_id, turn) in disagreement_keys else "COMPLETED", "reviewer_id": submission["reviewer_id"], "submitted_at": submission["submitted_at"], "confidence": submission["confidence"], "notes": submission["notes"]})
    for case in cases:
        if case["case_id"] not in affected_cases:
            continue
        case_turns = [row for row in turns if row.get("case_id") == case["case_id"]]
        if any((case["case_id"], int(row["turn"])) in disagreement_keys for row in case_turns):
            case.update({"artifact_state": "NEEDS_ADJUDICATION", "human_score": None, "final_score": None, "final_source": "unresolved", "review_status": "NEEDS_ADJUDICATION", "status": "NEEDS_ADJUDICATION"})
            continue
        values = [float(row["final_score"]) for row in case_turns if isinstance(row.get("final_score"), (int, float)) and not isinstance(row.get("final_score"), bool)]
        human_value = sum(values) / len(values) if values else None
        deterministic = [str(row.get("status", "")).upper() for row in metrics if row.get("case_id") == case["case_id"] and row.get("score_source") == "automatic" and row.get("metric_id") != "route_preference"
            and not str(row.get("metric_id", "")).startswith("judge:")]
        status = "ERROR" if "ERROR" in deterministic else "FAIL" if "FAIL" in deterministic else "PASS"
        case.update({"artifact_state": "FINAL", "human_score": human_value, "final_score": human_value, "final_source": "human", "review_status": "COMPLETED", "status": status})
    state = ("NEEDS_ADJUDICATION" if any(row.get("review_status") == "NEEDS_ADJUDICATION" for row in turns)
        else "PENDING_REVIEW" if any(row.get("review_status") == "NEEDS_REVIEW" for row in turns) else "FINAL")
    _refresh_reviewed_cases([case for case in cases if case["case_id"] in affected_cases], turns, metrics)
    return build_model_from_rows(
        manifest=facts.manifest, artifact_state=state, cases=cases, turns=turns, metrics=metrics,
        baseline=facts.baseline, experiments=facts.experiments, human_review=review_rows,
        stability=facts.stability, text_content=facts.text_content,
    )


def adjudicate_workbook(
    results: Path,
    *,
    case_id: str,
    turn: int,
    decision_source: str,
    adjudicator_id: str,
    rationale: str,
    rule: RatingRule,
):
    """Resolve one disputed turn using its immutable automatic or human facts."""
    facts = read_workbook(_results_path(results))
    if facts.artifact_state != "NEEDS_ADJUDICATION":
        raise ValueError(f"adjudication requires NEEDS_ADJUDICATION, got {facts.artifact_state}")
    if decision_source not in {"automatic", "human"}:
        raise ValueError("adjudication decision must be automatic or human")
    if not adjudicator_id.strip() or not rationale.strip():
        raise ValueError("adjudicator operational label and rationale are required")
    cases = [dict(row) for row in facts.cases]
    turns = [dict(row) for row in facts.turns]
    metrics = [dict(row) for row in facts.metrics]
    reviews = [dict(row) for row in facts.human_review]
    turn_row = next((row for row in turns if row.get("case_id") == case_id and int(row.get("turn")) == turn), None)
    if turn_row is None or turn_row.get("review_status") != "NEEDS_ADJUDICATION":
        raise ValueError(f"no disputed review exists for {case_id} turn {turn}")
    source_dimensions = [
        row for row in metrics
        if row.get("case_id") == case_id and int(row.get("turn")) == turn
        and row.get("score_source") == decision_source
        and str(row.get("metric_id", "")).startswith("judge:")
        and ":red_line:" not in str(row.get("metric_id", ""))
    ]
    source_red_lines = [
        row for row in metrics
        if row.get("case_id") == case_id and int(row.get("turn")) == turn
        and row.get("score_source") == decision_source
        and str(row.get("metric_id", "")).startswith("judge:red_line:")
    ]
    if {str(row.get("dimension")) for row in source_dimensions} != {module.name for module in rule.modules}:
        raise ValueError("adjudication source has incomplete dimension facts")
    triggered = any(row.get("raw_score") == 1 for row in source_red_lines)
    score = 0.0 if triggered else sum(float(row["raw_score"]) * _case_weights(next(case for case in cases if case["case_id"] == case_id), rule)[row["dimension"]] for row in source_dimensions)
    turn_row.update({"final_score": score, "final_source": "adjudicated", "review_status": "COMPLETED"})
    for row in source_dimensions + source_red_lines:
        copy = dict(row)
        copy.update({
            "row_key": str(copy["row_key"]).rsplit(":", 1)[0] + ":adjudicated",
            "comparison_key": canonical_digest([case_id, turn, copy["metric_id"], "adjudicated"]),
            "score_source": "adjudicated",
            "reason": f"{rationale} (selected {decision_source})",
        })
        metrics.append(copy)
    review_row = next(row for row in reviews if row.get("case_id") == case_id and int(row.get("turn")) == turn)
    review_row.update({"status": "ADJUDICATED", "reviewer_id": f"{review_row.get('reviewer_id', '')}; adjudicator={adjudicator_id}", "notes": f"{review_row.get('notes', '')}\nAdjudication: {rationale}; source={decision_source}".strip()})
    for case in cases:
        if case.get("case_id") != case_id:
            continue
        case_turns = [row for row in turns if row.get("case_id") == case_id]
        unresolved = any(row.get("review_status") == "NEEDS_ADJUDICATION" for row in case_turns)
        if not unresolved:
            values = [float(row["final_score"]) for row in case_turns if isinstance(row.get("final_score"), (int, float)) and not isinstance(row.get("final_score"), bool)]
            case.update({"artifact_state": "FINAL", "final_score": sum(values) / len(values) if values else None, "final_source": "adjudicated", "review_status": "COMPLETED", "adjudication_status": "COMPLETED", "status": "FAIL" if triggered else "PASS"})
    state = "NEEDS_ADJUDICATION" if any(row.get("review_status") == "NEEDS_ADJUDICATION" for row in turns) else "FINAL"
    _refresh_reviewed_cases([case for case in cases if case["case_id"] == case_id], turns, metrics)
    return build_model_from_rows(
        manifest=facts.manifest, artifact_state=state, cases=cases, turns=turns,
        metrics=metrics, baseline=facts.baseline, experiments=facts.experiments,
        human_review=reviews, stability=facts.stability, text_content=facts.text_content,
    )


def _reconstruct_text(rows: Sequence[Mapping[str, Any]]):
    grouped: dict[str, list[Mapping[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(str(row["text_id"]), []).append(row)
    result = {}
    for text_id, chunks in grouped.items():
        ordered = sorted(chunks, key=lambda row: int(row["chunk_index"]))
        expected_count = int(ordered[0]["chunk_count"])
        if len(ordered) != expected_count or [int(row["chunk_index"]) for row in ordered] != list(range(1, expected_count + 1)):
            raise ValueError(f"text chunks are incomplete or out of order: {text_id}")
        hashes = {str(row["text_sha256"]) for row in ordered}
        if len(hashes) != 1:
            raise ValueError(f"text chunks disagree on content hash: {text_id}")
        content = "".join(str(row.get("content") or "") for row in ordered)
        if hashlib.sha256(content.encode("utf-8")).hexdigest() not in hashes:
            raise ValueError(f"text chunk content hash does not match: {text_id}")
        result[text_id] = content
    return result


def _results_path(path: Path) -> Path:
    return path / "results.xlsx" if path.is_dir() else path


def _case_weights(case, rule):
    raw = case.get("quality_weights")
    weights = json.loads(raw) if isinstance(raw, str) and raw else raw
    if not weights:
        raise ValueError("workbook lacks dynamic quality weights; regenerate it from the original case/rule before human scoring")
    if not isinstance(weights, Mapping) or set(weights) != {module.name for module in rule.modules}:
        raise ValueError("workbook quality weights do not match rating rule")
    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 1 for value in weights.values()) or abs(sum(weights.values()) - 1) > 1e-9:
        raise ValueError("workbook quality weights must be normalized")
    return weights


def _refresh_reviewed_cases(cases, turns, metrics):
    """Preserve execution failure and select red lines across the whole episode."""
    for case in cases:
        if case.get("review_status") == "NEEDS_ADJUDICATION":
            continue
        case_turns = [row for row in turns if row.get("case_id") == case["case_id"]]
        triggered = False
        for turn in case_turns:
            source = turn.get("final_source", "automatic")
            triggered |= any(row.get("raw_score") == 1 for row in metrics
                if row.get("case_id") == case["case_id"] and row.get("turn") == turn.get("turn")
                and row.get("score_source") == source and str(row.get("metric_id", "")).startswith("judge:red_line:"))
        deterministic = [row for row in metrics if row.get("case_id") == case["case_id"]
            and row.get("score_source") == "automatic" and row.get("metric_id") != "route_preference"
            and not str(row.get("metric_id", "")).startswith("judge:")]
        error = case.get("execution_status") == "ERROR" or any(str(row.get("status", "")).upper() == "ERROR" for row in deterministic)
        failed = any(str(row.get("status", "")).upper() == "FAIL" for row in deterministic)
        complete = bool(case_turns) and all(isinstance(row.get("final_score"), (int, float)) for row in case_turns)
        case["hard_gate_passed"] = False if triggered or failed else None if error else True
        case["execution_status"] = "ERROR" if error else "FAIL" if failed else "PASS"
        case["quality_status"] = "AVAILABLE" if complete and not error else "UNAVAILABLE"
        if not complete or error:
            case["final_score"] = None
        elif triggered:
            case["final_score"] = 0.0
        threshold = case.get("quality_threshold")
        below = isinstance(threshold, (int, float)) and case.get("final_score") is not None and case["final_score"] < threshold
        case["status"] = "ERROR" if error else "FAIL" if failed or triggered or below else "PASS"
