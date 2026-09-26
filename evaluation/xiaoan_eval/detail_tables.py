"""Readable projections of frozen scenario facts; no model calls or score changes."""
from collections import defaultdict
import json
from math import ceil

NAMES = ('Claims', 'Evidence', 'Relevancy', 'Requirements', 'Groups')
ID = ('subject_id', 'judge_id', 'case_id', 'turn')
TURN_COLUMNS = ('subject_id','judge_id','annotation_status','task_family','task','topic','constraints','dialogue','risk','need_summary','answer_status','context_status','trace_checks','primary_faithfulness','strict_faithfulness','answer_relevancy','relevancy_status','required_coverage','diagnostic_next_check')
HEADERS = {
 'Claims': (*ID,'claim_id','kind','claim_text','support','unsupported_category','evidence_layers','demand_relevance','status','part','parts'),
 'Evidence': (*ID,'claim_id','evidence_number','claim_text','relation','layer','evidence_ref','evidence_text','start','end','occurrence_id','status','part','parts'),
 'Relevancy': (*ID,'question_number','original_question','reverse_question','similarity','answer_mean','n','query_mode','generator_model','embedding_model','embedding_revision','status','reason','part','parts'),
 'Requirements': (*ID,'requirement_id','kind','requirement_text','verdict','reason','answer_quote','oracle_status','status','part','parts'),
 'Groups': ('subject_id','judge_id','axis','label','n_cases','n_turns','answered_n','annotated_n','metric','numerator','denominator','mean','median','available_n','unavailable_n','interpretation','part','parts'),
}

def chunks(row):
    """Split long fields on the same row key; UTF-16-safe Excel cell sizes."""
    fields = {}
    for key, value in row.items():
        if isinstance(value, str):
            # 14000 Unicode characters <= 28000 UTF-16 units, below Excel limit.
            fields[key] = [value[i:i+14000] for i in range(0,len(value),14000)] or ['']
    count = max([len(x) for x in fields.values()] or [1])
    for i in range(count):
        yield {**row, **{k: (v[i] if i<len(v) else None) if len(v)>1 else v[0] for k,v in fields.items()}, 'part':i+1,'parts':count}


def analysis_from_metrics(metrics):
    result = {s: [] for s in ('turns','claims','groups','answer_groups','claim_groups')}
    pieces=defaultdict(list); order=[]
    for row in metrics:
        name=str(row.get('metric_id',''))
        if name.startswith('scenario.') and name[9:] in result:
            order.append((name[9:], json.loads(row['reason'])))
        elif name.startswith('scenario_chunk.'):
            key=(name[15:],row['comparison_key'])
            if key not in pieces: order.append((key[0], key))
            pieces[key].append(row)
    for section,value in order:
        if section in result:
            if isinstance(value,tuple):
                value=json.loads(''.join(r['reason'] for r in sorted(pieces[value],key=lambda x:int(x['raw_score']))))
            result[section].append(value)
    return result


def project(analysis):
    tables={n:[] for n in NAMES};turns=[];seen_answers=set()
    for t in analysis.get('turns',[]):
        ident={k:t.get(k) for k in ID}; ar=t.get('answer_relevancy',{});cov=t.get('required_coverage',{})
        turns.append({**ident, **{k:'；'.join(v) for k,v in t.get('tags',{}).items()},
          'annotation_status':t.get('annotation_status'),'need_summary':t.get('need_summary'),
          'answer_status':t.get('answer_status'),'context_status':t.get('context_status'),
          'trace_checks':'\n'.join(str(k)+': '+str(v.get('status','UNKNOWN'))+' '+str(v.get('reason','')) for k,v in t.get('trace_checks',{}).items() if isinstance(v,dict)),
          'primary_faithfulness':t.get('primary_binary',{}).get('score'),
          'strict_faithfulness':t.get('strict_entailment_proxy',{}).get('score'),
          'answer_relevancy':ar.get('score'),'relevancy_status':ar.get('status'),
          'required_coverage':cov.get('score'),
          'diagnostic_next_check':'\n'.join(h.get('stage','')+': '+h.get('next_check','') for h in t.get('hypotheses',[]))})
        answer_key=(t.get('subject_id'),t.get('case_id'),t.get('turn'),t.get('binding',{}).get('answer_sha256'),t.get('binding',{}).get('context_sha256'))
        questions = [] if answer_key in seen_answers else (ar.get('questions') or [None])
        seen_answers.add(answer_key)
        for i,q in enumerate(questions):
            tables['Relevancy'].append({**ident,'judge_id':'ANSWER_LEVEL','question_number':i+1 if q else None,
              'original_question':t.get('user'),'reverse_question':q,
              'similarity':ar.get('similarities',[None])[i] if q else None,
              'answer_mean':ar.get('score'),**{k:ar.get(k) for k in ('n','query_mode','generator_model','embedding_model','embedding_revision','status','reason')}})
        items=t.get('oracle_items',cov.get('items',[]))
        expected=t.get('observed_expected',{}).get('response_oracle',{}) or {}
        for item in items or [{'status':'UNAVAILABLE'}]:
            kind=item.get('kind'); idx='required_claims' if kind=='required' else 'forbidden_claims'
            values=expected.get(idx,[]) or []; rid=item.get('id',''); text=item.get('text') or item.get('claim')
            if not text and rid[1:].isdigit() and 0<int(rid[1:])<=len(values):text=values[int(rid[1:])-1]
            tables['Requirements'].append({**ident,'requirement_id':rid,'kind':kind,'requirement_text':text,
               'verdict':item.get('verdict'),'reason':item.get('reason'),
               'answer_quote':'\n'.join(s.get('quote',s.get('text','')) for s in item.get('spans',[])),
               'oracle_status':t.get('oracle_status','UNKNOWN'),'status':item.get('status','AVAILABLE')})
    for c in analysis.get('claims',[]):
        ident={k:c.get(k) for k in ID};base={**ident,'claim_id':c.get('claim_id'),'claim_text':c.get('text')}
        tables['Claims'].append({**base,'kind':c.get('kind'),'support':c.get('support_category'),
           'unsupported_category':c.get('unsupported_category'),'evidence_layers':'；'.join(c.get('layers',[])),
           'demand_relevance':c.get('demand_relevance'),'status':'AVAILABLE'})
        for i,r in enumerate(c.get('relations',[]),1):
            span=r.get('evidence_span') or {}
            tables['Evidence'].append({**base,'evidence_number':i,'relation':r.get('relation'),
                **{k:r.get(k) for k in ('layer','evidence_ref','occurrence_id')},
                'evidence_text':span.get('text'),'start':span.get('start'),'end':span.get('end'),'status':'AVAILABLE'})
    # Missing claim inventories remain explicit at turn level, never invented claims.
    for g in analysis.get('groups',[]):
        for metric in ('primary_binary','strict_entailment_proxy','required_coverage'):
            v=g.get(metric,{})
            tables['Groups'].append({**{k:g.get(k) for k in HEADERS['Groups']},'metric':metric,
                **{k:v.get(k) for k in ('numerator','denominator','median','available_n','unavailable_n')},
                'mean':v.get('turn_macro_mean'),'interpretation':'多標籤重疊；非因果結論；必要需求分母僅已判定項'})
    for g in analysis.get('answer_groups',[]):
        v=g.get('answer_relevancy',{})
        tables['Groups'].append({**{k:g.get(k) for k in HEADERS['Groups']},'metric':'answer_relevancy',
            **{k:v.get(k) for k in ('median','available_n','unavailable_n')},'mean':v.get('turn_macro_mean'),
            'interpretation':'每個回答只計一次；平均餘弦不是正確率'})
    return turns,{name:[chunk for row in rows for chunk in chunks(row)] for name,rows in tables.items()}


def write_table(book,name,headers,rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    if name in book: del book[name]
    sheet=book.create_sheet(name);sheet.append(list(headers))
    for row in rows:
        sheet.append([float(format(row[k], '.15g')) if isinstance(row.get(k),float) else row.get(k) for k in headers])
        for c in sheet[sheet.max_row]:
            if isinstance(c.value,str):c.data_type='s'
    for c in sheet[1]:
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78')
        c.comment=Comment('part/parts 為超長文字分段；相同識別鍵可依 part 重組。缺失不計品質零分。','XiaoAn')
    sheet.freeze_panes='E2';sheet.auto_filter.ref=sheet.dimensions
    for i,k in enumerate(headers,1):
        sheet.column_dimensions[get_column_letter(i)].width=60 if any(s in k for s in ('text','question','quote','reason','check')) else 22
    colors={'ENTAILS':'E2EFDA','PARTIAL_ONLY':'FFF2CC','PARTIAL':'FFF2CC','CONTRADICTS':'FCE4D6','UNAVAILABLE':'E7E6E6'}
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if c.value in colors if isinstance(c.value,str) else False:c.fill=PatternFill('solid',fgColor=colors[c.value])
    return sheet


def turn_projection(base_rows, analysis):
    summaries,_=project(analysis)
    index={(t.get('case_id'),t.get('turn')):t for t in summaries}
    seen=set();result=[]
    for row in base_rows:
        key=(row.get('case_id'),row.get('turn'));seen.add(key)
        result.append({**row,**{k:index.get(key,{}).get(k) for k in TURN_COLUMNS}})
    for key,t in index.items():
        if key not in seen:
            result.append({'row_kind':'diagnostic','case_id':key[0],'turn':key[1],**{k:t.get(k) for k in TURN_COLUMNS}})
    return result
