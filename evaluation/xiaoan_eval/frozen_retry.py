"""Selective retries from a sealed frozen result into a new immutable generation."""
from __future__ import annotations

from copy import deepcopy
from collections import Counter
from pathlib import Path
import hashlib
import json
import os
from contextlib import contextmanager
from uuid import uuid4

from xiaoan_eval_core.configuration import snapshot, workflow
from xiaoan_eval_core.contracts import digest
from xiaoan_eval_core.contracts import assessment_request, extraction_request, validate_inventory
from xiaoan_eval_core.runtime import prepare_rows
from xiaoan_eval_core.results import (aggregate_complete_results, seal_complete_results,
                                      validate_complete_results)
from xiaoan_eval_core.runtime import ResponseStore, identity
from xiaoan_eval_core.orchestration import run_orchestration
from .frozen_ingress import file_hash, validate_frozen_spec
from .frozen_provider import ConfiguredProvider, transport_options
from .frozen_cli import write_json
from .frozen_export import export_results_workbook


STAGES = {'extraction', 'rubric', 'assessment', 'relevancy'}


def replace_json(path, value):
    path = Path(path)
    temporary = path.with_name(f'.{path.name}.{uuid4().hex}.tmp')
    try:
        with temporary.open('x', encoding='utf-8') as stream:
            temporary.chmod(0o600)
            json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.write('\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


@contextmanager
def locked_retry_output(directory):
    import fcntl
    with (Path(directory)/'.retry.lock').open('a+') as stream:
        try:
            fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ValueError('An in-place retry is already running in this directory') from exc
        try:
            yield
        finally:
            fcntl.flock(stream, fcntl.LOCK_UN)


def assembled_retry_spec(parent):
    """Recover a retry plan only from a validated, complete balanced assembly."""
    parent = validate_complete_results(parent)
    manifest = parent['manifest']
    if manifest.get('schema_version') != 'balanced-assembly/v1' or manifest.get('manifest_digest') != digest(
            {key:value for key,value in manifest.items() if key != 'manifest_digest'}):
        raise ValueError('A sealed balanced-assembly/v1 result is required')
    plan = parent['plan']
    if (set(manifest['case_ids']) != set(plan['planned_turns']) or
            not manifest.get('source_manifest_digests') or
            any(not row.get('manifest_digest') for row in parent['answers'])):
        raise ValueError('Balanced assembly plan or source answer manifests are missing')
    for row in parent['answers']:
        if (not isinstance(row.get('answer'), str) or
                hashlib.sha256(row['answer'].encode()).hexdigest() != row.get('answer_sha256') or
                digest({key:value for key,value in row.items() if key != 'row_digest'}) != row.get('row_digest')):
            raise ValueError('Balanced assembly answer digest mismatch')
    spec = {**deepcopy(plan), 'plan':deepcopy(plan), 'manifest':deepcopy(manifest),
            'rows':deepcopy(parent['answers']), 'evaluation_config':deepcopy(parent['evaluation_config']),
            'artifacts':deepcopy(parent.get('artifacts', [])), 'provenance':deepcopy(parent.get('provenance', []))}
    prepare_rows(spec)
    return spec


def validate_retry_spec(spec, parent):
    if parent['manifest'].get('schema_version') != 'balanced-assembly/v1':
        return validate_frozen_spec(spec)
    expected = assembled_retry_spec(parent)
    if spec['manifest'] != expected['manifest'] or spec['plan'] != expected['plan'] or spec['rows'] != expected['rows']:
        raise ValueError('Balanced retry plan or frozen answers changed')
    if (spec.get('judges') != expected['judges'] or spec.get('extractor') != expected['extractor'] or
            spec.get('provider_options') != expected.get('provider_options')):
        raise ValueError('Balanced retry model or provider identity changed')
    return spec


def publish_in_place(result, out, spec):
    workbook = out/'.retry-results.xlsx'
    try:
        if (out/'results.xlsx').exists():
            from openpyxl import load_workbook
            existing = load_workbook(out/'results.xlsx', read_only=True, data_only=False)
            try:
                if 'Human Review' in existing:
                    rows = existing['Human Review'].iter_rows(values_only=True)
                    headers = next(rows, ())
                    editable = [i for i, name in enumerate(headers) if name in
                                {'decision','notes','reviewer','reviewed_at','revisions_json'}]
                    if any(any(row[i] is not None for i in editable) for row in rows):
                        raise ValueError('Existing workbook contains human edits; reconcile them before retry')
            finally:
                existing.close()
        export_results_workbook(result, workbook)
        replace_json(out/'results.json', result)
        workbook.replace(out/'results.xlsx')
        replace_json(out/'frozen-input.json', spec)
    finally:
        workbook.unlink(missing_ok=True)


def reconcile_assessment_conflicts(parent):
    """Reclassify complete, non-scoring differences; leave evidence/verdict conflicts for review."""
    resolved = 0
    needs_review = 0
    for env in parent['envelopes']:
        for cell in env['assessments']:
            assessment = cell.get('assessment') or {}
            conflicts = assessment.get('retry_conflicts') or []
            if cell.get('status') != 'PARTIAL' or assessment.get('validation_errors') or not conflicts:
                continue
            scoring_conflict = any(
                c['preserved'].get('verdict') != c['candidate'].get('verdict') or
                c['preserved'].get('evidence', c['preserved'].get('answer_spans', [])) !=
                c['candidate'].get('evidence', c['candidate'].get('answer_spans', []))
                for c in conflicts)
            if scoring_conflict:
                assessment['conflict_review'] = 'REQUIRED'
                needs_review += 1
            elif (assessment.get('dimension_status') == {'faithfulness':'AVAILABLE','correctness':'AVAILABLE'}
                  and assessment.get('requirements_status') == 'AVAILABLE'):
                cell['status'] = assessment['status'] = 'AVAILABLE'
                assessment['conflict_resolution'] = 'PRESERVED_PRIOR_NON_SCORING_FIELDS'
                resolved += 1
    return {'resolved_non_scoring': resolved, 'needs_review': needs_review}


def extend_judges(spec, additions):
    result = validate_frozen_spec(spec)
    if not additions:
        return result
    old = result['judges']
    if len({j['id'] for j in old + additions}) != len(old) + len(additions):
        raise ValueError('Judge identity already planned or duplicated')
    prior = result.get('judge_extension', {}).get('added_judges', [])
    result['plan'] = deepcopy(result['plan'])
    result['judges'] = deepcopy(old + additions)
    result['plan']['judges'] = deepcopy(result['judges'])
    result['judge_extension'] = {
        'source_manifest_digest':result['manifest']['manifest_digest'],
        'added_judges':deepcopy(prior + additions)}
    return validate_frozen_spec(result)


def import_successful_checkpoints(output, sources):
    destination = Path(output)/'checkpoint'
    for source in sources:
        folder = Path(source)/'checkpoint'
        if not folder.is_dir():
            raise ValueError(f'Checkpoint directory missing: {folder}')
        for path in sorted(folder.glob('*.json')):
            partial = path.name.endswith('.partial.json')
            key = path.name.removesuffix('.partial.json') if partial else path.stem
            payload = json.loads(path.read_text())
            if (payload.get('request_hash') != key or 'response' not in payload or
                    (partial and (not isinstance(payload['response'], dict) or
                                  not isinstance(payload['response'].get('claims'), list) or
                                  not isinstance(payload['response'].get('requirements'), list)))):
                raise ValueError(f'Invalid checkpoint: {path}')
            destination.mkdir(parents=True, exist_ok=True)
            target = destination/path.name
            if partial and (destination/(key+'.json')).exists():
                continue
            if target.exists():
                if target.read_bytes() != path.read_bytes():
                    raise ValueError(f'Conflicting checkpoint: {target}')
            else:
                try:
                    os.link(path, target)
                except FileExistsError:
                    if target.read_bytes() != path.read_bytes():
                        raise ValueError(f'Conflicting checkpoint: {target}')


def compatible_schema_update(previous, current):
    """Allow only strict object closure and missing array-item shape completion."""
    def is_refinement(before, after):
        if before == after:
            return True
        if not isinstance(before, dict) or not isinstance(after, dict):
            return False
        old = deepcopy(before)
        new = deepcopy(after)
        if old.get('type') == 'object' and 'additionalProperties' not in old:
            if new.pop('additionalProperties', None) is not False:
                return False
        if old.get('type') == 'array' and 'items' not in old:
            item = new.pop('items', None)
            if not isinstance(item, dict):
                return False
        old_props = old.pop('properties', None)
        new_props = new.pop('properties', None)
        if old != new or (old_props is None) != (new_props is None):
            return False
        return old_props is None or (old_props.keys() == new_props.keys() and
                all(is_refinement(old_props[k], new_props[k]) for k in old_props))

    if previous.keys() != current.keys():
        return False
    for name, old in previous.items():
        new = current[name]
        if old == new:
            continue
        if not name.startswith('schemas/') or not name.endswith('.json'):
            return False
        import json
        from hashlib import sha256
        before = json.loads(old['content'])
        after = json.loads(new['content'])
        # Definitions are new structural constraints, never a changed old contract.
        definitions = after.pop('$defs', {})
        if any(not isinstance(item, dict) for item in definitions.values()) or not is_refinement(before, after):
            return False
        if sha256(old['content'].encode()).hexdigest() != old['sha256'] or sha256(new['content'].encode()).hexdigest() != new['sha256']:
            return False
    return True


def retry_evaluation(spec, parent, output, *, stages=None, answer_ids=None,
                     judge_ids=None, provider=None, execute=False, max_workers=2,
                     new_evaluator_cohort=False, adopt_current_runtime=False, dry_run=False,
                     provider_max_inflight=None, provider_requests_per_second=None, checkpoint_sources=None,
                     unavailable_only=False, reconcile_only=False):
    parent = validate_complete_results(parent)
    spec = validate_retry_spec(spec, parent)
    parent = {**parent, 'envelopes': deepcopy(parent['envelopes'])}
    conflict_summary = reconcile_assessment_conflicts(parent)
    if parent['manifest'] != spec['manifest']:
        raise ValueError('Parent manifest differs from frozen input')
    old = {row['answer_id']:row for row in parent['answers']}
    new = {row['answer_id']:row for row in spec['rows']}
    current_config = snapshot()
    if adopt_current_runtime:
        if new_evaluator_cohort or not compatible_schema_update(parent['evaluation_config'], current_config):
            raise ValueError('Current evaluator changes are not the compatible schema-only update')
        spec['evaluation_config'] = current_config
        spec['evaluation_provider_options'] = transport_options(ConfiguredProvider().values)
    if not new_evaluator_cohort and (parent['evaluation_config'] != current_config or
            spec.get('evaluation_config', spec['manifest'].get('evaluator_config')) != current_config) and not adopt_current_runtime:
        raise ValueError('Evaluator configuration changed; start a new evaluation cohort')
    if not new_evaluator_cohort and not adopt_current_runtime and parent['plan'].get('provider_options') != spec.get('provider_options'):
        raise ValueError('Evaluator provider options differ from parent')
    if any(env.get('human_review') for env in parent['envelopes']):
        raise ValueError('Export and reconcile Human Review before changing automatic evaluations')
    old_inventories = {inv['answer_id']:inv for inv in parent['inventories']}
    for aid in old.keys() & new.keys():
        previous = old[aid]
        if 'inventory' in previous:
            inventory = old_inventories.get(aid)
            request = extraction_request(new[aid], identity(spec['extractor']))
            supplied = previous['inventory']
            if (inventory is None or not isinstance(supplied, dict) or
                    validate_inventory(supplied, request)['inventory_id'] != inventory['inventory_id']):
                raise ValueError('Unverified auxiliary inventory on frozen answer')
            previous = {k:v for k,v in previous.items() if k != 'inventory'}
        if previous != new[aid]:
            raise ValueError('Frozen answer changed without a new answer_id')
    changed = new.keys() - old.keys()
    removed = old.keys() - new.keys()
    if removed and {old[aid]['planned_unit_id'] for aid in removed} != {new[aid]['planned_unit_id'] for aid in changed}:
        raise ValueError('Subject retry must replace exactly the selected planned rows')
    if len(new) != len(old) or {r['planned_unit_id'] for r in new.values()} != {r['planned_unit_id'] for r in old.values()}:
        raise ValueError('Subject retry changed the planned matrix')
    if new_evaluator_cohort:
        if stages or answer_ids or judge_ids or changed:
            raise ValueError('New evaluator cohort requires every frozen answer and stage')
        spec['evaluation_config'] = current_config
        spec['evaluation_provider_options'] = transport_options(ConfiguredProvider().values)
    requested = set(stages or STAGES)
    if not requested or requested - STAGES:
        raise ValueError('Unknown retry stage')
    selected_answers = set(answer_ids or new)
    if selected_answers - new.keys():
        raise ValueError('Unknown answer_id')
    judges = {j['id'] for j in spec['judges']}
    selected_judges = set(judge_ids or judges)
    if selected_judges - judges:
        raise ValueError('Unknown judge_id')
    old_env = {env['answer_id']:env for env in parent['envelopes']}
    cached_extractions = {receipt['request_digest'] for receipt in parent['stages']
                          if receipt.get('task') == 'extract_claims' and
                          receipt.get('execution_status') == 'SUCCEEDED' and 'raw_response' in receipt}
    selected = set()
    supplied_inventories = {}
    for aid in selected_answers | changed:
        if new[aid]['status'] != 'AVAILABLE':
            continue
        env = old_env.get(aid, {})
        if new_evaluator_cohort or aid in changed or 'extraction' in requested and env.get('inventory_id') is None:
            selected.add(('extract_claims', aid, None))
        for name, task in (('rubric', 'rubric'), ('assessment', 'assess_claims')):
            if name not in requested and aid not in changed:
                continue
            branch = 'rubric' if name == 'rubric' else 'assessments'
            cells = {cell['judge_id']:cell for cell in env.get(branch, [])}
            for judge in selected_judges if aid not in changed else judges:
                status = cells.get(judge, {}).get('status', 'UNAVAILABLE')
                conflicts_only = (task == 'assess_claims' and status == 'PARTIAL' and
                                  cells[judge].get('assessment', {}).get('retry_conflicts') and
                                  not cells[judge]['assessment'].get('validation_errors')) if judge in cells else False
                if new_evaluator_cohort or aid in changed or (not conflicts_only and
                        (status == 'UNAVAILABLE' if unavailable_only else status != 'AVAILABLE')):
                    selected.add((task, aid, judge))
                    if task == 'assess_claims' and env.get('inventory_id') is None:
                        selected.add(('extract_claims', aid, None))
                    elif task == 'assess_claims' and aid in old_inventories:
                        request = extraction_request(new[aid], identity(spec['extractor']))
                        if digest(request) not in cached_extractions:
                            inventory = old_inventories[aid]
                            payload = {'binding':inventory['binding'], 'claims':inventory['claims']}
                            if validate_inventory(payload, request)['inventory_id'] != inventory['inventory_id']:
                                raise ValueError('Preserved inventory does not match frozen answer')
                            supplied_inventories[aid] = payload
        if (new_evaluator_cohort or 'relevancy' in requested or aid in changed) and (new_evaluator_cohort or aid in changed or env.get('relevancy', {}).get('status') != 'AVAILABLE'):
            selected.update({('relevancy_generation', aid, None), ('relevancy_embedding', aid, None)})
        if 'extraction' in requested and ('extract_claims', aid, None) in selected and aid not in supplied_inventories:
            selected.update(('assess_claims', aid, judge) for judge in judges)
    if not selected and not changed and not conflict_summary['resolved_non_scoring'] and not (dry_run and conflict_summary['needs_review']):
        raise ValueError('No unavailable selected stage cells to retry')
    out = Path(output)
    in_place = (out/'results.json').exists()
    if in_place and new_evaluator_cohort:
        raise ValueError('A full new evaluator cohort cannot reuse existing checkpoints in place')
    if reconcile_only and (not in_place or new_evaluator_cohort or adopt_current_runtime or checkpoint_sources):
        raise ValueError('Conflict-only reconciliation needs an unchanged sealed result in its own directory')
    if in_place and validate_complete_results(json.loads((out/'results.json').read_text()))['result_generation'] != parent['result_generation']:
        raise ValueError('In-place retry source changed; reload the current sealed results')
    if dry_run:
        return {'parent_generation':parent['result_generation'],
                'new_evaluator_cohort':new_evaluator_cohort,
                'adopt_current_runtime':adopt_current_runtime,
                'judges':[j['id'] for j in spec['judges']],
                'selected':{} if reconcile_only else dict(sorted(Counter(task for task,_,_ in selected).items())),
                'conflict_reconciliation':conflict_summary, 'in_place':in_place}
    if reconcile_only:
        if not conflict_summary['resolved_non_scoring']:
            raise ValueError('No non-scoring conflicts can be resolved automatically')
        parent['provenance'] = [*parent.get('provenance', []),
                                {'operation':'reconcile_assessment_conflicts',
                                 'parent_generation':parent['result_generation'],
                                 'conflict_reconciliation':conflict_summary}]
        parent['aggregates'] = aggregate_complete_results(parent)
        from xiaoan_eval_core.costs import build_answer_costs
        parent['aggregates']['answer_costs'] = build_answer_costs(parent['answers'])
        seal_complete_results(parent)
        validate_complete_results(parent)
        publish_in_place(parent,out,spec)
        return parent
    call = provider or (ConfiguredProvider(out/'provider-artifacts') if execute else None)
    if isinstance(call, ConfiguredProvider) and transport_options(call.values) != spec.get('evaluation_provider_options',spec.get('provider_options', {})):
        raise ValueError('Provider options differ from frozen input')
    import_successful_checkpoints(out, checkpoint_sources or [])
    selection={'parent_generation':parent['result_generation'],
               'new_evaluator_cohort':new_evaluator_cohort,
               'selected':sorted([list(x) for x in selected],key=str)}
    if in_place:
        replace_json(out/'retry-selection.json',selection)
    else:
        write_json(out/'retry-selection.json',selection)
    settings = workflow()
    def selected_request(request):
        task = request['task']
        judge = request.get('identity', {}).get('id') if task in {'rubric', 'assess_claims'} else None
        return (task,request.get('answer_id'),judge) in selected
    store = ResponseStore(out/'checkpoint',max_workers=max_workers,
                          provider_max_inflight=(provider_max_inflight if provider_max_inflight is not None else
                              spec.get('provider_max_inflight',settings['provider_max_inflight'])),
                          max_attempts=spec.get('max_attempts',settings['max_attempts']),
                          provider_options=spec.get('evaluation_provider_options',spec.get('provider_options')),
                          provider_requests_per_second=(provider_requests_per_second if provider_requests_per_second is not None else
                              spec.get('provider_requests_per_second',settings.get('provider_requests_per_second',0))),
                          request_filter=selected_request)
    for aid, payload in supplied_inventories.items():
        request = extraction_request(new[aid], identity(spec['extractor']))
        key = store.request_digest(request)
        store.memory[key] = {'request_hash':key, 'response':deepcopy(payload)}
    for receipt in ([] if new_evaluator_cohort else parent['stages']):
        if (receipt.get('execution_status') == 'SUCCEEDED' and
                receipt.get('answer_id') in new and 'raw_response' in receipt):
            key = receipt['request_digest']
            if digest(receipt['request']) != key or receipt.get('stage_id') != key:
                raise ValueError('Parent stage request digest mismatch')
            if receipt.get('availability') == 'PARTIAL':
                if receipt.get('task') == 'assess_claims' and selected_request(receipt['request']):
                    # The sealed parent can contain several generations of
                    # this same request. Its latest receipt is the retry base.
                    store.partial[key] = deepcopy(receipt['output'])
                if (receipt.get('task') == 'assess_claims' and
                      ('assess_claims', receipt['answer_id'], receipt.get('identity', {}).get('id')) in selected and
                      receipt.get('output', {}).get('validation_errors')):
                    aid = receipt['answer_id']
                    inventory = old_inventories.get(aid)
                    judge = next((j for j in spec['judges'] if j['id'] == receipt['identity']['id']), None)
                    if inventory and judge:
                        request = assessment_request(new[aid], {k:inventory[k] for k in
                            ('contract','binding','extractor','claims','inventory_id')}, judge)
                        if request['binding'] == receipt['request']['binding']:
                            store.partial[store.request_digest(request)] = deepcopy(receipt['output'])
                continue
            cached = {'request_hash':key, 'response':receipt['raw_response']}
            if key in store.memory and store.memory[key] != cached:
                raise ValueError('Conflicting successful parent stage outputs')
            store.memory[key] = cached

    def allowed(request):
        if not selected_request(request):
            raise LookupError('Stage not selected for retry')
        if call is None:
            raise LookupError('Provider not configured')
        return call(request)

    def embed(texts, task):
        if ('relevancy_embedding', task['answer_id'], None) not in selected:
            raise LookupError('Stage not selected for retry')
        if call is None or not hasattr(call, 'embeddings'):
            raise LookupError('Embedding provider not configured')
        return call.embeddings(texts,task)

    if not in_place:
        write_json(out/'frozen-input.json', spec)
    try:
        execution_spec = deepcopy(spec)
        for row in execution_spec['rows']:
            if row['answer_id'] in supplied_inventories:
                row['inventory'] = supplied_inventories[row['answer_id']]
        result = run_orchestration(execution_spec, faithfulness_provider=allowed, rubric_provider=allowed,
                                   generation_provider=allowed, embedding_provider=embed,
                                   checkpoint_dir=out/'checkpoint', max_workers=max_workers, store=store)
        for row in result['answers']:
            row.pop('inventory', None)
    finally:
        if provider is None and isinstance(call, ConfiguredProvider):
            call.close()
    for env in result['envelopes']:
        aid = env['answer_id']
        prior = old_env.get(aid)
        if prior is None or new_evaluator_cohort:
            continue
        if ('extract_claims', aid, None) not in selected:
            env['inventory_id'] = prior['inventory_id']
        for branch, task in (('rubric','rubric'), ('assessments','assess_claims')):
            attempted = {cell['judge_id'] for cell in env[branch]}
            preserved = {cell['judge_id']:deepcopy(cell) for cell in prior[branch]
                         if (task, aid, cell['judge_id']) not in selected}
            def retained(cell):
                judge = cell['judge_id']
                if judge in preserved:
                    return preserved[judge]
                if (task, aid, judge) in selected:
                    return cell
                # A newly appended Judge may have a checkpoint for an answer
                # outside --answer-ids. Do not attach it without its inventory.
                return {key:deepcopy(value) for key,value in cell.items()
                        if key in {'answer_id','case_id','turn','subject_id','judge_id','judge_identity','evaluator_role'}} | {
                            'status':'UNAVAILABLE','execution_status':'SKIPPED','reason':'NOT_SELECTED'}
            env[branch] = [retained(cell) for cell in env[branch]]
            if attempted != judges:
                raise ValueError('Retry dropped a planned Judge cell')
        if ('relevancy_generation', aid, None) not in selected:
            env['relevancy'] = deepcopy(prior['relevancy'])
        env['checks'] = deepcopy(prior['checks'])
    preserved_inventories = [inv for inv in ([] if new_evaluator_cohort else parent['inventories'])
                             if inv['answer_id'] in new and ('extract_claims',inv['answer_id'],None) not in selected]
    result['inventories'] = [*preserved_inventories,
                             *(inv for inv in result['inventories']
                               if ('extract_claims',inv['answer_id'],None) in selected)]
    result['stages'] = [r for r in ([] if new_evaluator_cohort else parent['stages']) if r.get('answer_id') in new] + [
        r for r in result['stages'] if (r.get('task'),r.get('answer_id'),
              r.get('identity',{}).get('id') if r.get('task') in {'rubric','assess_claims'} else None) in selected]
    for env in result['envelopes']:
        env['stage_refs'] = list(dict.fromkeys(r['stage_id'] for r in result['stages'] if r.get('answer_id') == env['answer_id']))
    result['provenance'] = [*parent.get('provenance', []),
                            {'operation':'new_evaluator_cohort' if new_evaluator_cohort else 'selective_retry',
                             'parent_generation':parent['result_generation'],
                             'judge_extension':deepcopy(spec.get('judge_extension')),
                             'checkpoint_sources':[str(Path(path).resolve()) for path in checkpoint_sources or []],
                             'runtime_migration': {'prior_evaluator_config': parent['evaluation_config'],
                                                   'prior_provider_options': parent['plan'].get('provider_options'),
                                                   'current_provider_options': spec['evaluation_provider_options'],
                                                   'rule': 'strict_object_closure_and_array_item_shape_only'} if adopt_current_runtime else None,
                             'selected':sorted([list(x) for x in selected], key=str)}]
    result['provenance'][-1]['execution_limits'] = {
        'max_workers':max_workers, 'provider_max_inflight':store.provider_max_inflight,
        'provider_requests_per_second':store.rate}
    result['provenance'][-1]['conflict_reconciliation'] = conflict_summary
    for folder in ('provider-artifacts','subject-checkpoint','checkpoint'):
        for artifact in sorted((out/folder).glob('*')):
            if artifact.is_file() and artifact.suffix in ('.json','.jsonl'):
                result['artifacts'].append({'id':file_hash(artifact),
                    'path':str(artifact.relative_to(out)),'sha256':file_hash(artifact),'type':folder})
    result['aggregates'] = aggregate_complete_results(result)
    from xiaoan_eval_core.costs import build_answer_costs
    result['aggregates']['answer_costs'] = build_answer_costs(result['answers'])
    seal_complete_results(result)
    validate_complete_results(result)
    if in_place:
        publish_in_place(result,out,spec)
    else:
        write_json(out/'results.json',result)
        export_results_workbook(result,out/'results.xlsx')
    return result
