"""Versioned, human-readable XLSX contract for evaluation deliverables."""

from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path, PurePosixPath
from typing import Any, Mapping, Sequence
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .report_model import ReportModel, SCHEMA_VERSION, canonical_digest


SHEETS = (
    "00_Overview", "01_Cases", "02_Turns", "03_Metrics", "04_Baseline",
    "05_Experiments", "06_Human_Review", "07_Stability", "08_Metadata",
    "09_Data_Dictionary",
)

HEADERS: Mapping[str, tuple[str, ...]] = {
    "00_Overview": ("section", "metric", "value", "status", "interpretation"),
    "01_Cases": ("component_run_id", "case_id", "comparison_key", "status", "artifact_state", "automatic_score", "human_score", "final_score", "final_source", "quality_status", "quality_verdict", "quality_threshold", "quality_weights", "execution_status", "oracle_approved", "hard_gate_passed", "review_status", "adjudication_status", "failure_stage", "cohorts", "total_ms", "evidence_refs"),
    "02_Turns": ("row_kind", "component_run_id", "case_id", "turn", "comparison_key", "status", "user_text_id", "assistant_text_id", "route_id", "safety_level", "ground_refs", "response_sha256", "total_ms", "ttft_ms", "review_status", "automatic_score", "human_score", "final_score", "final_source", "text_id", "role", "chunk_index", "chunk_count", "text_sha256", "content"),
    "03_Metrics": ("row_key", "comparison_key", "component_run_id", "case_id", "turn", "metric_id", "dimension", "raw_score", "normalized_score", "weight", "contribution", "status", "reason", "evidence_refs", "score_source", "baseline_value", "delta"),
    "04_Baseline": ("domain", "grain", "key", "metric", "baseline_value", "candidate_value", "delta", "status", "reason"),
    "05_Experiments": ("experiment_id", "hypothesis", "control", "candidate", "repetitions", "target_result", "non_target_regression", "guardrails", "verdict", "next_action"),
    "06_Human_Review": ("review_id", "case_id", "turn", "response_sha256", "status", "trigger", "reviewer_id", "submitted_at", "confidence", "notes"),
    "07_Stability": ("section", "metric", "value", "status"),
    "08_Metadata": ("key", "value", "value_type"),
    "09_Data_Dictionary": ("sheet", "column", "type", "description"),
}

SHEET_LABELS: Mapping[str, str] = {
    "00_Overview": "總覽",
    "01_Cases": "案例",
    "02_Turns": "對話輪次",
    "03_Metrics": "評估指標",
    "04_Baseline": "基準比較",
    "05_Experiments": "受控實驗",
    "06_Human_Review": "人工複核",
    "07_Stability": "穩定性",
    "08_Metadata": "中繼資料",
    "09_Data_Dictionary": "資料字典",
}

HEADER_LABELS: Mapping[str, str] = {
    "section": "區段", "metric": "指標", "value": "值", "status": "狀態",
    "interpretation": "解讀", "component_run_id": "元件執行 ID", "case_id": "案例 ID",
    "comparison_key": "比較鍵", "artifact_state": "交付物狀態", "automatic_score": "自動評分",
    "human_score": "人工評分", "final_score": "最終評分", "final_source": "最終評分來源",
    "oracle_approved": "Oracle已審閱", "quality_weights": "最終維度權重", "execution_status": "執行及門檻狀態", "quality_status": "品質可評狀態", "quality_verdict": "品質結論", "quality_threshold": "品質門檻",
    "hard_gate_passed": "硬性門檻是否通過", "review_status": "複核狀態",
    "adjudication_status": "裁決狀態", "failure_stage": "失敗階段", "cohorts": "群組",
    "total_ms": "總耗時（毫秒）", "evidence_refs": "證據引用", "turn": "輪次",
    "user_text_id": "使用者文字 ID", "assistant_text_id": "小安回應文字 ID", "route_id": "路由 ID",
    "safety_level": "安全等級", "ground_refs": "依據引用", "response_sha256": "回應 SHA-256",
    "ttft_ms": "首字延遲（毫秒）", "row_key": "資料列鍵", "metric_id": "指標 ID",
    "dimension": "評分維度", "raw_score": "原始分數", "normalized_score": "標準化分數",
    "weight": "權重", "contribution": "加權貢獻", "reason": "原因", "score_source": "評分來源",
    "baseline_value": "基準值", "delta": "差異", "domain": "比較領域", "grain": "比較粒度",
    "key": "鍵", "candidate_value": "候選值", "experiment_id": "實驗 ID", "hypothesis": "假設",
    "control": "控制組", "candidate": "候選組", "repetitions": "重複次數",
    "target_result": "目標結果", "non_target_regression": "非目標退步", "guardrails": "護欄",
    "verdict": "結論", "next_action": "下一步", "review_id": "複核 ID", "trigger": "觸發原因",
    "reviewer_id": "複核者 ID", "submitted_at": "提交時間", "confidence": "信心程度",
    "notes": "備註", "value_type": "值類型", "sheet": "工作表", "column": "欄位",
    "type": "類型", "description": "說明", "text_id": "文字 ID", "role": "角色",
    "chunk_index": "分段序號", "chunk_count": "分段總數", "text_sha256": "文字 SHA-256",
    "content": "內容",
    "row_kind": "資料列類型",
}

OVERVIEW_DISPLAY: Mapping[str, Mapping[str, str]] = {
    "section": {
        "Result": "結果", "Quality": "品質", "Coverage": "覆蓋範圍", "Performance": "效能",
        "Baseline": "基準比較", "Experiment": "實驗", "Stability": "穩定性", "Dimension": "評分維度", "Decision": "決策摘要",
    },
    "metric": {
        "Artifact state": "交付物狀態", "Evaluation verdict": "評估結論", "Overall score": "整體分數",
        "Pass rate": "通過率", "Cases": "案例數", "Metric rows": "指標資料列數",
        "Mean total latency (ms)": "平均總延遲（毫秒）", "Comparison": "比較結果",
        "Result": "實驗結果", "Classification": "分類",
    },
    "value": {
        "FINAL": "已完成", "PENDING_REVIEW": "待人工複核", "PASS": "通過", "FAIL": "未通過",
        "ATTACHED": "已附上", "NOT_RUN": "未執行", "NOT_MEASURED": "未測量",
        "DESCRIPTIVE_ONLY": "僅描述性結果",
    },
    "status": {
        "FINAL": "已完成", "PENDING_REVIEW": "待人工複核", "AVAILABLE": "可用",
        "UNAVAILABLE": "不可用", "NOT_RUN": "未執行",
    },
    "interpretation": {
        "Lifecycle state of this workbook.": "此工作簿目前的生命週期狀態。",
        "Quality verdict; speed is reported separately.": "品質結論；執行速度另行呈現。",
        "Mean case weighted score on the rating-rule scale.": "依評分規則量尺計算的案例加權平均分數。",
        "Mean available final case score.": "所有可用案例最終分數的平均值。",
        "Share of cases with PASS status.": "狀態為通過的案例比例。",
        "Number of subject cases.": "本次受測案例數。",
        "Number of turn-level metric facts.": "輪次層級的指標事實數。",
        "Number of metric facts including score-source rows.": "包含不同評分來源資料列的指標事實數。",
        "Mean over cases with telemetry; missing values are excluded.": "僅計算具有遙測資料的案例平均值，缺漏值不納入。",
        "Mean over cases with telemetry.": "具有遙測資料之案例的平均值。",
        "Baseline deltas are shown only when comparison is attached.": "只有附上基準比較時才呈現差異。",
        "Baseline evidence attachment state.": "基準比較證據的附加狀態。",
        "Recommendations remain hypotheses without controlled evidence.": "沒有受控證據時，改善建議仍只是待驗證假設。",
        "Controlled experiment attachment state.": "受控實驗結果的附加狀態。",
        "Repeatability only; a stable wrong answer is still wrong.": "只衡量可重複性；穩定地答錯仍然是錯。",
        "Repeatability only; correctness remains separate.": "只衡量可重複性；正確性需分開判斷。",
        "Mean automatic judge score across evaluated turns.": "所有已評估輪次的自動 Judge 平均分數。",
        "Mean recorded judge score.": "已記錄 Judge 分數的平均值。",
        "Router accepted accuracy": "Router 可接受路由準確率",
        "Router preferred accuracy": "Router 首選路由準確率",
        "Capsule claim alignment": "Capsule claim 對齊率",
        "Overall claim support rate": "整體 claim 支持率",
        "Actual route falls within the reviewed accepted route set; inspect 03_Metrics for case/turn evidence.": "實際路由落在已審核的可接受集合內；請到 03_Metrics 查看案例／輪次證據。",
        "Actual route equals the reviewed preferred route; use this to assess crisis_sop/baseline/capsule routing.": "實際路由等於已審核的首選路由；用於判斷 crisis_sop、baseline、capsule 分流。",
        "Judge claims with valid injected Capsule evidence; observational, not causal proof.": "具有有效注入 Capsule 證據的 Judge claims；屬觀察性證據，不是因果證明。",
        "Semantic attribution support across substantive claims; inspect layer support for Capsule/Ground detail.": "實質 claims 的語義支持率；Capsule／Ground 細節請查看 layer support。",
    },
}

MAX_PACKAGE_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_PARTS = 5000


@dataclass(frozen=True)
class WorkbookFacts:
    schema_version: str
    generation_id: str
    artifact_state: str
    core_digest: str
    lifecycle_digest: str
    manifest: Mapping[str, Any]
    cases: tuple[Mapping[str, Any], ...]
    turns: tuple[Mapping[str, Any], ...]
    metrics: tuple[Mapping[str, Any], ...]
    baseline: tuple[Mapping[str, Any], ...]
    experiments: tuple[Mapping[str, Any], ...]
    human_review: tuple[Mapping[str, Any], ...]
    stability: tuple[Mapping[str, Any], ...]
    text_content: tuple[Mapping[str, Any], ...]


def write_workbook(model: ReportModel, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows_by_sheet = {
        "00_Overview": model.overview,
        "01_Cases": model.cases,
        "02_Turns": _turn_sheet_rows(model.turns, model.text_content),
        "03_Metrics": model.metrics,
        "04_Baseline": model.baseline,
        "05_Experiments": model.experiments,
        "06_Human_Review": model.human_review,
        "07_Stability": model.stability,
        "08_Metadata": _metadata_rows(model),
        "09_Data_Dictionary": _dictionary_rows(),
    }
    for index, name in enumerate(SHEETS):
        sheet = workbook.create_sheet(name)
        _write_sheet(sheet, HEADERS[name], rows_by_sheet[name], index)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.properties.title = "小安評估結果"
    workbook.properties.subject = f"資料契約 {model.schema_version}；產生批次 {model.generation_id}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_workbook(path: Path) -> WorkbookFacts:
    preflight_ooxml(path)
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    if tuple(workbook.sheetnames) != SHEETS:
        raise ValueError("workbook sheet contract does not match the supported schema")
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            raise ValueError(f"hidden sheets are not allowed: {sheet.title}")
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise ValueError(f"formulas are not allowed in imported workbooks: {sheet.title}")
    metadata_rows = _read_rows(workbook["08_Metadata"], HEADERS["08_Metadata"])
    metadata = {str(row["key"]): _decode_metadata(row) for row in metadata_rows}
    if metadata.get("schema_version") not in {"2.0", SCHEMA_VERSION}:
        raise ValueError(f"unsupported workbook schema: {metadata.get('schema_version')!r}")
    manifest = metadata.get("manifest", {})
    if not isinstance(manifest, Mapping):
        raise ValueError("workbook manifest metadata must be an object")
    turns, text_content = _read_turn_sheet(workbook["02_Turns"])
    facts = WorkbookFacts(
        schema_version=str(metadata["schema_version"]),
        generation_id=str(metadata["generation_id"]),
        artifact_state=str(metadata["artifact_state"]),
        core_digest=str(metadata["core_digest"]),
        lifecycle_digest=str(metadata["lifecycle_digest"]),
        manifest=dict(manifest),
        cases=tuple(_read_rows(workbook["01_Cases"], tuple(name for name in HEADERS["01_Cases"] if metadata.get("schema_version") != "2.0" or name not in {"quality_status", "quality_verdict", "quality_threshold", "quality_weights", "execution_status", "oracle_approved"}))),
        turns=turns,
        metrics=tuple(_read_rows(workbook["03_Metrics"], HEADERS["03_Metrics"])),
        baseline=tuple(_read_rows(workbook["04_Baseline"], HEADERS["04_Baseline"])),
        experiments=tuple(_read_rows(workbook["05_Experiments"], HEADERS["05_Experiments"])),
        human_review=tuple(_read_rows(workbook["06_Human_Review"], HEADERS["06_Human_Review"])),
        stability=tuple(_read_rows(workbook["07_Stability"], HEADERS["07_Stability"])),
        text_content=text_content,
    )
    expected_core = canonical_digest({
        "schema_version": facts.schema_version, "manifest": facts.manifest,
        "cases": facts.cases, "turns": facts.turns, "metrics": facts.metrics,
        "text_content": facts.text_content,
    })
    expected_lifecycle = canonical_digest({
        "artifact_state": facts.artifact_state, "baseline": facts.baseline,
        "experiments": facts.experiments, "human_review": facts.human_review,
        "stability": facts.stability,
    })
    if facts.core_digest != expected_core or facts.lifecycle_digest != expected_lifecycle:
        raise ValueError("workbook logical digest validation failed")
    return facts


def preflight_ooxml(path: Path) -> None:
    if not path.is_file() or path.stat().st_size > MAX_PACKAGE_BYTES:
        raise ValueError("workbook package is missing or exceeds the size limit")
    try:
        with ZipFile(path) as package:
            entries = package.infolist()
            if len(entries) > MAX_PARTS:
                raise ValueError("workbook package contains too many parts")
            total = 0
            seen = set()
            for entry in entries:
                name = entry.filename
                pure = PurePosixPath(name)
                if name in seen or pure.is_absolute() or ".." in pure.parts or "\\" in name:
                    raise ValueError("workbook package contains an unsafe or duplicate path")
                seen.add(name)
                total += entry.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("workbook package exceeds the uncompressed size limit")
                lowered = name.lower()
                if lowered.endswith(("vbaproject.bin", ".ole", ".bin")) or "activex" in lowered or "embeddings" in lowered:
                    raise ValueError("active or embedded workbook content is not allowed")
                if lowered.endswith((".xml", ".rels")):
                    content = package.read(entry)
                    upper = content.upper()
                    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                        raise ValueError("DTD and entity declarations are not allowed")
                    if lowered.endswith(".rels") and b'TARGETMODE="EXTERNAL"' in upper:
                        raise ValueError("external workbook relationships are not allowed")
    except BadZipFile as exc:
        raise ValueError("invalid XLSX package") from exc


def _write_sheet(sheet, headers: Sequence[str], rows: Sequence[Mapping[str, Any]], index: int) -> None:
    sheet.append(list(headers))
    for row in rows:
        values = [_excel_value(_display_value(sheet.title, header, row.get(header))) for header in headers]
        sheet.append(values)
        for cell, value in zip(sheet[sheet.max_row], values, strict=True):
            if isinstance(value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{_column_letter(len(headers))}{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:1"
    sheet.oddHeader.center.text = f"小安評估結果｜{SHEET_LABELS[sheet.title]}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
        label = HEADER_LABELS.get(str(cell.value), str(cell.value))
        cell.comment = Comment(f"{label}（{cell.value}）", "XiaoAn Evaluation")
    for column_index, header in enumerate(headers, 1):
        width = 14
        if sheet.title == "00_Overview":
            width = {"section": 15, "metric": 30, "value": 22, "status": 22, "interpretation": 58}[header]
        if header in {"interpretation", "reason", "notes", "content", "trigger", "evidence_refs"}:
            width = 58 if sheet.title == "00_Overview" else 45
        elif "key" in header or "sha256" in header:
            width = 24
        sheet.column_dimensions[_column_letter(column_index)].width = width
        for cell in list(sheet.columns)[column_index - 1][1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        table = Table(displayName=f"tbl_{index:02d}", ref=f"A1:{_column_letter(len(headers))}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    if "status" in headers and sheet.max_row > 1:
        column = _column_letter(headers.index("status") + 1)
        sheet.conditional_formatting.add(
            f"{column}2:{column}{sheet.max_row}",
            CellIsRule(operator="equal", formula=['"FAIL"'], fill=PatternFill("solid", fgColor="F4CCCC")),
        )


def _excel_value(value: Any) -> Any:
    if isinstance(value, float):
        return float(format(value, ".15g"))
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _display_value(sheet: str, header: str, value: Any) -> Any:
    if sheet == "00_Overview" and isinstance(value, str):
        return OVERVIEW_DISPLAY.get(header, {}).get(value, value)
    return value


def _turn_sheet_rows(turns: Sequence[Mapping[str, Any]], text_content: Sequence[Mapping[str, Any]]):
    """Combine turn facts and ordered text chunks in one public worksheet."""
    chunks_by_turn: dict[tuple[str, int], list[Mapping[str, Any]]] = {}
    for chunk in text_content:
        key = (str(chunk.get("case_id")), int(chunk.get("turn", 0)))
        chunks_by_turn.setdefault(key, []).append(chunk)
    rows = []
    base_headers = tuple(header for header in HEADERS["02_Turns"] if header not in {
        "row_kind", "text_id", "role", "chunk_index", "chunk_count", "text_sha256", "content"
    })
    for turn in turns:
        base = {header: turn.get(header) for header in base_headers}
        rows.append({**base, "row_kind": "turn"})
        key = (str(turn.get("case_id")), int(turn.get("turn", 0)))
        for chunk in sorted(chunks_by_turn.get(key, ()), key=lambda item: (0 if item.get("role") == "user_input" else 1, int(item.get("chunk_index", 0)))):
            rows.append({
                **base, "row_kind": "text", "text_id": chunk.get("text_id"),
                "role": chunk.get("role"), "chunk_index": chunk.get("chunk_index"),
                "chunk_count": chunk.get("chunk_count"), "text_sha256": chunk.get("text_sha256"),
                "content": chunk.get("content"),
            })
    return tuple(rows)


def _read_turn_sheet(sheet):
    rows = _read_rows(sheet, HEADERS["02_Turns"])
    turns = []
    text_content = []
    for row in rows:
        kind = str(row.get("row_kind") or "turn").lower()
        if kind == "text":
            text_content.append({key: row.get(key) for key in (
                "text_id", "case_id", "turn", "role", "chunk_index", "chunk_count", "text_sha256", "content"
            )})
            continue
        turns.append({key: row.get(key) for key in HEADERS["02_Turns"] if key not in {
            "row_kind", "text_id", "role", "chunk_index", "chunk_count", "text_sha256", "content"
        }})
    return tuple(turns), tuple(text_content)


def _metadata_rows(model: ReportModel):
    values = {
        "schema_version": model.schema_version,
        "generation_id": model.generation_id,
        "artifact_state": model.artifact_state,
        "core_digest": model.core_digest,
        "lifecycle_digest": model.lifecycle_digest,
        "manifest": dict(model.manifest),
    }
    return tuple({"key": key, "value": _excel_value(value), "value_type": "json" if isinstance(value, Mapping) else "text"} for key, value in values.items())


def _decode_metadata(row: Mapping[str, Any]) -> Any:
    value = row.get("value")
    if row.get("value_type") == "json" and isinstance(value, str):
        return json.loads(value)
    return value


def _dictionary_rows():
    descriptions = {
        "comparison_key": "綁定內容與評分脈絡、可跨執行比較的穩定鍵。",
        "row_key": "工作簿中每筆資料列的唯一識別。",
        "status": "可用性或生命週期狀態；數值零仍保留為真正的數值。",
        "score_source": "分數來源：自動評分、人工評分或裁決結果。",
        "content": "原文分段；請依 chunk_index 順序還原，並以 text_sha256 驗證。",
    }
    return tuple(
        {
            "sheet": f"{SHEET_LABELS[sheet]}（{sheet}）",
            "column": f"{HEADER_LABELS.get(column, column)}（{column}）",
            "type": "具型別欄位",
            "description": descriptions.get(column, f"{SHEET_LABELS[sheet]}工作表的「{HEADER_LABELS.get(column, column)}」欄位。"),
        }
        for sheet, headers in HEADERS.items() for column in headers
    )


def _read_rows(sheet, headers: Sequence[str]):
    actual = tuple(cell.value for cell in sheet[1])
    if actual != tuple(headers):
        raise ValueError(f"unexpected columns in {sheet.title}")
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        rows.append(dict(zip(headers, values, strict=True)))
    return rows


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
