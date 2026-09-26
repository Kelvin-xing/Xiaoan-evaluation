import copy
import json
import pytest

from xiaoan_eval.frozen_review import build_review_rows, import_review_rows
from xiaoan_eval.frozen_calibration import (freeze_split, build_benchmark, judge_inputs,
    snapshot_calibration, record_adoption, compare_calibration, mark_validation_used_for_development)
from xiaoan_eval_core.contracts import digest


def result():
    out = {'schema_version': 'xiaoan-results/v2', 'result_generation': 'g1',
           'plan': {'judges': [{'id': 'j1'}, {'id': 'j2'}]},
           'answers': [{'answer_id': 'a1', 'case_id': 'TC-1', 'turn': 1, 'subject_id': 's1',
                        'question': 'q', 'answer': 'a', 'row_digest': 'r1'}],
           'inventories': [], 'envelopes': [{'answer_id': 'a1', 'rubric': [
               {'judge_id': 'j1', 'status': 'AVAILABLE', 'rubric': {'scores': {'care': 3},
                'dimension_details': [{'score': 3, 'reason': 'clear'}], 'red_lines': []}},
               {'judge_id': 'j2', 'status': 'AVAILABLE', 'rubric': {'scores': {'care': 1},
                'dimension_details': [{'score': 1, 'reason': 'miss'}], 'red_lines': []}}],
               'assessments': [], 'human_review': []}], 'provenance': []}
    from xiaoan_eval_core.results import seal_complete_results
    seal_complete_results(out)
    return out


def filled(source):
    rows = build_review_rows(source)
    rows[0].update(decision='APPROVE', reviewer='owner', reviewed_at='2026-09-24T12:00:00+08:00')
    return rows


def test_full_population_import_immutable_and_new_generation():
    source = result()
    rows = filled(source)
    assert len(rows) == 2
    assert 'RUBRIC_DISAGREEMENT:care' in rows[0]['priority_flags']
    updated = import_review_rows(source, rows, confirmed_by='owner')
    assert source['envelopes'][0]['human_review'] == []
    assert updated['result_generation'] != source['result_generation']
    assert updated['envelopes'][0]['rubric'] == source['envelopes'][0]['rubric']
    assert updated['provenance'][0]['coverage'] == {'planned': 2, 'filled': 1, 'approved': 1}
    with pytest.raises(ValueError, match='foreign'):
        import_review_rows(updated, rows)
    with pytest.raises(ValueError, match='full-population'):
        import_review_rows(source, rows[:1])
    rows[0]['answer'] = 'changed'
    with pytest.raises(ValueError, match='immutable'):
        import_review_rows(source, rows)


def test_reject_gold_requires_explicit_revision_and_human_confirmation():
    source = result()
    rows = filled(source)
    rows[0].update(decision='REJECT', notes='care should be 1', revisions_json=json.dumps([
        {'pointer': '/rubric/rubric/dimension_details/0/score', 'value': 1,
         'reason': 'missed safety', 'evidence': [{'text': 'a'}], 'scope': 'care'}]))
    split = freeze_split([f'TC-{i}' for i in range(1, 34)])
    partition = 'calibration' if 'TC-1' in split['calibration_case_ids'] else 'validation'
    revised = import_review_rows(source, rows)
    assert not build_benchmark(revised, split, partition=partition, purpose='final-validation')['items']
    revised = import_review_rows(source, rows, confirmed_by='owner')
    benchmark = build_benchmark(revised, split, partition=partition, purpose='final-validation')
    assert benchmark['items'][0]['gold']['rubric']['rubric']['dimension_details'][0]['score'] == 1
    assert revised['envelopes'][0]['rubric'][0]['rubric']['dimension_details'][0]['score'] == 3
    candidate = copy.deepcopy(source)
    candidate['envelopes'][0]['rubric'][0]['rubric']['dimension_details'][0]['score'] = 1
    from xiaoan_eval_core.results import seal_complete_results
    seal_complete_results(candidate)
    comparison = compare_calibration(benchmark, source, candidate)
    assert comparison['by_judge']['j1']['improved'] == 1
    candidate['answers'][0]['row_digest'] = 'changed'
    seal_complete_results(candidate)
    with pytest.raises(ValueError, match='identical'):
        compare_calibration(benchmark, source, candidate)


def test_split_and_leakage_guards():
    split = freeze_split([f'TC-{i}' for i in range(33)])
    assert len(split['calibration_case_ids']) == 22
    assert len(split['validation_case_ids']) == 11
    with pytest.raises(ValueError, match='holdout'):
        build_benchmark(result(), split, partition='validation')
    used = mark_validation_used_for_development(split, reason='read labels')
    with pytest.raises(ValueError, match='development'):
        build_benchmark(result(), used, partition='validation', purpose='final-validation')
    source = result()
    source['answers'][0]['human_gold'] = 'secret'
    assert 'human_gold' not in judge_inputs(source, ['TC-1'])[0]


def test_snapshot_and_manual_adoption(tmp_path):
    baseline = tmp_path / 'base'
    candidate = tmp_path / 'new'
    baseline.mkdir(); candidate.mkdir()
    (baseline / 'prompt.md').write_text('old')
    (candidate / 'prompt.md').write_text('new')
    target = tmp_path / 'cal'
    split = freeze_split([str(i) for i in range(33)])
    snapshot_calibration(target, baseline, candidate, split, changes='reason')
    assert not (target / 'adoption.json').exists()
    with pytest.raises(ValueError, match='confirmer'):
        record_adoption(target, confirmed_by='', scope=['j1'], reason='ok', comparison_digest='x')
    receipt = record_adoption(target, confirmed_by='owner', scope=['j1'], reason='ok', comparison_digest='x')
    assert receipt['activation'] == 'EXPLICIT_SEPARATE_ACTION'
    assert (baseline / 'prompt.md').read_text() == 'old'


def test_actual_eight_sheet_review_roundtrip(tmp_path):
    from test_frozen_outputs import fixture_result
    from xiaoan_eval.frozen_export import export_results_workbook
    from xiaoan_eval.frozen_review import read_review_workbook
    from xiaoan_eval_core.results import validate_complete_results
    from openpyxl import load_workbook
    source = fixture_result()
    path = export_results_workbook(source, tmp_path / 'results.xlsx')
    book = load_workbook(path)
    sheet = book['Human Review']
    headers = [c.value for c in sheet[1]]
    for name, value in [('decision', 'APPROVE'), ('reviewer', 'owner'), ('reviewed_at', '2026-09-24T12:00:00+08:00')]:
        sheet.cell(2, headers.index(name) + 1, value)
    book.save(path)
    rows = read_review_workbook(path)
    updated = import_review_rows(source, rows, confirmed_by='owner')
    validate_complete_results(updated)
    assert len(updated['envelopes'][0]['human_review']) == 1
    assert updated['aggregates'] == source['aggregates']
    assert len(rows) == 2
    rows[0]['assessment_digest'] = 'tampered'
    with pytest.raises(ValueError, match='immutable'):
        import_review_rows(source, rows)


def test_calibration_cli_no_holdout_and_no_overwrite(tmp_path):
    import argparse
    from xiaoan_eval.frozen_calibration_cli import register_calibration_commands, run_calibration_command
    parser = argparse.ArgumentParser()
    register_calibration_commands(parser.add_subparsers(dest='command'))
    cases = tmp_path / 'cases.json'
    cases.write_text(json.dumps([f'TC-{i}' for i in range(33)]))
    out = tmp_path / 'split.json'
    args = parser.parse_args(['frozen-calibration', 'split', '--cases', str(cases), '--output', str(out)])
    split = run_calibration_command(args)
    assert len(split['validation_case_ids']) == 11
    with pytest.raises(FileExistsError):
        run_calibration_command(args)
    results = tmp_path / 'results.json'
    results.write_text(json.dumps(result()))
    args = parser.parse_args(['frozen-calibration', 'inputs', '--results', str(results), '--split', str(out),
                              '--partition', 'validation', '--output', str(tmp_path / 'inputs.json')])
    with pytest.raises(ValueError, match='validation answers'):
        run_calibration_command(args)


def test_priority_flags_use_critical_definition_and_keep_disagreements_separate():
    from xiaoan_eval_core.results import seal_complete_results
    source = result()
    source['answers'][0]['requirements'] = [{'id': 'r', 'critical': True}]
    envelope = source['envelopes'][0]
    envelope['assessments'] = [
        {'judge_id': 'j1', 'inventory_id': 'i', 'status': 'AVAILABLE',
         'assessment': {'requirements': [{'id': 'r', 'verdict': 'VIOLATED'}],
                        'claims': [{'id': 'c', 'faithfulness': {'verdict': 'ENTAILED'}}]}},
        {'judge_id': 'j2', 'inventory_id': 'i', 'status': 'AVAILABLE',
         'assessment': {'requirements': [], 'claims': [{'id': 'c', 'faithfulness': {'verdict': 'CONTRADICTED'}}]},
         'requirements': {'release_gate': 'UNDETERMINED'}}]
    seal_complete_results(source)
    flags = json.loads(build_review_rows(source)[0]['priority_flags'])
    assert 'CRITICAL_REQUIREMENT:r' in flags
    assert 'CLAIM_CONFLICT:c:faithfulness' in flags
    assert 'CONTENT_UNCERTAIN' in flags


def test_calibration_comparison_drops_holdout_and_raw_telemetry_and_checks_digest():
    from xiaoan_eval_core.results import seal_complete_results
    source = result()
    source['stages'] = [
        {'answer_id': 'a1', 'task': 'rubric', 'request': {'answer': 'SECRET_PROMPT'}, 'output': 'SECRET_OUTPUT',
         'attempts': [{'usage': {'input_tokens': 3, 'raw_usage': {'sensitive': 'SECRET_USAGE'}}, 'latency_ms': 1}]},
        {'answer_id': 'holdout', 'task': 'rubric', 'request': {'answer': 'SECRET_HOLDOUT'},
         'attempts': [{'usage': {'input_tokens': 999}}]}]
    seal_complete_results(source)
    reviewed = import_review_rows(source, filled(source), confirmed_by='owner')
    split = freeze_split([f'TC-{i}' for i in range(1,34)])
    partition = 'calibration' if 'TC-1' in split['calibration_case_ids'] else 'validation'
    benchmark = build_benchmark(reviewed, split, partition=partition, purpose='final-validation')
    comparison = compare_calibration(benchmark, source, source)
    encoded = json.dumps(comparison)
    assert 'SECRET' not in encoded and '999' not in encoded
    assert comparison['usage']['baseline'] == [{'answer_id':'a1','task':'rubric','usage':{'input_tokens':3,'latency_ms':1}}]
    candidate = copy.deepcopy(source)
    candidate['envelopes'][0]['rubric'][0]['rubric']['scores']['care'] = 0
    with pytest.raises(ValueError, match='digest mismatch'):
        compare_calibration(benchmark, source, candidate)
    with pytest.raises(ValueError, match='digest mismatch'):
        build_benchmark(candidate, split)


def test_calibration_labels_use_ids_not_array_order():
    from xiaoan_eval.frozen_calibration import _labels
    value = {'claims':[{'id':'c1','faithfulness':{'verdict':'ENTAILED'}}, {'id':'c2','faithfulness':{'verdict':'UNKNOWN'}}],
             'dimension_details':[{'module':'care','score':1},{'module':'action','score':2}]}
    reordered = {key: list(reversed(items)) for key, items in value.items()}
    assert _labels(value) == _labels(reordered)


def test_human_rubric_revision_rejects_fractional_score():
    source = result()
    rows = filled(source)
    rows[0].update(decision='REJECT',notes='score',revisions_json=json.dumps([
        {'pointer':'/rubric/rubric/dimension_details/0/score','value':1.5,'reason':'reason','evidence':['a'],'scope':'care'}]))
    with pytest.raises(ValueError,match='score'):
        import_review_rows(source, rows, confirmed_by='owner')
