"""Offline refresh of derived results and the multi-model workbook projection."""
from copy import deepcopy
import json

from openpyxl import load_workbook
import pytest

from xiaoan_eval.cli import main
from xiaoan_eval_core.results import aggregate_complete_results, seal_complete_results, validate_complete_results
from test_frozen_outputs import fixture_result


def old_multimodel_result():
    result = fixture_result()
    first = result['answers'][0]
    first['oracle_source'] = {'route_ids': ['baseline'], 'preferred_route_id': 'baseline',
                              'reference_oracle': {'status': 'approved', 'scope': ['route_ids', 'preferred_route_id']}}
    first['trace'] = {'route': {'capsule_id': 'baseline'}}
    second = deepcopy(first)
    second.update(answer_id='b', subject_id='other-subject', trace={'route': {'capsule_id': 'crisis_sop'}})
    result['answers'].append(second)
    inventory = deepcopy(result['inventories'][0])
    inventory.update(answer_id='b', inventory_id='other-inventory')
    result['inventories'].append(inventory)
    envelope = deepcopy(result['envelopes'][0])
    envelope.update(answer_id='b', inventory_id='other-inventory')
    for cell in envelope['rubric'] + envelope['assessments']:
        cell['answer_id'] = 'b'
        if 'inventory_id' in cell:
            cell['inventory_id'] = 'other-inventory'
    result['envelopes'].append(envelope)
    result['stages'] = [{'task': 'rubric', 'answer_id': 'b', 'identity': {'id': 'j2'},
                         'attempts': [{'attempt_id': 'attempt-1', 'latency_ms': 10,
                                       'usage': {'input_tokens': 3, 'output_tokens': 2, 'total_tokens': 5}}]}]
    result['aggregates'] = aggregate_complete_results(result)
    result['aggregates'].pop('routing')
    seal_complete_results(result)
    return result


def test_refresh_preserves_frozen_evidence_and_adds_routing_and_identities(tmp_path):
    parent = old_multimodel_result()
    source = tmp_path/'source'/'results.json'
    source.parent.mkdir()
    source.write_text(json.dumps(parent, ensure_ascii=False))
    original = source.read_bytes()
    with pytest.raises(ValueError, match='refresh-derived-results'):
        main(['export-human-review', str(source), '--output', str(tmp_path/'bad.xlsx')])
    out = tmp_path/'refresh'
    assert main(['refresh-derived-results', str(source), '--output', str(out)]) == 0
    assert source.read_bytes() == original
    result = validate_complete_results(json.loads((out/'results.json').read_text()))
    assert result['answers'] == parent['answers']
    assert result['envelopes'] == parent['envelopes']
    assert result['stages'] == parent['stages']
    assert result['provenance'][-1]['parent_generation'] == parent['result_generation']
    assert result['result_generation'] != parent['result_generation']
    assert [s['subject_id'] for s in result['aggregates']['routing']['summary']] == ['other-subject', 's']

    book = load_workbook(out/'results.xlsx', read_only=True, data_only=True)
    overview = list(book['Overview'].values)
    assert overview[0][0].startswith('TC-01｜1 案例｜2 Subject × 2 Judge')
    rubric = next(index for index, row in enumerate(overview) if row[:3] == ('總體', '全部案例', 'rubric'))
    assert overview[rubric + 1][:3] == ('Subject', 'j1', 'j2')
    assert overview[rubric + 2][0] == 'other-subject'
    assert overview[rubric + 2][1:3] == (1, 3)
    score = list(book['Score Summary'].values)
    assert score[3][:5] == ('Subject', 'Judge', '軸', '細分', '指標')
    assert ('other-subject', 'j2', '總體', '全部案例', 'rubric') in [row[:5] for row in score]
    routing = list(book['Routing Summary'].values)
    assert routing[4][0] == 'other-subject' and routing[5][0] == 's'
    assert any(row[:2] == ('other-subject', '基礎回應') for row in routing)
    coverage = list(book['Coverage & Usage'].values)
    assert coverage[3][:4] == ('Subject', 'Judge', '指標', '可納入案例')
    assert not any('TC-01' in row for row in coverage)
    assert any(row[:5] == ('other-subject', 'j2', 'rubric', 1, '1/1') for row in coverage)
    book.close()


def test_refresh_rejects_same_directory_and_existing_generation(tmp_path):
    source = tmp_path/'results.json'
    source.write_text(json.dumps(old_multimodel_result()))
    with pytest.raises(ValueError, match='new output directory'):
        main(['refresh-derived-results', str(source), '--output', str(tmp_path)])
    out = tmp_path/'new'
    main(['refresh-derived-results', str(source), '--output', str(out)])
    with pytest.raises(ValueError, match='new output directory'):
        main(['refresh-derived-results', str(source), '--output', str(out)])


def test_refresh_attaches_versioned_case_taxonomy_without_changing_source(tmp_path):
    source = tmp_path/'results.json'
    source.write_text(json.dumps(old_multimodel_result(), ensure_ascii=False))
    original = source.read_bytes()
    from pathlib import Path
    taxonomy = Path(__file__).resolve().parents[1]/'test-cases/scenario-taxonomy.json'
    out = tmp_path/'classified'
    assert main(['refresh-derived-results', str(source), '--output', str(out), '--case-taxonomy', str(taxonomy)]) == 0
    result = validate_complete_results(json.loads((out/'results.json').read_text()))
    assert source.read_bytes() == original
    assert all(answer['scenario_category'] == 'immediate_safety' for answer in result['answers'])
    assert result['provenance'][-2]['operation'] == 'classify_cases'
    assert any(row['scope'] == 'scenario_tags:imminent_threat' for row in result['aggregates']['metrics'])
    book = load_workbook(out/'results.xlsx', read_only=True)
    assert ('test_type', 'emergency', 'rubric') in [row[:3] for row in book['Overview'].values]
    book.close()
