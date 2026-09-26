from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import asyncio
import pytest
from xiaoan_eval_core.costs import build_answer_costs,load_catalog
from xiaoan_eval.frozen_provider import usage


def answer(inp=1000,out=200,**extra):
    return {'answer_id':'a','case_id':'TC-35','turn':1,'subject_id':'s','subject_model':'gpt-5.6-luna','subject_provider':'gpt',
            'usage':{'input_tokens':inp,'output_tokens':out,**extra}}


def test_basic_formula_and_input_cache_not_double_counted():
    r=build_answer_costs([answer()])['rows'][0]
    assert r['input_cost']==pytest.approx(.0002)
    assert r['output_cost']==pytest.approx(.00024)
    r=build_answer_costs([answer(cached_input_tokens=600)])['rows'][0]
    assert r['total_cost']==pytest.approx((400*.2+600*.02+200*1.2)/1e6)


def test_long_context_threshold_applies_to_whole_request():
    r=build_answer_costs([answer(272001,10)])['rows'][0]
    assert r['context_tier']=='long'
    assert r['total_cost']==pytest.approx((272001*.4+10*1.8)/1e6)


def test_missing_zero_invalid_unknown_model_and_no_call():
    a=answer();b={**answer(None,200),'answer_id':'b'}
    summary=build_answer_costs([a,a,b])['summary'][0]
    assert summary['priced_answers']==1 and summary['planned_answers']==2 and summary['total_cost'] is None
    assert build_answer_costs([answer(0,0)])['rows'][0]['total_cost']==0
    assert build_answer_costs([answer(-1,0)])['rows'][0]['status']=='UNAVAILABLE'
    assert build_answer_costs([{**a,'subject_model':'unknown'}])['rows'][0]['reason']=='OFFICIAL_PRICE_NOT_FOUND'
    assert build_answer_costs([answer(None,None,no_model_call=True)])['rows'][0]['status']=='NOT_APPLICABLE'
    assert build_answer_costs([answer(cached_input_tokens=1001)])['rows'][0]['reason']=='CACHE_TOKENS_EXCEED_INPUT'


def test_anthropic_input_normalization_and_write_duration():
    raw={'input_tokens':100,'output_tokens':10,'cache_read_input_tokens':200,'cache_creation_input_tokens':50,
         'cache_creation':{'ephemeral_5m_input_tokens':20,'ephemeral_1h_input_tokens':30}}
    tokens=usage({'usage':raw})
    assert tokens['input_tokens']==350 and tokens['cache_write_tokens']==50
    a={**answer(),'subject_provider':'claude','subject_model':'claude-sonnet-5','usage':tokens}
    r=build_answer_costs([a])['rows'][0]
    assert r['total_cost']==pytest.approx((100*2+200*.2+20*2.5+30*4+10*10)/1e6)


def test_deepseek_needs_explicit_period():
    a={**answer(),'subject_provider':'deepseek','subject_model':'deepseek-flash'}
    assert build_answer_costs([a])['rows'][0]['reason']=='PRICING_PERIOD_REQUIRED'
    a['usage']['pricing_period']='peak'
    assert build_answer_costs([a])['rows'][0]['total_cost']==pytest.approx(.00054)


def test_stream_preserves_usage_and_response_id():
    from xiaoan_eval_core.llm_adapter import astream
    class Stream:
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        def __aiter__(self):return self.events()
        async def events(self):yield SimpleNamespace(type='response.output_text.delta',delta='answer')
        async def get_final_response(self):return SimpleNamespace(id='r',usage={'input_tokens':40,'output_tokens':5})
    client=SimpleNamespace(responses=SimpleNamespace(stream=lambda **kwargs:Stream()))
    captured=[]
    async def collect():return [x async for x in astream(client,{'model':'gpt-test','input':'question'},{},completion_callback=captured.append)]
    assert asyncio.run(collect())==['answer']
    assert captured[0]['usage']['input_tokens']==40 and captured[0]['id']=='r'


def test_local_composer_usage_reaches_frozen_answer(monkeypatch):
    from xiaoan_eval.frozen_generation import LocalChatflowSubject
    from xiaoan_eval.frozen_ingress import build_plan,freeze_answer
    subject=LocalChatflowSubject()
    import chat_service,openai_compat
    from context_snapshot import record_request,snapshot
    import json
    raw={'input_tokens':1000,'output_tokens':200,'input_tokens_details':{'cached_tokens':600}}
    monkeypatch.setattr(openai_compat,'completion_usage',lambda:raw)
    def turn(message,capsules,*,state,debug_callback,**kwargs):
        record_request('composer',{'model':'gpt-5.6-luna','instructions':'prompt','input':json.dumps({'user_message':message})})
        debug_callback({'effective_context_snapshot':snapshot(1)})
        return 'answer',state
    monkeypatch.setattr(chat_service,'run_turn',turn)
    identity={'id':'s','model':'gpt-5.6-luna','provider':'gpt','prompt_version':'v1'}
    row=build_plan([identity],[identity],identity,case_ids=['TC-35'])['rows'][0]
    response,_=subject.turn(identity,row,None,[])
    frozen=freeze_answer(row,response,[])
    assert frozen['usage']['input_tokens']==1000
    assert build_answer_costs([frozen])['rows'][0]['total_cost']==pytest.approx(.000332)


def test_cost_cli_creates_new_result_and_eight_sheets(tmp_path):
    import json
    from openpyxl import load_workbook
    from test_frozen_outputs import fixture_result
    from xiaoan_eval_core.results import seal_complete_results
    from xiaoan_eval.cli import main
    r=fixture_result();r['answers'][0].update(subject_provider='gpt',subject_model='gpt-5.6-luna',usage=answer()['usage'])
    seal_complete_results(r);source=tmp_path/'source.json';source.write_text(json.dumps(r))
    assert main(['cost',str(source),'--output',str(tmp_path/'priced')])==0
    updated=json.loads((tmp_path/'priced/results.json').read_text())
    assert updated['aggregates']['answer_costs']['rows'][0]['total_cost']==pytest.approx(.00044)
    assert json.loads(source.read_text())==r
    wb=load_workbook(tmp_path/'priced/results.xlsx',read_only=True)
    assert len(wb.sheetnames)==8 and 'official_answer_cost' in [c.value for c in wb['Answers'][1]]
