"""Regression examples from the 2026-09-13 measurement audit."""
import pytest

from test_pipeline import FakeRunner, RULE, CONFIG, _case, _judge_payload
from xiaoan_eval.pipeline import EvaluationPipeline
from xiaoan_eval.judge_client import JudgeClient
from xiaoan_eval.runner import CaseRunResult, TurnRunResult
from xiaoan_eval.report_model import build_report_model


def pipeline(runner=None, threshold=None):
    return EvaluationPipeline(
        runner or FakeRunner(), RULE, CONFIG,
        primary_judge=JudgeClient(lambda _: _judge_payload(), RULE),
        egress_validator=lambda _: True,
        authoritative_context_provider=lambda *_: {"refs": []},
        quality_threshold=threshold,
    )


class FailedRunner:
    def run_cases(self, cases):
        return [CaseRunResult("TC-01", "error", "c", [TurnRunResult(1, None, {}, "timeout")])]


def test_failed_execution_has_no_quality_value_in_report():
    failed = pipeline(FailedRunner()).evaluate_case(_case())
    assert failed["quality"]["weighted_total"] is None
    assert failed["quality"]["status"] == "UNAVAILABLE"
    good = pipeline().evaluate_case(_case())
    good["case_id"] = "TC-02"
    report = build_report_model([failed, good])
    overall = next(row for row in report.overview if row["metric"] == "Overall score")
    assert overall["value"] == pytest.approx(2)
    verdict = next(row for row in report.overview if row["metric"] == "Evaluation verdict")
    assert verdict["value"] == "UNAVAILABLE"


def test_quality_threshold_is_a_verdict_not_only_a_second_judge_trigger():
    record = pipeline(threshold=2.5).evaluate_case(_case())
    assert record["status"] == "FAIL"
    assert record["quality"]["verdict"] == "FAIL"
    assert record["quality"]["weighted_total"] == pytest.approx(2)
    overview = build_report_model([record]).overview
    assert next(row for row in overview if row["metric"] == "Execution gate pass rate")["value"] == 1


def test_unconfigured_threshold_does_not_claim_quality_pass():
    record = pipeline().evaluate_case(_case())
    assert record["quality"]["verdict"] == "NOT_CONFIGURED"
    overview = build_report_model([record]).overview
    assert next(row for row in overview if row["metric"] == "Evaluation verdict")["value"] == "NOT_CONFIGURED"


def test_empty_oracles_and_unsupported_required_claims_have_distinct_denominators():
    from xiaoan_eval.v3_metrics import summarize_v3
    summary = summarize_v3([{"pipeline": {"observations": [{
        "oracle_approved": True,
        "expected": {"route_ids": [], "safety_levels": [], "response_oracle": {"required_claims": ["fact"]}},
        "actual": {"route_id": "baseline", "safety_level": "baseline"},
        "judge": {"faithfulness_claims": [{"claim": "fact", "supported": False}]},
    }]}}])
    assert summary["route"]["accepted_accuracy"] is None
    assert summary["safety"]["accepted_accuracy"] is None
    assert summary["claims"]["unsupported_claim_rate"] == 1


@pytest.mark.parametrize("kind, state, status", [
    ("not_use", {"memory_facts": ["fact"], "memory_used": False}, "pass"),
    ("use", {}, "skip"), ("isolation", {}, "skip"),
    ("retrieve", {"memory_retrieved_facts": []}, "fail"),
    ("update", {"memory_updated_correctly": True}, "pass"),
    ("isolation", {"memory_isolated": True, "memory_contamination_candidates": ["other-session"]}, "fail"),
])
def test_memory_lifecycle_respects_type_and_missing_evidence(kind, state, status):
    from xiaoan_eval.cases import MemoryCheckpoint
    from xiaoan_eval.metrics import evaluate_memory_checkpoint
    result = evaluate_memory_checkpoint(MemoryCheckpoint(1, ("fact",), "synthetic", kind), {"state": state})
    assert result.status.value == status
    if status == "skip":
        assert result.score is None


def test_suite_and_case_use_all_dimensions_with_dynamic_focus():
    from xiaoan_eval.scoring import score_case, score_case_fact, TurnDimensionFact, TurnQuality
    scores = {module.name: 3 if module.name == "行动赋权" else 0 for module in RULE.modules}
    ordinary = score_case(RULE, [TurnQuality(scores)], ["行动赋权"])
    suite = score_case_fact(case_id="TC-01", expected_turns=[1], quality_focus=["行动赋权"],
        dimension_weights={module.name: module.weight for module in RULE.modules},
        turn_facts=[TurnDimensionFact("TC-01", 1, name, "AVAILABLE", score) for name, score in scores.items()])
    assert ordinary.weighted_total == pytest.approx(0.743119266055046)
    assert suite.value == pytest.approx(ordinary.weighted_total)


def test_agreement_does_not_impute_provider_failures_or_claim_constant_alpha():
    from xiaoan_eval.methodology_metrics import agreement_report, red_line_agreement_report
    rows = [dict(case_id="TC", turn=1, subject={"id": s}, judge={"id": j},
        status="PASS", scores={"q": 3}, self_judging=False,
        red_line_evidence={"RL": []}, triggered_red_lines=[])
        for s in ("a", "b", "c") for j in ("j1", "j2")]
    assert agreement_report(rows, "q")["krippendorff_alpha_ordinal"] is None
    assert red_line_agreement_report(rows, "RL")["krippendorff_alpha_nominal"] is None
    for i, row in enumerate(rows):
        row["scores"]["q"] = i // 2
    rows[-1]["status"] = "UNAVAILABLE"
    report = agreement_report(rows, "q")
    assert report["kendall_w"] is None
    assert report["kendall_strata"][0]["missing_policy"] == "NO_IMPUTATION"


@pytest.mark.parametrize("score, expected, verdict", [(1, 1, "FAIL"), (3, 3, "PASS"), ("focus", 0.743119266055046, "FAIL")])
def test_human_review_preserves_dynamic_weights_and_recomputes_verdict(tmp_path, score, expected, verdict):
    from openpyxl import load_workbook
    from xiaoan_eval.workbook import write_workbook, read_workbook
    from xiaoan_eval.review_workbook import export_review_workbook, import_review_workbook
    record = pipeline(threshold=2).evaluate_case(_case())
    record["review"].update(status="NEEDS_REVIEW", per_turn=["NEEDS_REVIEW"])
    record["status"] = "NEEDS_REVIEW"
    original = tmp_path / "original.xlsx"
    write_workbook(build_report_model([record]), original)
    packet = tmp_path / "review.xlsx"
    export_review_workbook(original, packet, RULE)
    wb = load_workbook(packet)
    queue = wb["01_Review_Queue"]
    queue["H2"] = "synthetic-reviewer"
    queue["I2"] = "2026-09-13T00:00:00+00:00"
    queue["J2"] = "high"
    items = wb["02_Review_Items"]
    for row in items.iter_rows(min_row=2):
        if row[1].value == "dimension":
            row[5].value = (3 if row[2].value == "行动赋权" else 0) if score == "focus" else score
            row[8].value = "synthetic evidence-based assessment"
        else:
            row[6].value = False
    wb.save(packet)
    wb.close()
    reviewed = import_review_workbook(packet, original, RULE)
    assert reviewed.cases[0]["final_score"] == pytest.approx(expected)
    assert reviewed.cases[0]["quality_verdict"] == verdict
    assert next(row for row in reviewed.overview if row["metric"] == "Evaluation verdict")["value"] == verdict
    final = tmp_path / "final.xlsx"
    write_workbook(reviewed, final)
    assert read_workbook(final).cases[0]["quality_verdict"] == verdict


def test_proposed_cases_have_authored_but_no_reviewed_coverage():
    from xiaoan_eval.cases import load_cases
    from xiaoan_eval.coverage import case_coverage, coverage_rows
    loaded = load_cases("test-cases/proposed", RULE)
    assert len(loaded) == 9
    assert all(item.case is not None and not item.case.oracle_gate_eligible for item in loaded)
    rows = coverage_rows([case_coverage(item.case) for item in loaded])
    assert all(row["value"] == 0 for row in rows)
    assert any("Authored=1;" in row["interpretation"] for row in rows if row["metric"] == "Oracle memory:isolation")


def test_weighted_automatic_turn_and_unreviewed_case_survive_review_import(tmp_path):
    from xiaoan_eval.workbook import write_workbook, read_workbook
    from xiaoan_eval.review_workbook import _apply_submissions
    record = pipeline(threshold=0.6).evaluate_case(_case())
    scores = {module.name: 3 if module.name == '行动赋权' else 0 for module in RULE.modules}
    for item in record['review']['judge_audit'][0]['primary']['dimensions']:
        item['score'] = scores[item['module']]
    record['quality'].update(weighted_total=0.743119266055046, dimensions=scores)
    model = build_report_model([record])
    assert model.turns[0]['automatic_score'] == pytest.approx(0.743119266055046)
    path = tmp_path / 'weighted.xlsx'
    write_workbook(model, path)
    rebuilt = _apply_submissions(read_workbook(path), [], RULE)
    assert rebuilt.cases[0]['final_score'] == pytest.approx(0.743119266055046)
    assert rebuilt.cases[0]['final_source'] == 'automatic'
    assert rebuilt.cases[0]['quality_verdict'] == 'PASS'


@pytest.mark.parametrize('red_line', [False, True])
def test_available_zero_remains_a_numeric_automatic_turn_score(red_line):
    record = pipeline().evaluate_case(_case())
    record['quality']['weighted_total'] = 0.0
    primary = record['review']['judge_audit'][0]['primary']
    for item in primary['dimensions']:
        item['score'] = 2 if red_line else 0
    if red_line:
        primary['red_lines'][0]['triggered'] = True
    assert build_report_model([record]).turns[0]['automatic_score'] == 0.0
