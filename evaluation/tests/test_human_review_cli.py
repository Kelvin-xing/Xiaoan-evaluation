"""Full-review CLI binding and no automatic rescoring."""
import json
from openpyxl import load_workbook
import pytest
from xiaoan_eval.cli import main
from xiaoan_eval.frozen_export import SHEETS
from xiaoan_eval_core.results import validate_complete_results
from test_frozen_outputs import fixture_result


def packet(tmp_path):
    source=tmp_path/'results.json';source.write_text(json.dumps(fixture_result(),ensure_ascii=False))
    workbook=tmp_path/'review.xlsx'
    assert main(['export-human-review',str(source),'--output',str(workbook)])==0
    return source,workbook


def edit(workbook,field,value):
    book=load_workbook(workbook);sheet=book['Human Review'];headers=[c.value for c in sheet[1]]
    sheet.cell(2,headers.index(field)+1).value=value;book.save(workbook)


def test_all_rows_exported_and_approval_preserves_scores(tmp_path):
    source,workbook=packet(tmp_path);book=load_workbook(workbook)
    assert tuple(book.sheetnames)==SHEETS
    sheet=book['Human Review'];headers=[c.value for c in sheet[1]]
    for row in range(2,sheet.max_row+1):
        for field,value in [('decision','APPROVE'),('reviewer','user'),('reviewed_at','2026-09-24T10:00:00+08:00')]:
            sheet.cell(row,headers.index(field)+1).value=value
    book.save(workbook);before=json.loads(source.read_text());out=tmp_path/'reviewed'
    assert main(['import-human-review',str(source),str(workbook),'--output',str(out),'--confirmed-by','user'])==0
    after=json.loads((out/'results.json').read_text());validate_complete_results(after)
    assert before['result_generation']!=after['result_generation']
    assert before['aggregates']==after['aggregates']


def test_immutable_edit_rejected(tmp_path):
    source,workbook=packet(tmp_path);edit(workbook,'answer','tampered')
    with pytest.raises(ValueError):main(['import-human-review',str(source),str(workbook),'--output',str(tmp_path/'out')])


def test_review_input_not_overwritten(tmp_path):
    source,workbook=packet(tmp_path);edit(workbook,'notes','keep this draft')
    with pytest.raises(ValueError,match='human edits'):
        main(['export-human-review',str(source),'--output',str(workbook)])
