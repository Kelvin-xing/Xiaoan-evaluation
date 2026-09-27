from copy import deepcopy

import pytest
from openpyxl import load_workbook

from test_frozen_end_to_end import FixtureProvider, IDENTITY
from xiaoan_eval.frozen_balanced import assemble_balanced
from xiaoan_eval.frozen_cli import execute_frozen
from xiaoan_eval.frozen_ingress import build_plan, freeze_answer
from xiaoan_eval_core.results import seal_complete_results, validate_complete_results


def cohorts(tmp_path):
    judges = [{**IDENTITY, 'id': f'judge-{n}'} for n in range(4)]
    results = []
    paths = []
    for n in range(8):
        subject = {**IDENTITY, 'id': f'subject-{n}'}
        spec = build_plan([subject], judges, IDENTITY, case_ids=['TC-35', 'TC-22'], relevancy_generator=IDENTITY)
        spec['rows'] = [freeze_answer(row, {'text': '這是離線合成回答。'}, [], generation_id=subject['id'])
                        for row in spec['rows']]
        path = tmp_path / f'subject-{n}'
        results.append(execute_frozen(spec, path, provider=FixtureProvider()))
        paths.append(path / 'results.json')
    return results, paths


def test_balanced_assembly_preserves_original_ids_and_filters_whole_cases(tmp_path):
    results, paths = cohorts(tmp_path)
    source = results[0]
    missing = next(row for row in source['answers'] if row['case_id'] == 'TC-22')
    env = next(env for env in source['envelopes'] if env['answer_id'] == missing['answer_id'])
    env['rubric'][0]['status'] = 'UNAVAILABLE'
    seal_complete_results(source)
    assembled = assemble_balanced(results, paths, tmp_path / 'balanced')
    validate_complete_results(assembled)
    assert assembled['plan']['case_ids'] == ['TC-35']
    assert len(assembled['answers']) == 16
    assert len(assembled['envelopes']) == 16
    assert len({(r['subject_id'], cell['judge_id']) for r in assembled['answers']
                for cell in next(e for e in assembled['envelopes'] if e['answer_id'] == r['answer_id'])['rubric']}) == 32
    assert {row['answer_id'] for row in assembled['answers']} <= {
        row['answer_id'] for source in results for row in source['answers']}
    assert assembled['manifest']['source_manifest_digests'] == [r['manifest']['manifest_digest'] for r in results]
    assert (tmp_path / 'balanced' / 'results.xlsx').exists()


def test_balanced_assembly_rejects_incomplete_matrix(tmp_path):
    results, paths = cohorts(tmp_path)
    for env in results[0]['envelopes']:
        env['assessments'][0]['status'] = 'UNAVAILABLE'
    seal_complete_results(results[0])
    with pytest.raises(ValueError, match='No case'):
        assemble_balanced(results, paths, tmp_path / 'none')
    assert not (tmp_path / 'none' / 'results.json').exists()


def test_rubric_comparison_uses_metric_specific_common_denominators(tmp_path):
    results, paths = cohorts(tmp_path)
    target = next(e for e in results[0]['envelopes'] if e['answer_id'] ==
                  next(r['answer_id'] for r in results[0]['answers'] if r['case_id'] == 'TC-22'))
    target['assessments'][0]['status'] = 'UNAVAILABLE'
    seal_complete_results(results[0])
    assembled = assemble_balanced(results, paths, tmp_path / 'rubric', require_assessment=False)
    assert assembled['plan']['case_ids'] == ['TC-22', 'TC-35'] or assembled['plan']['case_ids'] == ['TC-35', 'TC-22']
    by_metric = {name: [row for row in assembled['aggregates']['balanced_comparison'] if row['metric'] == name]
                 for name in ('rubric', 'faithfulness')}
    assert {row['effective_cases'] for row in by_metric['rubric']} == {2}
    assert {row['effective_cases'] for row in by_metric['faithfulness']} == {1}
    wb = load_workbook(tmp_path / 'rubric' / 'results.xlsx', read_only=True)
    assert 'Balanced Comparison' in wb.sheetnames


def test_answer_balanced_keeps_incomplete_judge_results_without_inventing_scores(tmp_path):
    results, paths = cohorts(tmp_path)
    target = next(env for env in results[0]['envelopes'] if env['answer_id'] == results[0]['answers'][0]['answer_id'])
    target['rubric'][0]['status'] = 'UNAVAILABLE'
    target['assessments'][0]['status'] = 'UNAVAILABLE'
    seal_complete_results(results[0])
    assembled = assemble_balanced(results, paths, tmp_path / 'answer-balanced',
                                  require_relevancy=False, require_assessment=False, require_rubric=False)
    assert len(assembled['answers']) == 8 * 4
    assert len(assembled['envelopes']) == 8 * 4
    retained = next(env for env in assembled['envelopes'] if env['answer_id'] == target['answer_id'])
    assert retained['rubric'][0]['status'] == 'UNAVAILABLE'
    assert retained['assessments'][0]['status'] == 'UNAVAILABLE'
    assert assembled['provenance'][-1]['case_selection'] == 'whole_case_answer_balanced'


def test_balanced_assembly_rejects_different_judge_contract(tmp_path):
    results, paths = cohorts(tmp_path)
    altered = deepcopy(results[1])
    altered['plan']['judges'][0]['prompt_version'] = 'other'
    seal_complete_results(altered)
    results[1] = altered
    with pytest.raises(ValueError, match='Judge identities'):
        assemble_balanced(results, paths, tmp_path / 'mismatch')


def test_subject_can_extend_matrix_with_disjoint_case_shards(tmp_path):
    results, paths = cohorts(tmp_path)
    judges = results[7]['plan']['judges']
    subject = results[7]['plan']['subjects'][0]
    results.pop()
    paths.pop()
    for case in ('TC-22', 'TC-35'):
        spec = build_plan([subject], judges, IDENTITY, case_ids=[case], relevancy_generator=IDENTITY)
        spec['rows'] = [freeze_answer(row, {'text': '這是離線合成回答。'}, [], generation_id=case)
                        for row in spec['rows']]
        path = tmp_path / ('shard-' + case)
        results.append(execute_frozen(spec, path, provider=FixtureProvider()))
        paths.append(path / 'results.json')
    assembled = assemble_balanced(results, paths, tmp_path / 'sharded')
    assert set(assembled['plan']['case_ids']) == {'TC-22', 'TC-35'}
    assert len(assembled['answers']) == 8 * 4
    with pytest.raises(ValueError, match='Overlapping subject case shards'):
        assemble_balanced([*results, results[-1]], [*paths, paths[-1]], tmp_path / 'overlap')
