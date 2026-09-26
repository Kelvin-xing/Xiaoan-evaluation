import copy
import json
import pytest
from xiaoan_eval_core.results import build_complete_results, seal_complete_results, validate_complete_results
from xiaoan_eval.frozen_export import export_results_workbook
from openpyxl import load_workbook


def fixture_result():
    rows = [{"answer_id": "a", "subject_id": "s", "case_id": "TC-01", "turn": 1, "question": "問題", "answer": "完整回答", "status": "AVAILABLE", "requirements": [], "context": []}]
    inv = {"answer_id": "a", "inventory_id": "i", "claims": [{"id": "c", "proposition": "完整主張", "kind": "FACTUAL", "conditions": [], "answer_span": {"start": 0, "end": 4, "text": "完整回答"}}]}
    cells = [{"answer_id": "a", "judge_id": j, "status": "AVAILABLE", "inventory_id": "i", "assessment": {"claims": [{"id": "c", "faithfulness": {"verdict": "UNKNOWN", "evidence": [], "reason": "缺少證據"}, "correctness": {"verdict": "UNKNOWN", "evidence": [], "reason": "缺少真值"}}], "requirements": []}} for j in ("j1", "j2")]
    rubrics = [{"answer_id": "a", "judge_id": j, "status": "AVAILABLE", "rubric": {"status": "AVAILABLE", "scores": {"同理": n}, "weighted_total": n, "final_weights": {"同理": 1}, "gate": "PASS", "red_lines": [], "dimension_details": [{"module": "同理", "score": n, "reason": "明確回應感受"}]}} for j, n in (("j1", 1), ("j2", 3))]
    return build_complete_results({"judges": [{"id": "j1"}, {"id": "j2"}]}, rows, {"inventories": [inv], "cells": cells}, rubrics, [], [], [])


def test_joins_judges_and_unknown_denominator(tmp_path):
    result = fixture_result()
    metric = [m for m in result["aggregates"]["metrics"] if m["metric"] == "rubric" and m["scope"] == "own_complete_cases"]
    assert [m["value"] for m in metric] == [1, 3]
    claims = [m for m in result["aggregates"]["metrics"] if m["metric"] == "faithfulness"]
    assert all(m["value"] == 0 and m["unknown_ratio"] == 1 for m in claims)
    output = export_results_workbook(result, tmp_path / "results.xlsx")
    book = load_workbook(output)
    assert tuple(book.sheetnames) == (
        "Overview", "Score Summary", "Routing Summary", "Coverage & Usage",
        "Case Eligibility", "Spec", "Answers", "Scores",
        "Claims", "Requirements", "Rating Details", "Human Review",
    )
    assert not any(
        "路由" in str(cell.value)
        for row in book["Score Summary"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    score_headers = [cell.value for cell in book["Scores"][1]]
    assert "call_telemetry" not in score_headers
    assert "metrics" not in score_headers
    assert "requirements" not in score_headers
    assert {"attempt_count", "latency_ms", "input_tokens", "output_tokens"}.issubset(score_headers)
    vals = list(book["Claims"].values)
    assert vals[1][vals[0].index("claim")] == "完整主張"
    assert vals[1][vals[0].index("answer_quote")] == "完整回答"
    review = book["Human Review"]
    headers = [c.value for c in review[1]]
    review.cell(2, headers.index("decision") + 1, "APPROVE")
    book.save(output)
    with pytest.raises(ValueError, match="human edits"):
        export_results_workbook(result, output)


def test_integrity_rejects_broken_join_and_span():
    result = fixture_result()
    result["envelopes"][0]["inventory_id"] = "bad"
    seal_complete_results(result)
    with pytest.raises(ValueError, match="join"):
        validate_complete_results(result)
    result = fixture_result()
    result["inventories"][0]["claims"][0]["answer_span"]["text"] = "錯誤"
    seal_complete_results(result)
    with pytest.raises(ValueError, match="span"):
        validate_complete_results(result)


def test_partial_dimensions_do_not_contaminate_other_metric():
    result = fixture_result()
    cell = result["envelopes"][0]["assessments"][0]
    cell["status"] = "PARTIAL"
    cell["assessment"]["dimension_status"] = {"faithfulness": "UNAVAILABLE", "correctness": "AVAILABLE"}
    from xiaoan_eval_core.results import aggregate_complete_results
    metrics = aggregate_complete_results(result)["metrics"]
    own = {m["metric"]: m for m in metrics if m["judge_id"] == "j1" and m["scope"] == "own_complete_cases"}
    assert own["faithfulness"]["value"] is None
    assert own["correctness"]["value"] == 0


def test_scenario_tags_overlap_without_double_counting_overall(tmp_path):
    from xiaoan_eval_core.results import aggregate_complete_results

    result = fixture_result()
    result["answers"][0].update(test_type="emergency", scenario_category="immediate_safety",
                                scenario_tags=["imminent_threat", "safety_planning"])
    second = copy.deepcopy(result["answers"][0])
    second.update(answer_id="b", case_id="TC-02", scenario_tags=["safety_planning"])
    result["answers"].append(second)
    inventory = copy.deepcopy(result["inventories"][0])
    inventory.update(answer_id="b", inventory_id="i-b")
    result["inventories"].append(inventory)
    envelope = copy.deepcopy(result["envelopes"][0])
    envelope["answer_id"] = "b"
    envelope["inventory_id"] = "i-b"
    for cell in envelope["rubric"] + envelope["assessments"]:
        cell["answer_id"] = "b"
        if "inventory_id" in cell:
            cell["inventory_id"] = "i-b"
    result["envelopes"].append(envelope)
    result["aggregates"] = aggregate_complete_results(result)
    seal_complete_results(result)

    rubric = {row["scope"]: row for row in result["aggregates"]["metrics"]
              if row["metric"] == "rubric" and row["judge_id"] == "j1"}
    assert rubric["own_complete_cases"]["planned_cases"] == 2
    assert rubric["test_type:emergency"]["case_ids"] == ["TC-01", "TC-02"]
    assert rubric["scenario_category:immediate_safety"]["case_ids"] == ["TC-01", "TC-02"]
    assert rubric["scenario_tags:imminent_threat"]["case_ids"] == ["TC-01"]
    assert rubric["scenario_tags:safety_planning"]["case_ids"] == ["TC-01", "TC-02"]

    book = load_workbook(export_results_workbook(result, tmp_path / "results.xlsx"), read_only=True)
    overview = list(book["Overview"].values)
    assert ("總體", "全部案例", "rubric") in [row[:3] for row in overview]
    assert ("scenario_tags", "imminent_threat", "rubric") in [row[:3] for row in overview]
    assert ("scenario_tags", "safety_planning", "rubric") in [row[:3] for row in overview]
    score = list(book["Score Summary"].values)
    assert score[3][:5] == ("Subject", "Judge", "軸", "細分", "指標")
    assert ("s", "j1", "scenario_tags", "safety_planning", "rubric") in [row[:5] for row in score]
    assert not any("own_complete_cases" in str(row) or "common_complete_cases" in str(row) for row in score)
    book.close()


def test_real_receipts_bind_stage_and_preserve_failure():
    from xiaoan_eval_core.runtime import ResponseStore
    store = ResponseStore(max_attempts=1)
    request = {"task": "rubric", "answer_id": "a", "binding": "string-digest", "identity": {"provider": "fixture"}}
    with pytest.raises(LookupError):
        store.call(request, None, lambda payload, request: payload)
    source = fixture_result()
    result = build_complete_results(source["plan"], source["answers"], {"inventories": source["inventories"], "cells": source["envelopes"][0]["assessments"]}, source["envelopes"][0]["rubric"], [], [], store.receipts)
    assert result["envelopes"][0]["stage_refs"] == [store.receipts[0]["stage_id"]]
    assert result["stages"][0]["reason"] == "PROVIDER_NOT_CONFIGURED"


@pytest.mark.parametrize('answer', ['- bullet answer', '=SUM(1,2)', '+1 is text', '@literal', '長' * 40000])
def test_review_literal_and_long_answer_roundtrip(tmp_path, answer):
    from xiaoan_eval.frozen_review import read_review_workbook, import_review_rows
    result = fixture_result()
    # Keep binding evidence valid while exercising displayed answer representation.
    result['answers'][0]['answer'] = answer
    result['answers'][0]['question'] = '=literal question'
    claim = result['inventories'][0]['claims'][0]
    claim['answer_span'] = {'start': 0, 'end': 1, 'text': answer[:1]}
    seal_complete_results(result)
    output = export_results_workbook(result, tmp_path / 'review.xlsx')
    book = load_workbook(output)
    sheet = book['Human Review']
    headers = [c.value for c in sheet[1]]
    for number in range(2, sheet.max_row + 1):
        assert sheet.cell(number, headers.index('answer') + 1).data_type == 's'
        assert sheet.cell(number, headers.index('question') + 1).data_type == 's'
        for field, value in [('decision', 'APPROVE'), ('reviewer', 'human'), ('reviewed_at', '2026-09-24T10:00:00+08:00')]:
            sheet.cell(number, headers.index(field) + 1, value)
    book.save(output)
    revised = import_review_rows(result, read_review_workbook(output), confirmed_by='human')
    assert revised['answers'][0]['answer'] == answer
    assert revised['envelopes'][0]['human_review'][0]['decision'] == 'APPROVE'
