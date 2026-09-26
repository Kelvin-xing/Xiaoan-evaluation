"""Full Minimal33-shaped pipeline with explicit synthetic provider, never model evidence."""
import json
from pathlib import Path
from openpyxl import load_workbook
from xiaoan_eval.frozen_ingress import build_plan, freeze_answer
from xiaoan_eval.frozen_cli import execute_frozen, generate_report
from xiaoan_eval.frozen_export import SHEETS
from xiaoan_eval_core.results import validate_complete_results

IDENTITY={'id':'offline-fixture','provider':'fixture','model':'fixture','prompt_version':'fixture/v1'}

class FixtureProvider:
    def __call__(self,request):
        task=request['task']
        if task=='extract_claims':
            value={'binding':request['binding'],'claims':[{'id':'c1','kind':'FACTUAL','proposition':request['answer'],
                'conditions':[],'answer_span':{'start':0,'end':len(request['answer']),'text':request['answer']}}]}
        elif task=='assess_claims':
            requirements=[]
            for r in request['requirements']:
                obs=request.get('observations',{}).get(r['id'],{})
                verdict='UNCERTAIN'
                if obs.get('status')=='NOT_APPLICABLE':verdict='NOT_APPLICABLE'
                elif isinstance(obs.get('value'),bool):verdict='SATISFIED' if obs['value'] else 'VIOLATED'
                requirements.append({'id':r['id'],'verdict':verdict,'reason':'離線替身不判定內容品質','answer_spans':[]})
            value={'binding':request['binding'],'claims':[{'id':c['id'],**{d:{'verdict':'UNKNOWN','evidence':[],'reason':'離線測試無獨立證據'} for d in ('faithfulness','correctness')}} for c in request['inventory']['claims']], 'requirements':requirements}
        elif task=='rubric':
            value={'dimensions':[{'module':m['name'],'score':2,'reason':'合成測試固定分數，不是模型品質','supporting_evidence':[],'deduction_evidence':[]} for m in request['rating_rule']['modules']],
                'red_lines':[{'id':r['id'],'triggered':False,'reason':'合成測試判定','evidence':[]} for r in request['rating_rule']['red_lines']]}
        elif task=='relevancy_generation':
            value={'questions':['合成問題'+str(i) for i in range(request.get('n',3))]}
        else:raise ValueError(task)
        return {'payload':value,'usage':{'input_tokens':1,'output_tokens':1,'total_tokens':2}}
    def embeddings(self,texts,task):
        return {'model_id':'fixture','revision':'v1','vectors':{t:[1.,float(len(t)),1.] for t in texts}}


class FixtureReport:
    model='fixture';endpoint='offline'
    def __call__(self,instructions,payload):
        if not payload['history']:
            value={'action':'inspect','requests':[{'tool':'read_evidence','args':{'ref':'/answers/0'}}]}
        else:
            value={'action':'finish','generation':payload['catalog']['generation'],'title':'離線端到端流程驗證（非模型品質報告）',
                'findings':[{'kind':'fact','conclusion':'此結果僅驗證合成資料流轉。','scope':'offline fixture',
                'quotes':[{'ref':'/answers/0','text':'這是離線合成回答。'}]}],
                'facts':[],'limitations':['未呼叫真實模型；分數、回答和 embedding 全部為測試替身。']}
        return json.dumps(value,ensure_ascii=False),{}


def run_fixture(output,case_ids=None):
    spec=build_plan([IDENTITY],[IDENTITY],IDENTITY,case_ids=case_ids,relevancy_generator=IDENTITY)
    spec['rows']=[freeze_answer(r,{'text':'這是離線合成回答。'},[],generation_id='synthetic-20260924') for r in spec['rows']]
    result=execute_frozen(spec,output,provider=FixtureProvider(),max_workers=2)
    generate_report(Path(output)/'results.json',Path(output)/'report',provider=FixtureReport())
    return result


def test_frozen_pipeline_to_current_sheets_and_report(tmp_path):
    result=run_fixture(tmp_path,case_ids=['TC-35'])
    validate_complete_results(result)
    assert len(result['answers'])==2
    wb=load_workbook(tmp_path/'results.xlsx',read_only=True)
    assert tuple(wb.sheetnames)==SHEETS and 'Human Review' in wb.sheetnames
    assert (tmp_path/'report/report.md').exists()
    assert all(e['rubric'][0]['status']=='AVAILABLE' for e in result['envelopes'])
    assert all(e['assessments'][0]['status']=='AVAILABLE' for e in result['envelopes'])
    assert all(e['relevancy']['status']=='AVAILABLE' for e in result['envelopes'])
