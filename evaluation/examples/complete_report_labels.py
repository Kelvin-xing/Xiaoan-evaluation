"""Offline label completion for the two inspected runs; never approves an oracle.

Writes derived workbooks and an identical 96-turn evidence pack to both projects.
Existing answers, scores, test cases, and historical workbooks are not modified.
"""
from pathlib import Path
import copy
import csv
import hashlib
import json
import re
import sys

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'evaluation'))
from xiaoan_eval.scenario_analysis import AXES, _groups, digest
from xiaoan_eval.detail_tables import write_table

PACK = 'report-labels-2026-09-22'
SOURCE = ROOT / 'evaluation/oracles/minimal32-remediation/scenario-analysis/turn-annotations.json'
SINGLE = ROOT / 'evaluation/runs/2026-09-22-minimal32-readable-details/results.xlsx'
MATRIX = ROOT / 'evaluation_multimodels/runs/2026-09-21-minimal32-five-self-included/results.xlsx'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def rows(book, name):
    data = book[name].iter_rows(values_only=True)
    headers = next(data)
    return [dict(zip(headers, row)) for row in data]


def cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def table(book, name, data, headers=None):
    headers = headers or list(data[0])
    write_table(book, name, headers, [{k: cell(v) for k, v in r.items()} for r in data])


def main():
    original_hashes = {str(p.relative_to(ROOT)): sha(p) for p in (SOURCE, SINGLE, MATRIX)}
    annotations = json.loads(SOURCE.read_text())['annotations']
    index = {(a['case_id'], a['turn']): a for a in annotations}
    assert len(index) == len(annotations) == 96
    cases, oracles = {}, []
    for cid in sorted({a['case_id'] for a in annotations}):
        path = ROOT / f'evaluation/test-cases/{cid}.yaml'
        twin = ROOT / f'evaluation_multimodels/test-cases/{cid}.yaml'
        assert path.read_bytes() == twin.read_bytes(), f'case copies differ: {cid}'
        case = yaml.safe_load(path.read_text())
        cases[cid] = case
        history = []
        for turn in case['turns']:
            key = (cid, turn['turn'])
            history.append({'turn': turn['turn'], 'user': turn['user']})
            a = index[key]
            assert a['user_history_sha256'] == digest(history), key
            assert [{'turn': x['turn'], 'text': x['user']} for x in history] == a['user_evidence']
            assert all(isinstance(a[axis], list) for axis in AXES)
            assert all(a[axis] for axis in AXES if axis != 'constraints'), key
            expected = turn['expected']
            ref = expected['reference_oracle']
            assert expected['preferred_route_id'] in expected['route_ids']
            assert set(expected['capsule_ids']) == set(expected['route_ids'])
            assert ref['status'] == 'provisional'
            for contract in ref['route_contracts'].values():
                if contract.get('capsule_ref'):
                    assert (ROOT / contract['capsule_ref']).is_file(), contract
            review = ref.get('semantic_review', {})
            oracles.append({
                'case_id': cid, 'turn': turn['turn'], 'user_history_sha256': digest(history),
                'user_prefix': copy.deepcopy(history), 'status': 'provisional',
                'route_ids': expected['route_ids'], 'preferred_route_id': expected['preferred_route_id'],
                'capsule_ids': expected['capsule_ids'], 'safety_levels': expected['safety_levels'],
                'rationale': ref['rationale'], 'risk_basis': review.get('risk_basis'),
                'allowed_alternatives': review.get('allowed_alternatives', []),
                'review_issue': review.get('review_issue'),
                'human_review_required': review.get('human_review_required', []),
                'route_contracts': ref['route_contracts'], 'ground': ref['ground'],
                'source_gaps': ref.get('source_gaps', []),
                'source_case_path': str(path.relative_to(ROOT)), 'source_case_sha256': sha(path),
                'source_snapshot_id': ref['snapshot_id'],
                'source_authored_by': ref.get('authored_by'), 'source_authored_at': ref.get('authored_at'),
                'completion_method': 'EXISTING_DRAFT_REUSED_WITH_USER_PREFIX_AND_STRUCTURE_VALIDATION',
                'reviewed_by': None, 'reviewed_at': None,
                'release_gate_eligible': False,
            })
    oracle_index = {(r['case_id'], r['turn']): r for r in oracles}
    assert len(cases) == 32 and len(oracles) == 96

    # Verify the matrix's saved user questions, not merely the current YAML IDs.
    saved = (MATRIX.parent / 'answers.md').read_text()
    frozen = {}
    for cid, block in re.findall(r'^## (TC-\d+)\n(.*?)(?=^## TC-|\Z)', saved, re.M | re.S):
        for turn, text in re.findall(r'^### 第 (\d+) 輪\n\n\*\*使用者：\*\* ([^\n]+)', block, re.M):
            frozen[(cid, int(turn))] = text
    assert len(frozen) == 17
    for (cid, turn), text in frozen.items():
        assert next(t['user'] for t in cases[cid]['turns'] if t['turn'] == turn) == text

    results = []
    for project, source in [('evaluation', SINGLE), ('evaluation_multimodels', MATRIX)]:
        book = openpyxl.load_workbook(source)
        if project == 'evaluation':
            # Match all reconstructed visible user text against the frozen prefix.
            texts = {}
            turns = rows(book, '02_Turns')
            for row in turns:
                if row['row_kind'] == 'text' and row.get('role') in {'user', 'user_input'}:
                    texts.setdefault(row['text_id'], []).append((row['chunk_index'], row['content'] or ''))
            for row in turns:
                if row['row_kind'] == 'turn':
                    text = ''.join(v for _, v in sorted(texts[row['user_text_id']]))
                    a = index[(row['case_id'], row['turn'])]
                    assert text == a['user_evidence'][-1]['text']
            keys = sorted(index)
            changed = 0
        else:
            scenario = rows(book, 'Scenario_turns')
            keys = sorted({(r['case_id'], r['turn']) for r in scenario})
            assert set(keys) == set(frozen)
            changed = sum(r['annotation_status'] == 'UNLABELED' for r in scenario)
            decoded = []
            for row in scenario:
                a = index[(row['case_id'], row['turn'])]
                row.update(annotation_status=a['annotation_status'],
                           annotation_reason='USER_PREFIX_VERIFIED_AGAINST_SAVED_RUN_QUESTIONS',
                           tags=json.dumps({axis: a[axis] for axis in AXES}, ensure_ascii=False),
                           need_summary=a['need_summary'])
                item = copy.deepcopy(row)
                for key in ('tags', 'primary_binary', 'strict_entailment_proxy', 'answer_relevancy', 'required_coverage'):
                    item[key] = json.loads(item[key])
                decoded.append(item)
            table(book, 'Scenario_turns', scenario)
            table(book, 'Scenario_groups', _groups(decoded))
            # Answer-only aggregation has no Judge duplication and includes unavailable answers.
            unique = {}
            for r in decoded:
                key = (r['subject_id'], r['case_id'], r['turn'])
                item = copy.deepcopy(r)
                item['judge_id'] = 'ANSWER_LEVEL'
                unique.setdefault(key, item)
            answer_groups = _groups(list(unique.values()))
            headers = ['subject_id', 'judge_id', 'axis', 'label', 'n_cases', 'n_turns',
                       'answered_n', 'annotated_n', 'reviewed_n', 'answer_relevancy']
            table(book, 'Scenario_answer_groups', answer_groups, headers)

        selected_a = [index[k] for k in keys]
        selected_o = [oracle_index[k] for k in keys]
        table(book, 'Scenario_Annotations', selected_a)
        table(book, 'Router_Oracle_Drafts', selected_o)
        metadata = [
            {'key': 'source_workbook', 'value': str(source.relative_to(ROOT))},
            {'key': 'source_workbook_sha256', 'value': sha(source)},
            {'key': 'annotation_source_sha256', 'value': sha(SOURCE)},
            {'key': 'annotation_unique_turns', 'value': len(keys)},
            {'key': 'router_oracle_unique_turns', 'value': len(keys)},
            {'key': 'matrix_judgement_rows_newly_labeled', 'value': changed},
            {'key': 'formal_router_accuracy', 'value': 'UNAVAILABLE: independent human approval pending'},
            {'key': 'label_scope', 'value': 'User-prefix scenario analysis and existing Router oracle drafts; no new legal/source adjudication'},
            {'key': 'run_binding', 'value': 'User text verified; current oracle draft is retrospective and does not replace the run-time oracle'},
        ]
        table(book, 'Label_Completeness', metadata)
        output = ROOT / project / 'runs' / PACK
        output.mkdir(parents=True, exist_ok=True)
        book.save(output / 'results.xlsx')
        # Core data must remain exactly unchanged, including operational failures.
        old = openpyxl.load_workbook(source, read_only=True)
        new = openpyxl.load_workbook(output / 'results.xlsx', read_only=True)
        excluded = {'Scenario_turns', 'Scenario_groups', 'Scenario_answer_groups'} if project.endswith('multimodels') else set()
        for name in old.sheetnames:
            if name not in excluded:
                assert list(old[name].values) == list(new[name].values), name
        for name in ('Scenario_Annotations', 'Router_Oracle_Drafts'):
            assert new[name].max_row - 1 == len(keys)
        if project.endswith('multimodels'):
            assert all(r['annotation_status'] == 'PROVISIONAL' for r in rows(new, 'Scenario_turns'))
            assert len(rows(new, 'All_Answers')) == 85 and len(rows(new, 'All_Judgements')) == 425
        old.close(); new.close(); book.close()
        results.append({'project': project, 'input': str(source.relative_to(ROOT)),
                        'output': str((output / 'results.xlsx').relative_to(ROOT)),
                        'unique_turns': len(keys), 'newly_labeled_judgement_rows': changed,
                        'original_score_and_answer_sheets_unchanged': True})

    report = {'schema_version': 'report-label-completion/v1', 'cases': 32, 'turns': 96,
              'annotation_prefix_matches': 96, 'router_drafts': 96, 'human_approved_router_drafts': 0,
              'source_hashes': original_hashes, 'outputs': results,
              'scope': 'The inspected Minimal32 single-model run and its five-case multimodel subset'}
    for p, expected_hash in original_hashes.items():
        assert sha(ROOT / p) == expected_hash, f'source changed: {p}'
    for project in ('evaluation', 'evaluation_multimodels'):
        directory = ROOT / project / 'oracles' / PACK
        directory.mkdir(parents=True, exist_ok=True)
        dump(directory / 'turn-annotations.json', {
            'schema_version': 'scenario-annotations/v1', 'annotation_status': 'PROVISIONAL',
            'source_path': str(SOURCE.relative_to(ROOT)), 'source_sha256': sha(SOURCE),
            'completion_method': 'EXISTING_96_ANNOTATIONS_REUSED_WITH_EXACT_PREFIX_VALIDATION',
            'annotations': annotations})
        dump(directory / 'router-oracles.json', {'schema_version': 'router-review-pack/v1', 'oracles': oracles})
        dump(directory / 'validation.json', report)
        fields = ['case_id', 'turn', 'user_prefix', 'route_ids', 'preferred_route_id', 'safety_levels',
                  'rationale', 'risk_basis', 'review_issue', 'status', 'reviewed_by', 'reviewed_at']
        with (directory / 'router-review.csv').open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows({k: cell(r[k]) for k in fields} for r in oracles)
        readme = '''# 報告標籤補齊包（2026-09-22）

範圍：已檢查的 Minimal32 單模型 32 案／96 輪，以及其中五案多模型 17 輪。

- `turn-annotations.json`：96 輪完整情境標籤，沿用既有逐輪草案；與當前兩套 YAML 的用戶前綴完全匹配，可用 `--scenario-annotations` 載入。
- `router-oracles.json`：96 輪可接受路由、首選路由、安全標籤、依據與爭議。為審核包，不是 evaluator 的新輸入格式。
- `router-review.csv`：逐輪人工校對入口；reviewed_by／reviewed_at 留空，未冒用既有 reviewer。
- `validation.json`：來源 hash、資料覆蓋、原始分數及回答不變的驗證結果。
- `../../runs/report-labels-2026-09-22/results.xlsx`：補齊後副本。新增 Scenario_Annotations、Router_Oracle_Drafts、Label_Completeness。多模型 Scenario_turns 及兩張分組表已重建，425 筆 Judge 記錄均有情境標籤；同一回答按 85 份去重聚合。

## 狀態與使用限制

資料完整性核對不等於 domain review。既有情境標籤及 Router oracle 全部保留 PROVISIONAL／provisional；正式 Router 準確率仍為 UNAVAILABLE。這次沒有重新裁定法律、醫療或來源正確性，也沒有重新呼叫 Subject／Judge。空 constraints 表示未標註到明示限制，並非缺失輪次。

單模型 96 輪原本已有標籤；本次把 oracle 草案與其一起帶入工作簿。多模型缺少的是標籤接入，沒有從 Subject 回答或實際路由反推期望。用戶前綴另外與該 run 保存的 answers.md 17 輪問題核對。

Router 草案為事後分析附件，不覆寫歷史 run 的 oracle 或評分；引用的原始 snapshot_id 保留，沒有冒稱是重新審核的當前內容快照。未补齐全部 74 案及歷史所有 runs，也未重建 25 份 pair 工作簿。

人工審核後須明確記錄核准欄位範圍，再更新 source oracle 並離線重新計算適用指標。現有 reference_oracle.status 涵蓋 route、安全與 source 等多欄位，不可只審過 route 就把整個 reference_oracle 一律 approved。

## 重建

在 repo 根目錄執行 `python evaluation/examples/complete_report_labels.py`（需要 PyYAML、openpyxl）。讀入原工作簿，輸出到獨立目录；請勿把私人 runs 或用戶對話直接發布。
'''
        (directory / 'README.zh-HK.md').write_text(readme, encoding='utf-8')
        (ROOT / project / 'runs' / PACK / 'README.md').write_text(readme, encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
